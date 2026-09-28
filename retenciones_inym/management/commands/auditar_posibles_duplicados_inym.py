"""
Auditoría de sólo lectura -- pedido de Gastón (28/09/2026), como red de
seguridad para cuando siga importando Excel de INYM de años anteriores.

Contexto: el importador matchea únicamente por (N° cert. INYM, Tipo de
tarifa) exactos (ver importador.py::importar_filas, `existentes_por_clave`).
Las retenciones viejas que no tenían "N° cert. INYM" cargado ya se
completaron con `corregir_duplicados_alta_simple_inym`, poniéndoles
N° cert. INYM = su propio id -- Gastón confirmó que ésa era la convención
histórica ("antes el id era el mismo o debería ser el mismo que el
número"), así que para la enorme mayoría ese número completado debería
coincidir con el real.

Pero si para alguna retención puntual esa convención NO se cumplió (por la
razón que sea -- carga manual con otro criterio en su momento, error, etc.),
el número que quedó cargado es falso. Cuando se importe el Excel real de ese
período (con el número VERDADERO de INYM), el importador no la va a
reconocer -- va a crear una copia nueva, con el número real. Ese caso
puntual ya NO lo agarra `corregir_duplicados_alta_simple_inym`, porque esa
retención ya tiene *algún* N° de certificado cargado (aunque sea el
equivocado) -- el filtro de ese comando es "sin N° cert. INYM", y acá ya no
aplica.

Este comando busca ese caso por CONTENIDO en lugar de por número: agrupa las
retenciones por (operador retenido, tipo de tarifa, total, fecha) -- los
mismos 4 datos que ya usa `corregir_duplicados_alta_simple_inym` para
encontrar la copia de una retención -- y reporta cualquier grupo con más de
una retención:
  - Si dentro del grupo los N° de certificado son DISTINTOS: es la señal de
    que pasó justo este caso -- dos filas que en el fondo son la misma
    retención, con dos números distintos -- hay que revisar a mano cuál
    conservar (en general, la de origen Excel trae el número real; la vieja
    quedaría para borrar, salvo que ya esté vinculada a una Liquidación).
  - Si coinciden (no debería pasar -- el importador nunca inserta dos filas
    con el mismo (número, tipo de tarifa) -- pero no hay una restricción real
    en la base que lo impida si se cargó a mano): se reporta igual, para no
    dejar pasar nada raro.

No se agrupan las retenciones con total = $0 (pedido de Gastón, 28/09/2026):
un total en $0 es una coincidencia demasiado común como para ser señal de
duplicado -- pueden existir varias retenciones distintas con ese mismo
"total vacío" que coinciden en operador/tipo de tarifa/fecha sin ser en
absoluto la misma retención repetida. Se ignoran sin importar cuántas
aparezcan agrupadas.

Sólo informa -- no borra ni completa nada. Pensado para correrlo después de
cada importación de un Excel de un año anterior, como chequeo de rutina
(es inofensivo si no encuentra nada).

Con `--excel` además genera un .xlsx (pedido de Gastón, 28/09/2026) con TODOS
los datos de cada retención involucrada -- no sólo lo que se ve por consola
-- para poder revisar/filtrar cómodo o mandarlo. Una fila por retención,
agrupadas (columna "Grupo") y con el motivo (números distintos / mismo
número) y si ya está vinculada a una Liquidación.

Uso:
    python manage.py auditar_posibles_duplicados_inym
    python manage.py auditar_posibles_duplicados_inym --fecha-desde 2020-01-01 --fecha-hasta 2020-12-31
    python manage.py auditar_posibles_duplicados_inym --excel
    python manage.py auditar_posibles_duplicados_inym --excel auditoria_2020.xlsx
"""
from collections import defaultdict
from datetime import date

from django.core.management.base import BaseCommand

from retenciones_inym.models import RetencionInym

NOMBRE_EXCEL_DEFAULT = 'auditoria_posibles_duplicados_inym_{fecha}.xlsx'


class Command(BaseCommand):
    help = (
        'Busca, por contenido (operador retenido + tipo de tarifa + total + fecha), posibles '
        'retenciones INYM duplicadas que ya tengan las dos un N° cert. INYM cargado (y que por eso no '
        'detecta corregir_duplicados_alta_simple_inym). Sólo informa, no modifica nada. Con --excel '
        'además genera un .xlsx con todos los datos de las retenciones involucradas.'
    )

    def add_arguments(self, parser):
        parser.add_argument(
            '--fecha-desde', type=str, default=None,
            help='Filtra desde esta fecha, formato AAAA-MM-DD (default: sin límite).',
        )
        parser.add_argument(
            '--fecha-hasta', type=str, default=None,
            help='Filtra hasta esta fecha, formato AAAA-MM-DD (default: sin límite).',
        )
        parser.add_argument(
            '--excel', nargs='?', const='', default=None, metavar='RUTA',
            help='Además del informe por consola, genera un .xlsx con todos los datos de las '
                 'retenciones encontradas. Sin indicar ruta, usa '
                 '"auditoria_posibles_duplicados_inym_<fecha de hoy>.xlsx" en la carpeta actual.',
        )

    def handle(self, *args, **options):
        fecha_desde = date.fromisoformat(options['fecha_desde']) if options['fecha_desde'] else None
        fecha_hasta = date.fromisoformat(options['fecha_hasta']) if options['fecha_hasta'] else None

        # Mismo requisito que corregir_duplicados_alta_simple_inym: sin total
        # y/o fecha no se puede agrupar con confianza. Además (pedido de
        # Gastón, 28/09/2026) se descarta total=0 directamente: un total en
        # $0 no es una coincidencia significativa -- puede haber muchas
        # retenciones distintas con total $0 (por ejemplo eliminadas o
        # corregidas por INYM) que coinciden por casualidad en operador/tipo
        # de tarifa/fecha sin ser para nada la misma retención duplicada. Da
        # lo mismo si el grupo resultante tiene 2 o 9 -- se ignora igual.
        qs = RetencionInym.objects.filter(total__isnull=False, fecha__isnull=False).exclude(total=0)
        if fecha_desde:
            qs = qs.filter(fecha__gte=fecha_desde)
        if fecha_hasta:
            qs = qs.filter(fecha__lte=fecha_hasta)
        qs = qs.select_related('operador_retenido__entidad', 'id_tipo_tarifa')

        grupos = defaultdict(list)
        for r in qs:
            clave = (r.operador_retenido_id, r.id_tipo_tarifa_id, r.total, r.fecha)
            grupos[clave].append(r)

        sospechosos = {clave: miembros for clave, miembros in grupos.items() if len(miembros) > 1}

        if not sospechosos:
            self.stdout.write(self.style.SUCCESS(
                f'No se encontró ningún grupo de posibles duplicadas{self._describir_rango(fecha_desde, fecha_hasta)}.'
            ))
            if options['excel'] is not None:
                self.stdout.write('No se generó ningún .xlsx (no había nada para reportar).')
            return

        from liquidaciones.models import LiquidacionRetencionInym

        self.stdout.write(self.style.WARNING(
            f'{len(sospechosos)} grupo(s) de retenciones con el mismo operador retenido, tipo de tarifa, '
            f'total y fecha (posibles duplicadas){self._describir_rango(fecha_desde, fecha_hasta)}:'
        ))
        for (operador_id, tipo_id, total, fecha), miembros in sorted(sospechosos.items(), key=lambda kv: kv[0][3]):
            miembros = sorted(miembros, key=lambda r: r.id)
            certs = {r.id_certificado_inym for r in miembros}
            etiqueta = 'N° cert. INYM DISTINTOS -- revisar' if len(certs) > 1 else 'mismo N° cert. INYM (raro, revisar igual)'
            retenido = str(miembros[0].operador_retenido) if operador_id else '(sin operador retenido)'
            tipo_tarifa = str(miembros[0].id_tipo_tarifa) if tipo_id else '(sin tipo de tarifa)'
            self.stdout.write('')
            self.stdout.write(
                f'  fecha={fecha}  total={total}  operador_retenido={retenido}  tipo_tarifa={tipo_tarifa}  '
                f'-- {etiqueta}'
            )
            for r in miembros:
                liquidada = LiquidacionRetencionInym.objects.filter(retencion_inym_id=r.id).exists()
                aviso_liquidacion = '  [YA LIQUIDADA]' if liquidada else ''
                self.stdout.write(
                    f'    id={r.id}  cert.={r.id_certificado_inym}  agregado_desde={r.agregado_desde or "(sin dato)"}'
                    f'{aviso_liquidacion}'
                )

        self.stdout.write('')
        self.stdout.write('No se tocó nada -- esto es sólo un informe, revisar y decidir cada caso a mano.')

        if options['excel'] is not None:
            ruta = options['excel'] or NOMBRE_EXCEL_DEFAULT.format(fecha=date.today().isoformat())
            self._generar_excel(sospechosos, ruta)
            self.stdout.write(self.style.SUCCESS(f'Excel generado: {ruta}'))

    @staticmethod
    def _generar_excel(sospechosos, ruta):
        import openpyxl
        from openpyxl.utils import get_column_letter

        from liquidaciones.models import LiquidacionRetencionInym
        from services.gestorexcel import definir_estilo_general

        try:
            from zoneinfo import ZoneInfo
            zona_horaria = ZoneInfo('America/Argentina/Buenos_Aires')
        except ImportError:
            zona_horaria = None

        def a_hora_local(momento):
            # fecha_agregado/fecha_modificado se guardan en UTC (settings.
            # TIME_ZONE del proyecto) -- para que el Excel muestre la hora
            # real en la que se cargó/modificó cada fila (como la vería
            # Gastón), se convierte a hora de Argentina y se le saca el
            # tzinfo antes de escribirla (openpyxl no admite datetimes con
            # zona horaria). Mismo criterio ya usado para imprimir la hora
            # de emisión en solicitudes_compra/documentos.py.
            if momento is None:
                return None
            if zona_horaria is not None:
                momento = momento.astimezone(zona_horaria)
            return momento.replace(tzinfo=None)

        columnas = [
            'Grupo', 'Motivo', 'Id', 'Fecha', 'Período', 'Tipo de tarifa',
            'CUIT emisor', 'Emisor', 'Tipo oper. emisor',
            'CUIT retenido', 'Retenido', 'Tipo oper. retenido',
            'Kgs', 'Tarifa', 'Total', 'Eliminación (INYM)', 'N° cert. INYM',
            'Agregado desde', 'Agregado el', 'Modificado el', 'Vinculada a liquidación',
        ]
        columnas_numericas = {12, 13, 14}  # Kgs, Tarifa, Total

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Posibles duplicadas'
        ws.append(columnas)

        for numero_grupo, (clave, miembros) in enumerate(
            sorted(sospechosos.items(), key=lambda kv: kv[0][3]), start=1
        ):
            miembros = sorted(miembros, key=lambda r: r.id)
            certs = {r.id_certificado_inym for r in miembros}
            motivo = 'N° cert. INYM distintos' if len(certs) > 1 else 'Mismo N° cert. INYM (raro)'
            for r in miembros:
                liquidada = LiquidacionRetencionInym.objects.filter(retencion_inym_id=r.id).exists()
                ws.append([
                    numero_grupo, motivo, r.id, r.fecha, r.periodo,
                    r.id_tipo_tarifa.nombre if r.id_tipo_tarifa_id else '',
                    r.operador_emisor.entidad.cuit if r.operador_emisor_id else '',
                    r.operador_emisor.entidad.nombre if r.operador_emisor_id else '',
                    r.operador_emisor.tipo_operador.nombre if r.operador_emisor_id else '',
                    r.operador_retenido.entidad.cuit if r.operador_retenido_id else '',
                    r.operador_retenido.entidad.nombre if r.operador_retenido_id else '',
                    r.operador_retenido.tipo_operador.nombre if r.operador_retenido_id else '',
                    float(r.kgs) if r.kgs is not None else None,
                    float(r.tarifa) if r.tarifa is not None else None,
                    float(r.total) if r.total is not None else None,
                    r.eliminacion, r.id_certificado_inym, r.agregado_desde or '',
                    a_hora_local(r.fecha_agregado), a_hora_local(r.fecha_modificado),
                    'Sí' if liquidada else 'No',
                ])

        fila_encabezado = 1
        definir_estilo_general(ws)
        for indice in columnas_numericas:
            letra_columna = get_column_letter(indice + 1)
            for celda in ws[letra_columna]:
                if celda.row > fila_encabezado:
                    celda.number_format = '#,##0.00'
        for columna in ws.columns:
            letra = columna[0].column_letter
            largo_max = max((len(str(c.value)) for c in columna if c.value is not None), default=10)
            ws.column_dimensions[letra].width = min(max(largo_max + 2, 10), 40)

        wb.save(ruta)

    @staticmethod
    def _describir_rango(fecha_desde, fecha_hasta):
        if fecha_desde and fecha_hasta:
            return f' entre {fecha_desde} y {fecha_hasta}'
        if fecha_desde:
            return f' desde {fecha_desde}'
        if fecha_hasta:
            return f' hasta {fecha_hasta}'
        return ''
