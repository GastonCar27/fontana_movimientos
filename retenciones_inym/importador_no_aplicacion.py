"""Importador del "Listado Cert. de No Aplicación" que se exporta desde el
sistema de INYM (pedido de Gastón, 02/10/2026). Ej. de columnas:

    IDCERT_NO_APLICACION, IDOPERADOR, IDEMPRESA, TIPO_OPER, FECHA, PERIODO,
    VENCIMIENTO, TOTAL, ID_OPER_VALIDA, TIPO_OPER_VALIDA, OPER_VALIDA,
    FECHA_VALIDACION, FECHA_ELIMINACION

Crea o actualiza (por número = IDCERT_NO_APLICACION) los
CertificadoNoAplicacionInym. Los datos de INYM mandan: si el certificado ya
existía (por ejemplo creado sólo con el número al importar el Excel de
retenciones), se completan/pisan sus datos con los del listado. No toca los
vínculos con retenciones (eso lo hace el importador de retenciones).

Si el listado algún día trae también TIPO / TARIFA / KILOS (o KGS), se
toman; si no, quedan como estaban.
"""
from django.db import transaction

from .importador import (
    ErrorImportacion, _parse_decimal, _parse_fecha, _parse_int, _valor_documento, _valor_texto,
)
from .models import CertificadoNoAplicacionInym

AGREGADO_DESDE_LISTADO = 'listado_cert_inym'

COLUMNAS_REQUERIDAS = ['IDCERT_NO_APLICACION', 'IDOPERADOR', 'FECHA', 'TOTAL']

# columna del Excel -> (campo del modelo, parser)
MAPEO = {
    'IDOPERADOR': ('id_operador_emisor', _parse_int),
    'IDEMPRESA': ('cuit_emisor', _valor_documento),
    'TIPO_OPER': ('tipo_oper_emisor', _valor_texto),
    'FECHA': ('fecha', _parse_fecha),
    'PERIODO': ('periodo', _parse_fecha),
    'VENCIMIENTO': ('vencimiento', _parse_fecha),
    'TOTAL': ('total', _parse_decimal),
    'ID_OPER_VALIDA': ('id_operador_valida', _parse_int),
    'TIPO_OPER_VALIDA': ('tipo_oper_valida', _valor_texto),
    'OPER_VALIDA': ('nombre_valida', _valor_texto),
    'FECHA_VALIDACION': ('fecha_validacion', _parse_fecha),
    'FECHA_ELIMINACION': ('fecha_eliminacion', _parse_fecha),
    'TIPO': ('tipo', _valor_texto),
    'TARIFA': ('tarifa', _parse_decimal),
    'KILOS': ('kgs', _parse_decimal),
    'KGS': ('kgs', _parse_decimal),
}


def _filas_crudas(archivo, nombre_archivo):
    extension = (nombre_archivo or '').rsplit('.', 1)[-1].lower()
    if extension == 'xls':
        import xlrd
        contenido = archivo.read() if hasattr(archivo, 'read') else archivo
        wb = xlrd.open_workbook(file_contents=contenido)
        ws = wb.sheet_by_index(0)
        # xlrd devuelve las fechas de celdas tipo fecha como número de serie;
        # en el export de INYM vienen como texto dd/mm/aaaa, pero por las
        # dudas se convierten.
        filas = []
        for i in range(ws.nrows):
            fila = []
            for j in range(ws.ncols):
                celda = ws.cell(i, j)
                if celda.ctype == xlrd.XL_CELL_DATE:
                    fila.append(xlrd.xldate.xldate_as_datetime(celda.value, wb.datemode))
                else:
                    fila.append(celda.value)
            filas.append(fila)
        return filas
    import openpyxl
    wb = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
    return list(wb.worksheets[0].iter_rows(values_only=True))


def leer_certificados(archivo, nombre_archivo):
    """Devuelve una lista de dicts (uno por certificado) o entradas con
    '_error'. Lanza ErrorImportacion si el archivo no tiene el formato."""
    encabezados = None
    resultado = []
    for numero_fila, fila in enumerate(_filas_crudas(archivo, nombre_archivo), start=1):
        if fila is None:
            continue
        celdas = list(fila)
        if encabezados is None:
            textos = [_valor_texto(c).upper() for c in celdas]
            if 'IDCERT_NO_APLICACION' in textos:
                if 'IDCERTIFICADO' in textos:
                    raise ErrorImportacion(
                        'Este archivo parece el Excel de RETENCIONES de INYM (tiene la columna IDCERTIFICADO). '
                        'Importalo desde "Importar desde Excel de INYM" del listado de retenciones; acá va el '
                        '"Listado Cert. de No Aplicación".'
                    )
                encabezados = {nombre: i for i, nombre in enumerate(textos) if nombre}
                faltantes = [c for c in COLUMNAS_REQUERIDAS if c not in encabezados]
                if faltantes:
                    raise ErrorImportacion(
                        'El Excel no tiene el formato esperado -- faltan las columnas: ' + ', '.join(faltantes)
                    )
            continue
        valor = lambda col: celdas[encabezados[col]] if col in encabezados and encabezados[col] < len(celdas) else None  # noqa: E731
        try:
            numero = _parse_int(valor('IDCERT_NO_APLICACION'))
            if not numero:
                continue
            datos = {'fila_excel': numero_fila, 'numero': numero}
            for columna, (campo, parser) in MAPEO.items():
                if columna in encabezados:
                    datos[campo] = parser(valor(columna))
            resultado.append(datos)
        except ValueError as exc:
            resultado.append({'fila_excel': numero_fila, '_error': str(exc)})
    if encabezados is None:
        raise ErrorImportacion(
            'No se encontró la fila de encabezados (con la columna "IDCERT_NO_APLICACION") en el Excel.'
        )
    return resultado


def importar_certificados(filas):
    resultado = {
        'total_en_archivo': len([f for f in filas if '_error' not in f]),
        'creados': 0,
        'actualizados': 0,
        'sin_cambios': 0,
        'filas_con_error': [(f['fila_excel'], f['_error']) for f in filas if '_error' in f],
    }
    existentes = {c.numero: c for c in CertificadoNoAplicacionInym.objects.all()}
    with transaction.atomic():
        for fila in filas:
            if '_error' in fila:
                continue
            datos = {k: v for k, v in fila.items() if k not in ('fila_excel', 'numero')}
            cert = existentes.get(fila['numero'])
            if cert is None:
                cert = CertificadoNoAplicacionInym.objects.create(
                    numero=fila['numero'], agregado_desde=AGREGADO_DESDE_LISTADO,
                    **{k: (v if v is not None else ('' if isinstance(
                        CertificadoNoAplicacionInym._meta.get_field(k).default, str) else None))
                       for k, v in datos.items()},
                )
                existentes[cert.numero] = cert
                resultado['creados'] += 1
                continue
            cambios = []
            for campo, valor in datos.items():
                if valor is None:
                    valor = '' if isinstance(CertificadoNoAplicacionInym._meta.get_field(campo).default, str) else None
                if getattr(cert, campo) != valor:
                    setattr(cert, campo, valor)
                    cambios.append(campo)
            if cambios:
                cert.save()
                resultado['actualizados'] += 1
            else:
                resultado['sin_cambios'] += 1
    return resultado
