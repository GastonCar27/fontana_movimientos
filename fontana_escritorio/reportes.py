"""
Helpers genéricos para exportar un reporte tabular (columnas + filas ya
armadas por la pantalla) a Excel o PDF, para guardar en un archivo local.

Es la adaptación a escritorio de services/reportes.py del proyecto Django
(mismo formato de 'resultado' y mismo criterio de estilo) -- ahí se arma
una HttpResponse para descargar desde el navegador; acá se escribe
directo a un archivo en disco (la ruta la elige el usuario con un
diálogo "Guardar como" en la pantalla que llama a estas funciones).

'resultado' es siempre un dict con:
    columnas:           lista de encabezados (str)
    filas:              lista de listas, un valor por columna
    columnas_numericas: set() con los índices (0-based) de columnas que son
                         montos/números (se formatean con separador de miles)
    anchos:             (solo para PDF) lista de anchos relativos de columna,
                         mismo largo que 'columnas'; si se omite, todas iguales
"""


def exportar_excel(ruta, resultado):
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font, PatternFill

    columnas = resultado['columnas']
    filas = resultado['filas']
    columnas_numericas = resultado.get('columnas_numericas', set())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Reporte'
    ws.append(columnas)
    fila_encabezado = ws.max_row
    for fila in filas:
        ws.append(fila)

    fuente_encabezado = Font(bold=True, color='FFFFFF')
    relleno_encabezado = PatternFill('solid', fgColor='343A40')
    for celda in ws[fila_encabezado]:
        celda.font = fuente_encabezado
        celda.fill = relleno_encabezado

    for indice in columnas_numericas:
        letra_columna = get_column_letter(indice + 1)
        for celda in ws[letra_columna]:
            if celda.row > fila_encabezado:
                celda.number_format = '#,##0.00'

    for columna in ws.columns:
        letra = columna[0].column_letter
        largo_max = max((len(str(c.value)) for c in columna if c.value is not None), default=10)
        ws.column_dimensions[letra].width = min(max(largo_max + 2, 10), 40)

    wb.save(ruta)


def _formatear_miles(valor):
    try:
        return f'{float(valor):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    except (TypeError, ValueError):
        return str(valor)


def exportar_pdf(ruta, titulo, resultado):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    columnas = resultado['columnas']
    filas = resultado['filas']
    columnas_numericas = resultado.get('columnas_numericas', set())
    anchos_relativos = resultado.get('anchos') or [1] * len(columnas)

    margen = 1 * cm
    doc = SimpleDocTemplate(
        ruta, pagesize=landscape(A4),
        topMargin=margen, bottomMargin=margen, leftMargin=margen, rightMargin=margen,
    )
    ancho_disponible = landscape(A4)[0] - doc.leftMargin - doc.rightMargin

    estilos = getSampleStyleSheet()
    estilo_celda = ParagraphStyle('celda_reporte', parent=estilos['Normal'], fontSize=7, leading=8.5)
    estilo_encabezado = ParagraphStyle(
        'encabezado_reporte', parent=estilo_celda, textColor=colors.white, fontName='Helvetica-Bold',
    )

    elementos = [Paragraph(titulo, estilos['Title']), Spacer(1, 0.4 * cm)]

    def formatear_valor(indice, valor):
        if valor is None or valor == '':
            return ''
        if indice in columnas_numericas:
            return _formatear_miles(valor)
        return str(valor)

    fila_encabezado = [Paragraph(str(col), estilo_encabezado) for col in columnas]
    datos = [fila_encabezado]
    for fila in filas:
        datos.append([
            Paragraph(formatear_valor(indice, valor), estilo_celda)
            for indice, valor in enumerate(fila)
        ])

    total_relativo = sum(anchos_relativos) or 1
    col_widths = [ancho_disponible * (peso / total_relativo) for peso in anchos_relativos]

    tabla = Table(datos, colWidths=col_widths, repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#343a40')),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]))
    elementos.append(tabla)
    doc.build(elementos)
