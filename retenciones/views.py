from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from types import SimpleNamespace

from django.contrib import messages
from django.db import IntegrityError, transaction
from django.db.models import Count, Max, Q, Sum
from django.http import JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse

from comprobantes.models import Comprobante, ComprobanteTipo
from entidades.models import Entidad
from services.ordenamiento import aplicar_orden_lista, aplicar_orden_queryset
from services.permisos import requiere_grupo
from services.reportes import excel_response, pdf_response

from .forms import (
    RankingEntidadesForm,
    RetencionHeaderForm,
    RetencionRenglonFormSet,
    RetencionTipoImpuestoForm,
    RetencionTipoRegimenForm,
)
from .models import Retencion, RetencionRenglon, RetencionTipoImpuesto, RetencionTipoRegimen

# Textos que cambian según la dirección de la retención (Retencion.es_emisor
# -- ver el comentario del campo en models.py). Mismo criterio que
# liquidaciones/documentos.py::_TEXTOS_POR_TIPO.
_TEXTOS_POR_DIRECCION = {
    Retencion.ES_EMISOR: {
        'entidad_label': 'Proveedor',
        'titulo': 'CONSTANCIA DE RETENCIÓN',
        'hoja_excel': 'Constancia de Retención',
        'mostrar_firma': True,
    },
    Retencion.NO_ES_EMISOR: {
        'entidad_label': 'Cliente',
        'titulo': 'RETENCIÓN RECIBIDA (registro interno)',
        'hoja_excel': 'Retención recibida',
        'mostrar_firma': False,
    },
}


def _textos_retencion(es_emisor):
    return _TEXTOS_POR_DIRECCION.get(es_emisor, _TEXTOS_POR_DIRECCION[Retencion.ES_EMISOR])


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _siguiente_id_retencion():
    """La tabla 'retencion' no tiene AUTO_INCREMENT en 'id', así que el
    próximo id se calcula a mano (mismo patrón que
    liquidaciones.views._siguiente_id_liquidacion)."""
    ultimo = Retencion.objects.aggregate(Max('id'))['id__max'] or 0
    return ultimo + 1


def _siguiente_numero_retencion(anio):
    ultimo = Retencion.objects.filter(año=anio).aggregate(Max('numero'))['numero__max'] or 0
    return ultimo + 1


def _formatear_cuit(cuit):
    """Devuelve el CUIT separado como XX-XXXXXXXX-X. Si no tiene el formato
    esperado (11 dígitos), lo muestra tal cual está guardado."""
    if not cuit:
        return ''
    digitos = ''.join(c for c in str(cuit) if c.isdigit())
    if len(digitos) == 11:
        return f'{digitos[:2]}-{digitos[2:10]}-{digitos[10:]}'
    return str(cuit)


def _formatear_domicilio(entidad):
    if not entidad:
        return ''
    partes = [entidad.direccion, entidad.localidad, entidad.provincia]
    domicilio = ' - '.join(p for p in partes if p)
    if entidad.codpos:
        domicilio = f'{domicilio} ({entidad.codpos})' if domicilio else f'({entidad.codpos})'
    return domicilio


def _combinar_comprobante_origen(punto_venta, numero_comprobante):
    if punto_venta is None or numero_comprobante is None:
        return ''
    return f'{int(punto_venta):05d}-{int(numero_comprobante):08d}'


def _parsear_comprobante_origen(valor):
    """Inverso de _combinar_comprobante_origen, para precargar el form de
    Modificación a partir de lo guardado en comprobante_origen."""
    if not valor or '-' not in valor:
        return None, None
    izquierda, derecha = valor.split('-', 1)
    try:
        return int(izquierda), int(derecha)
    except ValueError:
        return None, None


def _calcular_total(subtotal, porcentaje):
    if subtotal is None or porcentaje is None:
        return None
    return (Decimal(str(subtotal)) * Decimal(str(porcentaje)) / Decimal('100')).quantize(
        Decimal('0.01'), rounding=ROUND_HALF_UP
    )


def _parse_decimal_post(valor):
    """Convierte un valor de POST (puede venir con coma decimal) a Decimal,
    o None si viene vacío/inválido -- para los campos numéricos del form de
    vincular/editar renglón (neto_gravado, total), que no pasan por un
    Django Form."""
    if valor is None:
        return None
    valor = str(valor).strip().replace(',', '.')
    if not valor:
        return None
    try:
        return Decimal(valor)
    except InvalidOperation:
        return None


def _parse_float_post(valor):
    if valor is None:
        return None
    valor = str(valor).strip().replace(',', '.')
    if not valor:
        return None
    try:
        return float(valor)
    except ValueError:
        return None


class _SumaRenglonesExcedeTotal(Exception):
    """Se lanza (y se atrapa) dentro de una transacción para deshacer un
    vincular/editar de renglón si, después de esa operación, la suma de los
    renglones de la retención superaría su total (pedido de Gastón,
    23/09/2026: la suma NUNCA puede superar el total, pero sí puede quedar
    por debajo -- puede haber más renglones por cargar)."""
    def __init__(self, suma, total):
        self.suma = suma
        self.total = total
        super().__init__(f'La suma de renglones (${suma}) superaría el total de la retención (${total}).')


def _chequear_suma_renglones(retencion_obj):
    """Suma los RetencionRenglon ya guardados de esta retención (se llama
    DESPUÉS de crear/editar el renglón en curso, dentro de la misma
    transacción) y lanza _SumaRenglonesExcedeTotal si superan el total de la
    Retencion. Devuelve (suma, total_retencion) si no hay problema."""
    suma = RetencionRenglon.objects.filter(retencion=retencion_obj).aggregate(s=Sum('total'))['s'] or Decimal('0')
    total_retencion = retencion_obj.total
    if total_retencion is not None and suma > total_retencion:
        raise _SumaRenglonesExcedeTotal(suma, total_retencion)
    return suma, total_retencion


def _tiene_liquidacion(retencion_id):
    """True si esta retención ya está incluida en una liquidación
    (liquidacion_retencion tiene ON DELETE RESTRICT hacia retencion: borrarla
    para 'reemplazarla' en una modificación rompería esa liquidación, así que
    no se permite)."""
    from liquidaciones.models import LiquidacionRetencion
    return LiquidacionRetencion.objects.filter(retencion_id=retencion_id).exists()


def _tipos_comprobante_por_id():
    return {t.id: t for t in ComprobanteTipo.objects.all()}


def _regimen_impuesto_map():
    """Mapa {id_regimen: [id_impuesto, ...]} para que el JS del form
    condicione el combo de Régimen según el Impuesto elegido (y
    viceversa). Un régimen puede estar vinculado a varios impuestos a la
    vez (ver RetencionTipoRegimen.impuestos en models.py)."""
    mapa = {}
    for r in RetencionTipoRegimen.objects.prefetch_related('impuestos'):
        ids = [str(i.id) for i in r.impuestos.all()]
        if not ids and r.impuesto_id is not None:
            # Régimen viejo, todavía no migrado a la relación M2M nueva.
            ids = [str(r.impuesto_id)]
        mapa[str(r.id)] = ids
    return mapa


def _texto_comprobante_origen(retencion, tipos_por_id):
    tipo = tipos_por_id.get(retencion.tipo_comp_origen)
    abreviatura = tipo.abreviatura if tipo and tipo.abreviatura else ''
    origen = retencion.comprobante_origen or ''
    return f'{abreviatura} {origen}'.strip()


def _texto_comprobante(comprobante):
    """Texto para mostrar un Comprobante REAL en el buscador de 'Detalle de
    las operaciones' y para precargar el campo de sólo-lectura de un
    RetencionRenglon ya vinculado -- mismo formato en los dos lugares
    (antes sólo vivía adentro de comprobante_buscar_para_retencion)."""
    if not comprobante:
        return ''
    from movimientos.templatetags.movimientos_extras import separador_miles

    identificador = _combinar_comprobante_origen(
        comprobante.punto_de_venta, comprobante.numero
    ) or (comprobante.comprobante_string or '')
    tipo = comprobante.tipo_comprobante
    partes = [
        (tipo.abreviatura or tipo.nombre) if tipo else None,
        identificador or None,
        comprobante.fecha.strftime('%d/%m/%Y') if comprobante.fecha else None,
        f'${separador_miles(comprobante.total)}' if comprobante.total is not None else None,
    ]
    return ' - '.join(p for p in partes if p)


def _resolver_comprobante_legacy(retencion):
    """Puente para migrar de forma transparente una Retencion VIEJA (de
    antes del rediseño de 24/09/2026, con sus propios tipo_comp_origen/
    comprobante_origen/fecha_comp_origen sueltos y sin ningún
    RetencionRenglon todavía) al modelo nuevo: intenta encontrar el
    Comprobante real que le corresponde, con el mismo criterio de
    dirección/entidad que ya usa comprobante_buscar_para_retencion. Si lo
    encuentra, Modificación puede guardarse sin que el usuario tenga que
    volver a buscar y elegir la factura a mano. Devuelve None si no hay
    datos suficientes o no hay ningún Comprobante que coincida exacto (nunca
    inventa un vínculo) -- en ese caso el usuario tiene que buscarla y
    elegirla de nuevo."""
    if not retencion.entidad_id or not retencion.comprobante_origen:
        return None
    punto_venta, numero_comprobante = _parsear_comprobante_origen(retencion.comprobante_origen)
    if punto_venta is None or numero_comprobante is None:
        return None

    es_emisor_retencion = retencion.es_emisor if retencion.es_emisor is not None else Retencion.ES_EMISOR
    qs = Comprobante.objects.filter(
        entidad_emisor_id=retencion.entidad_id,
        punto_de_venta=punto_venta,
        numero=numero_comprobante,
    )
    if es_emisor_retencion == Retencion.ES_EMISOR:
        qs = qs.filter(Q(es_emisor=1) | Q(es_emisor__isnull=True))
    else:
        qs = qs.filter(es_emisor=0)
    if retencion.tipo_comp_origen:
        qs = qs.filter(tipo_comprobante_id=retencion.tipo_comp_origen)
    return qs.select_related('tipo_comprobante').first()


def _armar_formset_inicial(retencion):
    """Arma el initial del formset de renglones para Modificación.

    Caso normal (retención creada con la pantalla nueva, o ya migrada):
    un renglón por cada RetencionRenglon ya vinculado.

    Caso retención VIEJA (de antes del rediseño de 24/09/2026: un solo
    renglón suelto en la propia Retencion, todavía sin ningún
    RetencionRenglon) -- para no perder esos datos al abrir Modificar, se
    arma UN renglón inicial con lo que ya tenía cargado, intentando
    resolver el Comprobante real que le corresponde (ver
    _resolver_comprobante_legacy) para que quede prellenado y listo para
    guardar sin tocar nada; si no se pudo resolver, se prellena igual el
    texto (para no mostrar la fila en blanco) pero el usuario tiene que
    buscar y elegir la factura de nuevo antes de poder guardar."""
    renglones = (
        RetencionRenglon.objects.filter(retencion=retencion)
        .select_related('comprobante', 'comprobante__tipo_comprobante')
        .order_by('id')
    )
    inicial = []
    for rn in renglones:
        c = rn.comprobante
        inicial.append({
            'comprobante': c.id if c else None,
            'comprobante_texto': _texto_comprobante(c),
            'tipo_comp_origen': c.tipo_comprobante_id if c else None,
            'punto_venta': c.punto_de_venta if c else None,
            'numero_comprobante': c.numero if c else None,
            'fecha_comp_origen': c.fecha if c else None,
            'subtotal': rn.neto_gravado,
            'porcentaje': rn.porcentaje,
            'total': rn.total,
        })
    if inicial:
        return inicial

    tiene_dato_legacy = retencion.subtotal is not None or retencion.tipo_comp_origen is not None \
        or retencion.comprobante_origen
    if not tiene_dato_legacy:
        return []

    comprobante_legacy = _resolver_comprobante_legacy(retencion)
    punto_venta, numero_comprobante = _parsear_comprobante_origen(retencion.comprobante_origen)
    if comprobante_legacy:
        texto = _texto_comprobante(comprobante_legacy)
    else:
        texto = (
            f'{_texto_comprobante_origen(retencion, _tipos_comprobante_por_id())} '
            '(factura no vinculada todavía -- buscala y elegila de nuevo)'
        ).strip()
    return [{
        'comprobante': comprobante_legacy.id if comprobante_legacy else None,
        'comprobante_texto': texto,
        'tipo_comp_origen': retencion.tipo_comp_origen,
        'punto_venta': punto_venta,
        'numero_comprobante': numero_comprobante,
        'fecha_comp_origen': retencion.fecha_comp_origen,
        'subtotal': retencion.subtotal,
        'porcentaje': retencion.porcentaje,
        'total': retencion.total,
    }]


def comprobante_buscar_para_retencion(request):
    """Devuelve, en JSON, comprobantes ya cargados para elegir como renglón
    de 'Detalle de las operaciones' en vez de tipear a mano tipo/punto de
    venta/número/fecha/importe. Se filtran según la entidad elegida en el
    encabezado y la dirección de la retención (?entidad= id de Entidad,
    ?es_emisor= '1' o '0', mismos valores que Retencion.es_emisor):

      - Practicada por Fontana (es_emisor=1): comprobantes que la entidad
        NOS emitió (Comprobante.entidad_emisor=entidad, con
        Comprobante.es_emisor=1 o vacío -- comportamiento histórico) --
        son las facturas de compra sobre las que Fontana le retiene al
        pagarle.
      - Sufrida (es_emisor=0): comprobantes que NOSOTROS le emitimos a la
        entidad (Comprobante.entidad_emisor=entidad,
        Comprobante.es_emisor=0) -- son las facturas de venta sobre las
        que la entidad nos retiene al pagarnos.

    Ojo: Comprobante.entidad_emisor SIEMPRE guarda la contraparte (nunca a
    Fontana); Comprobante.es_emisor es lo que indica si esa entidad fue la
    que realmente emitió el comprobante, o la que lo recibió (nosotros lo
    emitimos). Mismo criterio direccional que ya usa
    liquidaciones.views._armar_items para ofrecer comprobantes según el
    tipo (pago/cobro) de una liquidación."""
    entidad_id = request.GET.get('entidad', '').strip()
    es_emisor_retencion = request.GET.get('es_emisor', '').strip()
    q = request.GET.get('q', '').strip()

    if not entidad_id.isdigit() or es_emisor_retencion not in ('0', '1'):
        return JsonResponse({'resultados': []})

    comprobantes = Comprobante.objects.filter(entidad_emisor_id=int(entidad_id))
    if es_emisor_retencion == '1':
        # Practicada: la entidad nos emitió el comprobante (compra).
        comprobantes = comprobantes.filter(Q(es_emisor=1) | Q(es_emisor__isnull=True))
    else:
        # Sufrida: nosotros le emitimos el comprobante a la entidad (venta).
        comprobantes = comprobantes.filter(es_emisor=0)

    if q:
        filtro = Q(comprobante_string__icontains=q) | Q(tipo_comprobante__nombre__icontains=q)
        if q.isdigit():
            filtro |= Q(numero=int(q)) | Q(id=int(q))
        comprobantes = comprobantes.filter(filtro)

    comprobantes = comprobantes.select_related('tipo_comprobante').order_by('-fecha', '-id')[:20]

    resultados = []
    for c in comprobantes:
        resultados.append({
            'id': c.id,
            'text': _texto_comprobante(c),
            'tipo_comp_origen': c.tipo_comprobante_id,
            'punto_venta': c.punto_de_venta,
            'numero_comprobante': c.numero,
            'fecha_comp_origen': c.fecha.isoformat() if c.fecha else '',
            # El "Importe" que se precarga es la base sobre la que se
            # calcula la retención: el neto gravado del comprobante (no el
            # total, que incluye IVA y no es la base imponible).
            'subtotal': str(c.neto_gravado) if c.neto_gravado is not None else '',
        })
    return JsonResponse({'resultados': resultados})


# ---------------------------------------------------------------------------
# Alta / Modificación (comparten casi toda la lógica de guardado)
# ---------------------------------------------------------------------------

def _guardar_grupo(header_form, formset):
    """Crea UNA Retencion (el encabezado -- un solo comprobante/certificado,
    con su propio Total cargado a mano) y un RetencionRenglon por cada línea
    con datos del formset, vinculado al Comprobante real elegido en el
    buscador (rediseño de 24/09/2026, pedido de Gastón: "cuando cargo una
    retención con dos renglones me debería guardar solo una retención con
    el total, y el vínculo nomás debería ser con dos renglones distintos" --
    antes cada renglón del formset generaba su propia Retencion suelta).

    Devuelve (retencion, cantidad_de_renglones_creados). Las filas marcadas
    "Quitar" (DELETE) se ignoran -- antes del rediseño esa marca sólo
    afectaba la suma que se mostraba en pantalla (JS), pero el guardado
    igual las creaba; se corrige acá de paso, ya que se estaba reescribiendo
    este mismo loop."""
    entidad = header_form.cleaned_data.get('entidad')
    entidad_nombre = header_form.cleaned_data.get('entidad_nombre') or (entidad.nombre if entidad else '')
    id_impuesto = header_form.cleaned_data.get('id_impuesto')
    id_regimen = header_form.cleaned_data.get('id_regimen')
    es_emisor = header_form.cleaned_data.get('es_emisor', Retencion.ES_EMISOR)
    anio = header_form.cleaned_data['año']
    numero = header_form.cleaned_data['numero']
    total_header = header_form.cleaned_data.get('total')

    # Fecha de la retención en sí, cargada a mano en el encabezado (no se
    # deriva más de la fecha del comprobante origen de cada renglón -- son
    # dos cosas distintas, ver comentario en RetencionHeaderForm.fecha).
    fecha_retencion = header_form.cleaned_data.get('fecha')

    retencion = Retencion.objects.create(
        id=_siguiente_id_retencion(),
        entidad=entidad,
        entidad_nombre=entidad_nombre,
        es_emisor=es_emisor,
        total=total_header,
        comprobante_string=f'{anio}-{numero:04d}',
        fecha=fecha_retencion,
        id_impuesto=id_impuesto,
        id_regimen=id_regimen,
        año=anio,
        numero=numero,
        agregado_desde='retenciones_app',
    )

    cantidad = 0
    for form in formset:
        datos = form.cleaned_data
        if not datos or datos.get('_vacio') or datos.get('DELETE'):
            continue
        comprobante = datos.get('comprobante')
        neto_gravado = datos.get('subtotal')
        porcentaje = datos.get('porcentaje')
        # El campo "Retención" (total) se autocompleta en el JS con Importe
        # x Porcentaje / 100, pero queda editable a mano (puede haber una
        # diferencia de centavos con lo que realmente retuvo la otra
        # parte). Se respeta lo que vino cargado en el form; sólo se
        # recalcula acá como red de seguridad si llegara vacío.
        total_renglon = datos.get('total')
        if total_renglon is None:
            total_renglon = _calcular_total(neto_gravado, porcentaje)

        RetencionRenglon.objects.create(
            retencion=retencion,
            comprobante=comprobante,
            neto_gravado=neto_gravado,
            porcentaje=porcentaje,
            total=total_renglon,
        )
        cantidad += 1

    return retencion, cantidad


def retencion_alta(request):
    if request.method == 'POST':
        header_form = RetencionHeaderForm(request.POST)
        formset = RetencionRenglonFormSet(request.POST, prefix='form')

        if header_form.is_valid() and formset.is_valid():
            lineas_validas = [
                f for f in formset
                if f.cleaned_data and not f.cleaned_data.get('_vacio') and not f.cleaned_data.get('DELETE')
            ]
            if not lineas_validas:
                messages.error(request, 'Cargá al menos un renglón con los datos de la operación.')
            else:
                try:
                    with transaction.atomic():
                        retencion, cantidad = _guardar_grupo(header_form, formset)
                        suma, total_retencion = _chequear_suma_renglones(retencion)
                except IntegrityError:
                    messages.error(
                        request,
                        'Una de las facturas elegidas ya tiene esta misma retención (mismo impuesto y '
                        'régimen) vinculada en otra retención -- no se puede repetir. Revisá los '
                        'renglones cargados.'
                    )
                except _SumaRenglonesExcedeTotal as exc:
                    messages.error(
                        request,
                        f'No se guardó: la suma de los renglones (${exc.suma}) superaría el Total de la '
                        f'retención (${exc.total}). Corregí el Total del encabezado o los renglones.'
                    )
                else:
                    anio = header_form.cleaned_data['año']
                    numero = header_form.cleaned_data['numero']
                    accion = request.POST.get('accion')
                    messages.success(
                        request,
                        f'La retención {anio}-{numero:04d} se guardó correctamente '
                        f'({cantidad} renglón{"es" if cantidad != 1 else ""} vinculado'
                        f'{"s" if cantidad != 1 else ""}).'
                    )
                    if total_retencion is not None and suma < total_retencion:
                        messages.warning(
                            request,
                            f'La suma de los renglones (${suma}) todavía no llega al Total cargado '
                            f'(${total_retencion}). Podés agregar los renglones que falten después, '
                            'volviendo a Modificar esta retención.'
                        )
                    if accion == 'pdf':
                        return retencion_pdf(request, retencion.id)
                    if accion == 'excel':
                        return retencion_excel(request, retencion.id)

                    # "Guardar" (sin ir a PDF/Excel): en vez de mandar al
                    # listado, se vuelve a mostrar la propia alta lista para
                    # cargar la siguiente retención -- pedido de Gastón
                    # (25/09/2026): "por lo general se tienen varias
                    # retenciones a la vez para cargar en la misma dirección
                    # y fecha". Se conservan Entidad/Dirección/Fecha/Año (lo
                    # que se repite) y sólo se limpian Número (recalculado al
                    # siguiente disponible de ese año), Impuesto/Régimen y
                    # Total, más los renglones. Sigue hasta el render() de
                    # abajo (mismo que usa el GET inicial).
                    datos_previos = header_form.cleaned_data
                    header_form = RetencionHeaderForm(initial={
                        'entidad': datos_previos.get('entidad'),
                        'entidad_nombre': datos_previos.get('entidad_nombre'),
                        'es_emisor': datos_previos.get('es_emisor'),
                        'fecha': datos_previos.get('fecha'),
                        'año': datos_previos.get('año'),
                        'numero': _siguiente_numero_retencion(datos_previos.get('año')),
                    })
                    formset = RetencionRenglonFormSet(prefix='form')
    else:
        hoy = __import__('datetime').date.today()
        anio_actual = hoy.year
        header_form = RetencionHeaderForm(initial={
            'año': anio_actual,
            'numero': _siguiente_numero_retencion(anio_actual),
            'es_emisor': Retencion.ES_EMISOR,
            'fecha': hoy,
        })
        formset = RetencionRenglonFormSet(prefix='form')

    return render(request, 'retenciones/retencion_form.html', {
        'header_form': header_form,
        'formset': formset,
        'modo': 'alta',
        'regimen_impuesto_map': _regimen_impuesto_map(),
    })


def _retencion_vincular_renglones(request, id, lineas, renglones_nuevos):
    """Modo restringido de Modificación: sólo permite vincular (o quitar)
    RetencionRenglon reales para una retención que ya tiene liquidación
    asignada -- ver el comentario en retencion_modificar. No toca el total ni
    ningún otro dato de la Retencion; sólo agrega/borra filas de la tabla
    nueva `retencion_renglon`."""
    primera = lineas[0]

    if request.method == 'POST':
        accion = request.POST.get('accion')

        if accion == 'vincular_renglon':
            retencion_id = request.POST.get('retencion_id', '')
            comprobante_id = request.POST.get('comprobante_id', '')
            retencion_obj = next((l for l in lineas if str(l.id) == retencion_id), None)

            if not retencion_obj:
                messages.error(request, 'Renglón inválido.')
            elif not comprobante_id:
                messages.error(request, 'Elegí un comprobante de la lista antes de vincular.')
            else:
                comprobante = Comprobante.objects.filter(pk=comprobante_id).first()
                if not comprobante:
                    messages.error(request, 'No se encontró el comprobante elegido.')
                else:
                    # El neto/porcentaje/total de ESTE renglón son propios de
                    # este comprobante puntual -- no se copian sin más del
                    # comprobante ni de la cabecera, porque una misma
                    # retención puede tener varios comprobantes vinculados,
                    # cada uno con su propia base/porcentaje/monto retenido
                    # (pedido de Gastón, 23/09/2026). Si no se cargan a mano,
                    # se usa el neto del comprobante como base de resguardo, y
                    # el total se recalcula de neto x porcentaje si no vino.
                    neto_gravado = _parse_decimal_post(request.POST.get('neto_gravado'))
                    if neto_gravado is None:
                        neto_gravado = comprobante.neto_gravado
                    porcentaje = _parse_float_post(request.POST.get('porcentaje'))
                    total = _parse_decimal_post(request.POST.get('total'))
                    if total is None:
                        total = _calcular_total(neto_gravado, porcentaje)
                    try:
                        with transaction.atomic():
                            RetencionRenglon.objects.create(
                                retencion=retencion_obj,
                                comprobante=comprobante,
                                neto_gravado=neto_gravado,
                                porcentaje=porcentaje,
                                total=total,
                            )
                            suma, total_retencion = _chequear_suma_renglones(retencion_obj)
                    except IntegrityError:
                        messages.error(
                            request,
                            'Ese comprobante ya tiene esta misma retención (mismo impuesto y régimen) '
                            'vinculada -- no se puede repetir.'
                        )
                    except _SumaRenglonesExcedeTotal as exc:
                        messages.error(
                            request,
                            f'No se vinculó: la suma de los renglones (${exc.suma}) superaría el total de '
                            f'la retención (${exc.total}).'
                        )
                    else:
                        messages.success(
                            request,
                            f'Comprobante {comprobante.comprobante_string or comprobante.id} vinculado a la '
                            f'retención {retencion_obj.id}.'
                        )
                        if total_retencion is not None and suma < total_retencion:
                            messages.warning(
                                request,
                                f'La suma de los renglones vinculados (${suma}) todavía no llega al total '
                                f'de la retención (${total_retencion}). Podés seguir agregando renglones.'
                            )
        elif accion == 'editar_renglon':
            # Corrige el neto/porcentaje/total de un renglón YA vinculado
            # (pedido de Gastón, 23/09/2026) -- sigue sin tocar nada de la
            # Retencion/cabecera, sólo esta fila de retencion_renglon.
            renglon_id = request.POST.get('renglon_id', '')
            renglon = RetencionRenglon.objects.filter(pk=renglon_id, retencion_id__in=[l.id for l in lineas]).first()
            if not renglon:
                messages.error(request, 'No se encontró ese vínculo.')
            else:
                neto_gravado = _parse_decimal_post(request.POST.get('neto_gravado'))
                porcentaje = _parse_float_post(request.POST.get('porcentaje'))
                total = _parse_decimal_post(request.POST.get('total'))
                if total is None:
                    total = _calcular_total(neto_gravado, porcentaje)
                retencion_del_renglon = next((l for l in lineas if l.id == renglon.retencion_id), None) or renglon.retencion
                try:
                    with transaction.atomic():
                        renglon.neto_gravado = neto_gravado
                        renglon.porcentaje = porcentaje
                        renglon.total = total
                        renglon.save()
                        suma, total_retencion = _chequear_suma_renglones(retencion_del_renglon)
                except _SumaRenglonesExcedeTotal as exc:
                    messages.error(
                        request,
                        f'No se guardó: la suma de los renglones (${exc.suma}) superaría el total de la '
                        f'retención (${exc.total}).'
                    )
                else:
                    messages.success(
                        request,
                        f'Se actualizó el renglón del comprobante '
                        f'{renglon.comprobante.comprobante_string or renglon.comprobante_id}.'
                    )
                    if total_retencion is not None and suma < total_retencion:
                        messages.warning(
                            request,
                            f'La suma de los renglones vinculados (${suma}) todavía no llega al total de '
                            f'la retención (${total_retencion}). Podés seguir agregando renglones.'
                        )
        elif accion == 'quitar_renglon':
            renglon_id = request.POST.get('renglon_id', '')
            renglon = RetencionRenglon.objects.filter(pk=renglon_id, retencion_id__in=[l.id for l in lineas]).first()
            if renglon:
                renglon.delete()
                messages.success(request, 'Vínculo quitado.')
            else:
                messages.error(request, 'No se encontró ese vínculo.')

        return redirect('retenciones:modificar', id=id)

    renglones_por_retencion = {}
    for rn in renglones_nuevos:
        renglones_por_retencion.setdefault(rn.retencion_id, []).append(rn)
    # Se cuelga como atributo de cada línea (en vez de pasar el dict aparte)
    # porque el template necesita indexar por l.id, y eso no se puede hacer
    # con la sintaxis de puntos de Django (sólo permite claves fijas).
    for l in lineas:
        l.renglones_vinculados = renglones_por_retencion.get(l.id, [])

    return render(request, 'retenciones/retencion_vincular_renglones.html', {
        'anio': primera.año,
        'numero': primera.numero,
        'lineas': lineas,
        'entidad_texto': primera.entidad_nombre or (str(primera.entidad) if primera.entidad else ''),
        'entidad_id': primera.entidad_id,
        'es_emisor': primera.es_emisor if primera.es_emisor is not None else Retencion.ES_EMISOR,
        'total': sum((l.total or Decimal('0')) for l in lineas),
    })


def retencion_modificar(request, id):
    primera = (
        Retencion.objects.filter(pk=id)
        .select_related('entidad', 'id_impuesto', 'id_regimen')
        .first()
    )
    if primera is None:
        messages.error(request, f'No se encontró la retención con id {id}.')
        return redirect('retenciones:listado')

    anio = primera.año
    numero = primera.numero

    if _tiene_liquidacion(id):
        # Modo restringido (pedido por Gastón, 23/09/2026): mientras haya una
        # liquidación asignada, la cabecera y los datos viejos
        # (subtotal/porcentaje/total/comprobante_origen) quedan de sólo
        # lectura -- eso nunca se puede tocar acá. Lo único que se permite es
        # vincular (o quitar) el/los Comprobante(s) real(es) en la tabla
        # nueva RetencionRenglon: es sólo agregar documentación, no toca el
        # total ni ningún otro dato de la Retencion.
        renglones_nuevos = list(
            RetencionRenglon.objects.filter(retencion_id=id)
            .select_related('comprobante', 'comprobante__tipo_comprobante', 'retencion')
        )
        return _retencion_vincular_renglones(request, id, [primera], renglones_nuevos)

    if request.method == 'POST':
        header_form = RetencionHeaderForm(request.POST)
        formset = RetencionRenglonFormSet(request.POST, prefix='form')

        if header_form.is_valid() and formset.is_valid():
            lineas_validas = [
                f for f in formset
                if f.cleaned_data and not f.cleaned_data.get('_vacio') and not f.cleaned_data.get('DELETE')
            ]
            if not lineas_validas:
                messages.error(request, 'Cargá al menos un renglón con los datos de la operación.')
            else:
                try:
                    with transaction.atomic():
                        # Sólo se toca esta fila puntual (por id) -- nunca otras
                        # filas que puedan compartir año+numero (con otra
                        # entidad, u otro impuesto/régimen de la misma
                        # entidad). Al borrar la Retencion se van en cascada
                        # sus RetencionRenglon (on_delete=CASCADE), así que se
                        # recrean todos de cero con los datos del form.
                        Retencion.objects.filter(pk=id).delete()
                        retencion, cantidad = _guardar_grupo(header_form, formset)
                        suma, total_retencion = _chequear_suma_renglones(retencion)
                except IntegrityError:
                    messages.error(
                        request,
                        'Una de las facturas elegidas ya tiene esta misma retención (mismo impuesto y '
                        'régimen) vinculada en otra retención -- no se puede repetir. Revisá los '
                        'renglones cargados.'
                    )
                except _SumaRenglonesExcedeTotal as exc:
                    messages.error(
                        request,
                        f'No se guardó: la suma de los renglones (${exc.suma}) superaría el Total de la '
                        f'retención (${exc.total}). Corregí el Total del encabezado o los renglones.'
                    )
                else:
                    nuevo_anio = header_form.cleaned_data['año']
                    nuevo_numero = header_form.cleaned_data['numero']
                    accion = request.POST.get('accion')
                    messages.success(
                        request,
                        f'La retención {nuevo_anio}-{nuevo_numero:04d} se modificó correctamente '
                        f'({cantidad} renglón{"es" if cantidad != 1 else ""} vinculado'
                        f'{"s" if cantidad != 1 else ""}).'
                    )
                    if total_retencion is not None and suma < total_retencion:
                        messages.warning(
                            request,
                            f'La suma de los renglones (${suma}) todavía no llega al Total cargado '
                            f'(${total_retencion}).'
                        )
                    if accion == 'pdf':
                        return retencion_pdf(request, retencion.id)
                    if accion == 'excel':
                        return retencion_excel(request, retencion.id)
                    return redirect('retenciones:listado')
    else:
        header_form = RetencionHeaderForm(initial={
            'entidad': primera.entidad,
            'entidad_nombre': primera.entidad_nombre or (primera.entidad.nombre if primera.entidad else ''),
            'id_impuesto': primera.id_impuesto,
            'id_regimen': primera.id_regimen,
            'es_emisor': primera.es_emisor if primera.es_emisor is not None else Retencion.ES_EMISOR,
            'fecha': primera.fecha,
            'año': primera.año,
            'numero': primera.numero,
            'total': primera.total,
        })
        formset = RetencionRenglonFormSet(
            prefix='form',
            initial=_armar_formset_inicial(primera),
        )
        formset.extra = 1

    return render(request, 'retenciones/retencion_form.html', {
        'header_form': header_form,
        'formset': formset,
        'modo': 'modificar',
        'id': id,
        'anio': anio,
        'numero': numero,
        'entidad_texto': primera.entidad_nombre or (str(primera.entidad) if primera.entidad else ''),
        'regimen_impuesto_map': _regimen_impuesto_map(),
    })


def retencion_eliminar(request, id):
    primera = Retencion.objects.filter(pk=id).select_related('entidad').first()
    if primera is None:
        messages.error(request, f'No se encontró la retención con id {id}.')
        return redirect('retenciones:listado')

    anio = primera.año
    numero = primera.numero

    if request.method == 'POST':
        if _tiene_liquidacion(id):
            messages.error(
                request,
                f'Esta retención ({anio}-{numero:04d}) ya está incluida en una '
                'liquidación y no se puede eliminar desde acá.' if anio and numero else
                'Esta retención ya está incluida en una liquidación y no se puede eliminar desde acá.'
            )
            return redirect('retenciones:listado')

        Retencion.objects.filter(pk=id).delete()
        messages.success(request, 'La retención se eliminó correctamente.')
        return redirect('retenciones:listado')

    total = primera.total or Decimal('0')
    return render(request, 'retenciones/retencion_eliminar_confirm.html', {
        'anio': anio,
        'numero': numero,
        'lineas': [primera],
        'total': total,
        'entidad_nombre': primera.entidad_nombre or (str(primera.entidad) if primera.entidad else ''),
    })


# ---------------------------------------------------------------------------
# Listado
# ---------------------------------------------------------------------------

def retencion_listado(request):
    q_anio = request.GET.get('anio', '').strip()
    q_numero = request.GET.get('numero', '').strip()
    q_entidad = request.GET.get('entidad', '').strip()

    qs = Retencion.objects.select_related('entidad')
    if q_anio.isdigit():
        qs = qs.filter(año=int(q_anio))
    if q_numero.isdigit():
        qs = qs.filter(numero=int(q_numero))
    if q_entidad:
        qs = qs.filter(Q(entidad_nombre__icontains=q_entidad) | Q(entidad__nombre__icontains=q_entidad))

    # Cada fila de Retencion es un comprobante independiente (no se agrupan
    # por año+numero: ese agrupamiento fue la causa del bug reportado por
    # Gastón el 23/09/2026 -- distintas entidades, o distintos
    # impuesto/régimen de una misma entidad, pueden compartir numero, y
    # agruparlas mezclaba/ocultaba retenciones ajenas bajo una sola fila).
    filas = [
        {
            'id': r.id,
            'año': r.año,
            'numero': r.numero,
            'entidad_nombre': r.entidad_nombre or (str(r.entidad) if r.entidad else ''),
            'fecha': r.fecha,
            'total': r.total or Decimal('0'),
            'es_emisor': r.es_emisor if r.es_emisor is not None else Retencion.ES_EMISOR,
        }
        for r in qs.order_by('-año', '-numero', '-id')
    ]

    campos_orden = {
        'anio': lambda f: f['año'] or 0,
        'numero': lambda f: f['numero'] or 0,
        'entidad': lambda f: (f['entidad_nombre'] or '').lower(),
        'fecha': lambda f: f['fecha'],
        'total': lambda f: f['total'],
        'direccion': lambda f: f['es_emisor'],
    }
    if request.GET.get('orden') in campos_orden:
        lista = aplicar_orden_lista(request, filas, campos_orden)
    else:
        # Sin orden pedido por columna: más recientes primero (año y número
        # descendente), mismo criterio que antes de poder ordenar por columna.
        lista = sorted(filas, key=lambda f: (f['año'] or 0, f['numero'] or 0, f['id']), reverse=True)
    lista = lista[:500]

    return render(request, 'retenciones/retencion_listado.html', {
        'grupos': lista,
        'q_anio': q_anio,
        'q_numero': q_numero,
        'q_entidad': q_entidad,
    })


# ---------------------------------------------------------------------------
# Impresión (PDF / Excel), formato "Constancia de Retención"
# ---------------------------------------------------------------------------

def _lineas_impresion(retencion, tipos_por_id):
    """Filas de 'Detalle de las operaciones' a imprimir: una Retencion NUEVA
    (creada con la pantalla rediseñada del 24/09/2026, sin datos en sus
    propios campos sueltos subtotal/tipo_comp_origen) imprime una fila por
    cada RetencionRenglon vinculado, con los datos de ESE renglón (su
    comprobante real, su neto/porcentaje/total propios). Una Retencion
    VIEJA (o cualquiera que todavía tenga cargados sus propios subtotal/
    tipo_comp_origen -- comportamiento histórico, de antes de que existiera
    RetencionRenglon) sigue imprimiendo su única línea de siempre, aunque
    también tenga RetencionRenglon agregados después por 'Vincular
    renglones' (eso es sólo documentación adicional, nunca reemplaza el
    total/porcentaje ya cargado en el encabezado)."""
    if retencion.subtotal is None and retencion.tipo_comp_origen is None:
        renglones = list(
            retencion.renglones.select_related('comprobante', 'comprobante__tipo_comprobante').order_by('id')
        )
        if renglones:
            return [
                SimpleNamespace(
                    fecha_comp_origen=rn.comprobante.fecha if rn.comprobante else None,
                    texto_factura=_texto_comprobante(rn.comprobante),
                    subtotal=rn.neto_gravado,
                    porcentaje=rn.porcentaje,
                    total=rn.total,
                )
                for rn in renglones
            ]
    return [SimpleNamespace(
        fecha_comp_origen=retencion.fecha_comp_origen,
        texto_factura=_texto_comprobante_origen(retencion, tipos_por_id),
        subtotal=retencion.subtotal,
        porcentaje=retencion.porcentaje,
        total=retencion.total,
    )]


def _contexto_impresion(id):
    primera = Retencion.objects.filter(pk=id).select_related('entidad', 'id_impuesto', 'id_regimen').first()
    if primera is None or primera.año is None or primera.numero is None:
        return None
    tipos_por_id = _tipos_comprobante_por_id()
    lineas = _lineas_impresion(primera, tipos_por_id)

    # El total impreso es siempre el de la Retencion (el dato maestro,
    # cargado a mano en el encabezado -- ver RetencionHeaderForm.total) y no
    # la suma de las líneas: para una retención nueva la suma de renglones
    # nunca puede superarlo (_chequear_suma_renglones) pero puede quedar
    # por debajo si falta vincular algún comprobante, y en ese caso el
    # total impreso tiene que seguir siendo el real, no uno parcial.
    total = primera.total if primera.total is not None else sum((l.total or Decimal('0')) for l in lineas)
    es_emisor = primera.es_emisor if primera.es_emisor is not None else Retencion.ES_EMISOR

    return {
        'anio': primera.año,
        'numero': primera.numero,
        'entidad': primera.entidad,
        'entidad_nombre': primera.entidad_nombre or (str(primera.entidad) if primera.entidad else ''),
        'domicilio': _formatear_domicilio(primera.entidad),
        'cuit': _formatear_cuit(primera.entidad.cuit) if primera.entidad else '',
        'regimen': primera.id_regimen,
        'lineas': lineas,
        'total': total,
        'es_emisor': es_emisor,
        'textos': _textos_retencion(es_emisor),
    }


def retencion_pdf(request, id):
    from django.http import HttpResponse
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from movimientos.templatetags.movimientos_extras import separador_miles

    contexto = _contexto_impresion(id)
    if contexto is None:
        messages.error(request, 'No se encontró esa retención, o le falta año/número para poder imprimirla.')
        return redirect('retenciones:listado')

    anio = contexto['anio']
    numero = contexto['numero']
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename=retencion_{anio}_{numero:04d}.pdf'

    doc = SimpleDocTemplate(
        response, pagesize=A4,
        topMargin=1.2 * cm, bottomMargin=1.2 * cm, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
    )
    estilos = getSampleStyleSheet()
    titulo_estilo = ParagraphStyle('titulo_retencion', parent=estilos['Title'], fontSize=13)
    subtitulo_estilo = ParagraphStyle('subtitulo_retencion', parent=estilos['Normal'], alignment=1, fontSize=9)

    regimen_nombre = str(contexto['regimen']) if contexto['regimen'] else ''
    textos = contexto['textos']

    elementos = [
        Paragraph(textos['titulo'], titulo_estilo),
    ]
    if regimen_nombre:
        elementos.append(Paragraph(regimen_nombre, subtitulo_estilo))
    elementos.append(Paragraph(
        f"<b>COMPROBANTE Nº</b> {anio}- {numero:04d}", subtitulo_estilo
    ))
    elementos.append(Spacer(1, 0.5 * cm))

    datos_proveedor = [
        [textos['entidad_label'], contexto['entidad_nombre']],
        ['Domicilio', contexto['domicilio']],
        ['CUIT Nº', contexto['cuit']],
    ]
    tabla_proveedor = Table(datos_proveedor, colWidths=[3 * cm, 14 * cm])
    tabla_proveedor.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('BOX', (0, 0), (-1, -1), 0.75, colors.black),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
    ]))
    elementos.append(tabla_proveedor)
    elementos.append(Spacer(1, 0.5 * cm))
    elementos.append(Paragraph('<b>DETALLE DE LAS OPERACIONES</b>', estilos['Normal']))
    elementos.append(Spacer(1, 0.2 * cm))

    filas = [['Fecha', 'Factura', 'Importe', 'Porcentaje', 'Retención']]
    for l in contexto['lineas']:
        filas.append([
            l.fecha_comp_origen.strftime('%d/%m/%Y') if l.fecha_comp_origen else '-',
            l.texto_factura or '-',
            separador_miles(l.subtotal) if l.subtotal is not None else '-',
            f'{l.porcentaje:g}%' if l.porcentaje is not None else '-',
            separador_miles(l.total) if l.total is not None else '-',
        ])

    tabla_detalle = Table(filas, colWidths=[2.5 * cm, 5 * cm, 4 * cm, 3 * cm, 4.5 * cm], repeatRows=1)
    tabla_detalle.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#343a40')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ALIGN', (2, 0), (-1, -1), 'RIGHT'),
        ('ALIGN', (0, 0), (1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elementos.append(tabla_detalle)
    elementos.append(Spacer(1, 0.3 * cm))

    tabla_total = Table([['TOTAL $', separador_miles(contexto['total'])]], colWidths=[15 * cm, 4 * cm])
    tabla_total.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 11),
        ('ALIGN', (0, 0), (0, 0), 'RIGHT'),
        ('ALIGN', (1, 0), (1, 0), 'RIGHT'),
        ('BOX', (1, 0), (1, 0), 0.75, colors.black),
    ]))
    elementos.append(tabla_total)
    elementos.append(Spacer(1, 1.2 * cm))
    if textos['mostrar_firma']:
        elementos.append(Paragraph('_' * 30, estilos['Normal']))
        elementos.append(Paragraph('FIRMA', ParagraphStyle('firma', parent=estilos['Normal'], alignment=1)))
    else:
        elementos.append(Paragraph(
            'Registro interno de una retención que nos practicó la entidad -- '
            'no reemplaza la constancia que debe emitir ella.',
            ParagraphStyle('nota_sufrida', parent=estilos['Normal'], fontSize=8, textColor=colors.grey),
        ))

    doc.build(elementos)
    return response


def retencion_excel(request, id):
    import openpyxl
    from django.http import HttpResponse
    from openpyxl.styles import Font, Alignment
    from services.gestorexcel import definir_estilo_general, formatear_celda_fecha, formatear_celda_numero

    contexto = _contexto_impresion(id)
    if contexto is None:
        messages.error(request, 'No se encontró esa retención, o le falta año/número para poder imprimirla.')
        return redirect('retenciones:listado')

    anio = contexto['anio']
    numero = contexto['numero']
    textos = contexto['textos']
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = textos['hoja_excel']

    regimen_nombre = str(contexto['regimen']) if contexto['regimen'] else ''
    ws.append([textos['titulo']])
    if regimen_nombre:
        ws.append([regimen_nombre])
    ws.append([f'COMPROBANTE Nº {anio}- {numero:04d}'])
    ws.append([''])
    ws.append([textos['entidad_label'], contexto['entidad_nombre']])
    ws.append(['Domicilio', contexto['domicilio']])
    ws.append(['CUIT Nº', contexto['cuit']])
    ws.append([''])
    fila_encabezado = ws.max_row + 1
    ws.append(['Fecha', 'Factura', 'Importe', 'Porcentaje', 'Retención'])

    for l in contexto['lineas']:
        ws.append([
            l.fecha_comp_origen,
            l.texto_factura or '',
            float(l.subtotal) if l.subtotal is not None else None,
            l.porcentaje,
            float(l.total) if l.total is not None else None,
        ])

    fila_totales = ws.max_row + 1
    ws.append(['', '', '', 'Total', float(contexto['total'])])

    formatear_celda_fecha(wb, ws, 'A')
    for columna in ('C', 'E'):
        for cell in ws[columna]:
            if cell.row > fila_encabezado:
                cell.number_format = '#,##0.00'
    for cell in ws['D']:
        if cell.row > fila_encabezado:
            cell.number_format = '0.00"%"'

    definir_estilo_general(ws)

    fuente_titulo = Font(name='Arial', size=12, bold=True)
    ws['A1'].font = fuente_titulo
    for fila in range(1, fila_encabezado):
        for cell in ws[fila]:
            cell.font = Font(name='Arial', size=9, bold=(cell.column_letter == 'A'))

    fuente_total = Font(name='Arial', size=9, bold=True)
    for cell in ws[fila_totales]:
        cell.font = fuente_total

    ws.column_dimensions['B'].width = 30
    for cell in ws['B']:
        cell.alignment = Alignment(wrap_text=True, vertical='top')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=retencion_{anio}_{numero:04d}.xlsx'
    wb.save(response)
    return response


# ---------------------------------------------------------------------------
# Ret. Impuestos: alta / modificación / listado
# ---------------------------------------------------------------------------

def _siguiente_id_tipo_impuesto():
    ultimo = RetencionTipoImpuesto.objects.aggregate(Max('id'))['id__max'] or 0
    return ultimo + 1


def retencion_tipo_impuesto_listado(request):
    impuestos = RetencionTipoImpuesto.objects.all().order_by('nombre')
    impuestos = aplicar_orden_queryset(request, impuestos, {
        'id': 'id',
        'nombre': 'nombre',
    })
    return render(request, 'retenciones/retencion_tipo_impuesto_listado.html', {
        'impuestos': impuestos,
    })


def retencion_tipo_impuesto_alta(request):
    if request.method == 'POST':
        form = RetencionTipoImpuestoForm(request.POST)
        if form.is_valid():
            impuesto = form.save(commit=False)
            # Si Gastón cargó un ID a mano (para que coincida con el código
            # que ya tiene en AFIP u otro sistema), se respeta ese; si no,
            # se asigna el próximo disponible como hasta ahora.
            impuesto.id = form.cleaned_data.get('id') or _siguiente_id_tipo_impuesto()
            impuesto.save(force_insert=True)
            messages.success(request, f'El impuesto "{impuesto.nombre}" se creó correctamente.')
            return redirect('retenciones:tipo_impuesto_listado')
    else:
        form = RetencionTipoImpuestoForm()

    return render(request, 'retenciones/retencion_tipo_impuesto_form.html', {
        'form': form,
        'modo': 'alta',
    })


def retencion_tipo_impuesto_modificar(request, pk):
    impuesto = get_object_or_404(RetencionTipoImpuesto, pk=pk)

    if request.method == 'POST':
        form = RetencionTipoImpuestoForm(request.POST, instance=impuesto)
        if form.is_valid():
            form.save()
            messages.success(request, f'El impuesto "{impuesto.nombre}" se modificó correctamente.')
            return redirect('retenciones:tipo_impuesto_listado')
    else:
        form = RetencionTipoImpuestoForm(instance=impuesto)

    return render(request, 'retenciones/retencion_tipo_impuesto_form.html', {
        'form': form,
        'modo': 'modificar',
        'impuesto': impuesto,
    })


def retencion_tipo_impuesto_eliminar(request, pk):
    """Confirmación + baja de un RetencionTipoImpuesto. Se bloquea si está en
    uso (por una retención directamente, o por un régimen que lo referencia),
    ya que la FK está declarada DO_NOTHING y no lo impediría en la base."""
    impuesto = get_object_or_404(RetencionTipoImpuesto, pk=pk)

    en_uso = (
        Retencion.objects.filter(id_impuesto=impuesto).exists()
        or RetencionTipoRegimen.objects.filter(impuesto=impuesto).exists()
    )
    if en_uso:
        messages.error(
            request,
            f'El impuesto "{impuesto.nombre}" está siendo usado en una retención o en un '
            'régimen y no se puede eliminar.'
        )
        return redirect('retenciones:tipo_impuesto_listado')

    if request.method == 'POST':
        nombre = impuesto.nombre
        impuesto.delete()
        messages.success(request, f'El impuesto "{nombre}" se eliminó correctamente.')
        return redirect('retenciones:tipo_impuesto_listado')

    return render(request, 'retenciones/retencion_tipo_impuesto_eliminar_confirm.html', {
        'impuesto': impuesto,
    })


# ---------------------------------------------------------------------------
# Ret. Regimenes: alta / modificación / listado
# ---------------------------------------------------------------------------

def _siguiente_id_tipo_regimen():
    ultimo = RetencionTipoRegimen.objects.aggregate(Max('id'))['id__max'] or 0
    return ultimo + 1


def retencion_tipo_regimen_listado(request):
    regimenes = RetencionTipoRegimen.objects.prefetch_related('impuestos').order_by('nombre')
    regimenes = aplicar_orden_queryset(request, regimenes, {
        'id': 'id',
        'nombre': 'nombre',
    })
    return render(request, 'retenciones/retencion_tipo_regimen_listado.html', {
        'regimenes': regimenes,
    })


def retencion_tipo_regimen_alta(request):
    if request.method == 'POST':
        form = RetencionTipoRegimenForm(request.POST)
        if form.is_valid():
            regimen = form.save(commit=False)
            # Mismo criterio que en el alta de Impuestos: si se cargó un ID
            # a mano se respeta (para que coincida con AFIP u otro sistema),
            # si no se asigna el próximo disponible.
            regimen.id = form.cleaned_data.get('id') or _siguiente_id_tipo_regimen()
            # La columna legacy id_impuesto sigue siendo NOT NULL en la base
            # real (aunque el modelo la declare opcional) -- hay que
            # completarla igual al crear un régimen nuevo o la base rechaza
            # el INSERT. Se usa el primero de los impuestos tildados como
            # valor de resguardo; el vínculo real, que sí permite varios, se
            # guarda aparte en la tabla M2M (guardar_impuestos, más abajo).
            regimen.impuesto = form.cleaned_data['impuestos'][0]
            regimen.save(force_insert=True)
            form.guardar_impuestos(regimen)
            messages.success(request, f'El régimen "{regimen.nombre}" se creó correctamente.')
            return redirect('retenciones:tipo_regimen_listado')
    else:
        form = RetencionTipoRegimenForm()

    return render(request, 'retenciones/retencion_tipo_regimen_form.html', {
        'form': form,
        'modo': 'alta',
    })


def retencion_tipo_regimen_modificar(request, pk):
    regimen = get_object_or_404(RetencionTipoRegimen, pk=pk)

    if request.method == 'POST':
        form = RetencionTipoRegimenForm(request.POST, instance=regimen)
        if form.is_valid():
            form.save()
            form.guardar_impuestos(regimen)
            messages.success(request, f'El régimen "{regimen.nombre}" se modificó correctamente.')
            return redirect('retenciones:tipo_regimen_listado')
    else:
        form = RetencionTipoRegimenForm(instance=regimen)

    return render(request, 'retenciones/retencion_tipo_regimen_form.html', {
        'form': form,
        'modo': 'modificar',
        'regimen': regimen,
    })


def retencion_tipo_regimen_eliminar(request, pk):
    """Confirmación + baja de un RetencionTipoRegimen. Se bloquea si está en
    uso por alguna retención, ya que la FK está declarada DO_NOTHING y no lo
    impediría en la base."""
    regimen = get_object_or_404(RetencionTipoRegimen, pk=pk)

    if Retencion.objects.filter(id_regimen=regimen).exists():
        messages.error(
            request,
            f'El régimen "{regimen.nombre}" está siendo usado en una retención y no se puede eliminar.'
        )
        return redirect('retenciones:tipo_regimen_listado')

    if request.method == 'POST':
        nombre = regimen.nombre
        regimen.delete()
        messages.success(request, f'El régimen "{nombre}" se eliminó correctamente.')
        return redirect('retenciones:tipo_regimen_listado')

    return render(request, 'retenciones/retencion_tipo_regimen_eliminar_confirm.html', {
        'regimen': regimen,
    })


# ---------------------------------------------------------------------------
# Ranking de entidades (por monto total de retenciones, filtrando por un
# lapso de fecha) — mismo patrón que
# comprobantes.views.comprobante_ranking_entidades /
# movimientos_caja.views.movimiento_caja_ranking_entidades.
# ---------------------------------------------------------------------------

# Entidad que representa a la propia empresa (Fontana); mismo id que usan
# comprobantes.views.ENTIDAD_PROPIA_ID, movimientos_caja.views.ENTIDAD_PROPIA_ID
# y liquidaciones.views.ENTIDAD_PROPIA_ID.
ENTIDAD_PROPIA_ID = 100


def _retenciones_ranking_filtrados(request):
    """Aplica a Retencion los filtros de RankingEntidadesForm (fecha y,
    opcionalmente, excluir a Fontana). Devuelve (form, queryset,
    filtros_activos), centralizado para que la pantalla y las
    exportaciones (Excel / PDF) usen siempre los mismos criterios."""
    form = RankingEntidadesForm(request.GET or None)
    retenciones = Retencion.objects.all()

    filtros_activos = False
    if form.is_valid():
        fecha_desde = form.cleaned_data.get('fecha_desde')
        fecha_hasta = form.cleaned_data.get('fecha_hasta')
        excluir_fontana = form.cleaned_data.get('excluir_fontana')
        if fecha_desde:
            retenciones = retenciones.filter(fecha__gte=fecha_desde)
        if fecha_hasta:
            retenciones = retenciones.filter(fecha__lte=fecha_hasta)
        if excluir_fontana:
            retenciones = retenciones.exclude(entidad_id=ENTIDAD_PROPIA_ID)
        filtros_activos = bool(fecha_desde or fecha_hasta or excluir_fontana)

    return form, retenciones, filtros_activos


def _calcular_ranking_retenciones(retenciones):
    """A partir de un queryset de Retencion, arma el ranking de entidades
    por monto total (de mayor a menor) y el total general. Devuelve
    (ranking, total_general)."""
    ranking = list(
        retenciones.values('entidad_id', 'entidad__nombre')
        .annotate(total_monto=Sum('total'), cantidad=Count('id'))
        .order_by('-total_monto')
    )

    total_general = sum(
        (fila['total_monto'] for fila in ranking if fila['total_monto'] is not None), Decimal('0')
    )
    for posicion, fila in enumerate(ranking, start=1):
        fila['posicion'] = posicion
        fila['porcentaje'] = (
            fila['total_monto'] / total_general * 100
            if total_general and fila['total_monto'] is not None else Decimal('0')
        )

    return ranking, total_general


@requiere_grupo('Rankings')
def retencion_ranking_entidades(request):
    """Ranking de entidades por monto total de retenciones, de mayor a
    menor, filtrando opcionalmente por un rango de fecha (y excluyendo, si
    se pide, a Fontana)."""
    form, retenciones, filtros_activos = _retenciones_ranking_filtrados(request)
    ranking, total_general = _calcular_ranking_retenciones(retenciones)
    ranking = aplicar_orden_lista(request, ranking, {
        'posicion': lambda f: f['posicion'],
        'entidad': lambda f: (f['entidad__nombre'] or '').lower(),
        'cantidad': lambda f: f['cantidad'],
        'monto': lambda f: f['total_monto'] if f['total_monto'] is not None else Decimal('0'),
        'porcentaje': lambda f: f['porcentaje'],
    })

    return render(request, 'retenciones/retencion_ranking_entidades.html', {
        'form': form,
        'ranking': ranking,
        'total_general': total_general,
        'filtros_activos': filtros_activos,
    })


def _filas_ranking_retenciones(ranking):
    columnas = ['#', 'Entidad', 'Retenciones', 'Monto total', 'Participación %']
    filas = [
        [
            fila['posicion'],
            fila['entidad__nombre'] or 'Sin nombre',
            fila['cantidad'],
            float(fila['total_monto']) if fila['total_monto'] is not None else None,
            float(fila['porcentaje']) if fila['porcentaje'] is not None else None,
        ]
        for fila in ranking
    ]
    return {
        'columnas': columnas,
        'filas': filas,
        'columnas_numericas': {3, 4},  # Monto total, Participación %
        'anchos': [0.4, 2.2, 1.0, 1.2, 1.2],
    }


@requiere_grupo('Rankings')
def retencion_ranking_entidades_excel(request):
    _form, retenciones, _filtros_activos = _retenciones_ranking_filtrados(request)
    ranking, _total_general = _calcular_ranking_retenciones(retenciones)
    resultado = _filas_ranking_retenciones(ranking)
    return excel_response('ranking_entidades_retenciones', resultado)


@requiere_grupo('Rankings')
def retencion_ranking_entidades_pdf(request):
    _form, retenciones, _filtros_activos = _retenciones_ranking_filtrados(request)
    ranking, _total_general = _calcular_ranking_retenciones(retenciones)
    resultado = _filas_ranking_retenciones(ranking)
    return pdf_response('ranking_entidades_retenciones', 'Ranking de entidades por monto de retenciones', resultado)
