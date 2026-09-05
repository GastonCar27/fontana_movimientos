from django.shortcuts import render, get_object_or_404

from entidades.models import Entidad, Inym_Operador
from .forms import IngresoHvYerbaMateForm,MovimientoHvYerbaMate
from .forms import MovimientoPesajeForm
from .forms import MovimientoForm,MovimientoHvYerbaMateForm,SalidaForm,MovimientoReporteForm,ReporteHvYerbaMateForm,RankingProductoresForm
from .forms import PesajeInlineForm, MovimientoInymOperadoresForm
from .models import MovimientoHvYerbaMate, ProductoDetalle, ComprobanteUnidadDeMedida, MovimientoPesaje
from django.http import HttpResponseRedirect
from django.urls import reverse_lazy
from django.shortcuts import redirect, get_object_or_404
from django.contrib import messages
from django.views.generic import ListView
from django.views.generic import DetailView, UpdateView
from .models import Movimiento
from comprobantes.models import ComprobanteRenglon
from django.db import IntegrityError, transaction
from django.db.models import Q, Sum, Count, DecimalField, Value, Exists, OuterRef
from django.db.models.functions import Coalesce
from django.http import Http404
from decimal import Decimal
from datetime import date
import json

from services.buscadores import texto_entidad_buscador
from services.ordenamiento import aplicar_orden_queryset, aplicar_orden_lista
from services.reportes import excel_response, pdf_response
from services.permisos import requiere_grupo
def recepcion(request):
    return render(request,'movimientos/movimiento.html',{

    })

def ___recepcion_hv_yerba_mate(request):
    form = IngresoHvYerbaMateForm()
    return render(request,'movimientos/recepcion_hv_yerba_mate.html',{'form':form
    })

def recepcion_hv_yerba_mate(request,**kwargs):
    form_movimiento_hv = IngresoHvYerbaMateForm(request.POST or None)
    form_pesaje = MovimientoPesajeForm(request.POST or None)
    if request.method == 'POST' and form_movimiento_hv.is_valid():
        # El campo del modelo permite blank/null (por eso el ModelForm no lo exige
        # solo), pero acá la entidad emisora es obligatoria: sin este chequeo,
        # guardar sin seleccionarla rompía con AttributeError ('NoneType' object
        # has no attribute 'entidad') al armar el movimiento más abajo.
        if not form_movimiento_hv.cleaned_data.get('inym_operador_origen'):
            form_movimiento_hv.add_error(None, 'Seleccioná el operador INYM emisor (entidad) antes de guardar.')
        else:
            # 1. Crea la instancia pero no la guarda en la BD
            form_pesaje.bruto=form_movimiento_hv.cleaned_data['bruto']
            form_pesaje.tara=form_movimiento_hv.cleaned_data['tara']
            form_pesaje.descuento=form_movimiento_hv.cleaned_data['descuento']
            movimiento_hv = form_movimiento_hv.save(commit=False)
            # 2. Modifica el campo necesario
            #movimiento.entidad_receptor=Entidad.objects.get(id=100)
            #movimiento.unidad_de_medida=ComprobanteUnidadDeMedida.objects.get(id='01')
            #movimiento.entidad_emisor = movimiento_hv.inym_operador.entidad

            # 3. Guarda finalmente en la BD
            movimiento_hv.entidad_emisor = movimiento_hv.inym_operador_origen.entidad
            movimiento_hv.entidad_receptor = movimiento_hv.inym_operador_destino.entidad
            try:
                    # Intenta guardar el formulario
                    movimiento_hv.save()
                    # Aquí puedes agregar un redirect o mensaje de éxito
                    pesaje = form_pesaje.save(commit=False)

                    if form_pesaje.is_bound: # si contiene datos guardo
                        pesaje.movimiento = movimiento_hv
                        pesaje.save()
                    messages.success(request, f'La recepción número {movimiento_hv.numero} se ha registrado correctamente')
                    return redirect('movimientos:recepcion_hv_yerba_mate')
            except IntegrityError:
                    # Añade un error global al formulario para que Django lo muestre en el HTML
                    form_movimiento_hv.add_error(None, "Ya existe un registro con estos datos o viola una restricción de integridad.")

    return render(request, 'movimientos/recepcion_hv_yerba_mate.html', {'form_pesaje': form_pesaje,'form_movimiento_hv':form_movimiento_hv})

def salida_yerba_mate_canchada(request,**kwargs):
    from .forms import SalidaYerbaMateCanchadaForm
    form_salida_canchada = SalidaYerbaMateCanchadaForm(request.POST or None)
    form_pesaje = MovimientoPesajeForm(request.POST or None)
    if request.method == 'POST' and form_salida_canchada.is_valid():
        # Mismo caso que en recepcion_hv_yerba_mate: acá el campo editable es el
        # receptor (inym_operador_destino), y sin este chequeo se rompía igual
        # con AttributeError si se guardaba sin seleccionarlo.
        if not form_salida_canchada.cleaned_data.get('inym_operador_destino'):
            form_salida_canchada.add_error(None, 'Seleccioná el operador INYM receptor (entidad) antes de guardar.')
        else:
            # 1. Crea la instancia pero no la guarda en la BD
            form_pesaje.bruto=form_salida_canchada.cleaned_data['bruto']
            form_pesaje.tara=form_salida_canchada.cleaned_data['tara']
            form_pesaje.descuento=form_salida_canchada.cleaned_data['descuento']
            movimiento_hv = form_salida_canchada.save(commit=False)
            # 2. Modifica el campo necesario
            #movimiento.entidad_receptor=Entidad.objects.get(id=100)
            #movimiento.unidad_de_medida=ComprobanteUnidadDeMedida.objects.get(id='01')
            #movimiento.entidad_emisor = movimiento_hv.inym_operador.entidad

            # 3. Guarda finalmente en la BD
            movimiento_hv.entidad_emisor = movimiento_hv.inym_operador_origen.entidad
            movimiento_hv.entidad_receptor = movimiento_hv.inym_operador_destino.entidad
            try:
                    # Intenta guardar el formulario
                    movimiento_hv.save()
                    # Aquí puedes agregar un redirect o mensaje de éxito
                    pesaje = form_pesaje.save(commit=False)

                    if form_pesaje.is_bound: # si contiene datos guardo
                        pesaje.movimiento = movimiento_hv
                        pesaje.save()
                    messages.success(request, f'La Salida de canchada {movimiento_hv.numero} se ha registrado correctamente')
                    return redirect('movimientos:recepcion_hv_yerba_mate')
            except IntegrityError:
                    # Añade un error global al formulario para que Django lo muestre en el HTML
                    form_salida_canchada.add_error(None, "Ya existe un registro con estos datos o viola una restricción de integridad.")

    return render(request, 'movimientos/salida_canchada.html', {'form_pesaje': form_pesaje,'form_salida_canchada':form_salida_canchada})

class RecepcionHvYerbaMateListView(ListView):
    """Listado de recepciones de HV de Yerba Mate; es la puerta de entrada de
    'Modificación' del submenú Ingreso H.V. de Yerba Mate."""
    template_name = 'movimientos/listado_recepcion_hv_yerba_mate.html'

    def get_queryset(self):
        from .models import Movimiento
        movimientos = (
            Movimiento.objects.select_related('entidad_emisor', 'unidad_de_medida')
            .filter(producto_id=2)
            .order_by('-id_movimiento')
        )

        self.q_id = self.request.GET.get('id', '').strip()
        self.q_numero = self.request.GET.get('numero', '').strip()
        self.q_productor = self.request.GET.get('productor', '').strip()

        if self.q_id:
            movimientos = movimientos.filter(id_movimiento=self.q_id) if self.q_id.isdigit() else movimientos.none()
        if self.q_numero:
            movimientos = movimientos.filter(numero=self.q_numero) if self.q_numero.isdigit() else movimientos.none()
        if self.q_productor:
            movimientos = movimientos.filter(
                Q(entidad_emisor__nombre__icontains=self.q_productor) | Q(entidad_emisor__cuit__icontains=self.q_productor)
            )

        movimientos = aplicar_orden_queryset(self.request, movimientos, {
            'id': 'id_movimiento',
            'fecha': 'fecha',
            'numero': 'numero',
            'productor': 'entidad_emisor__nombre',
            'total': 'total',
        })

        return movimientos[:200]

    def get_context_data(self, **kwargs): ## pasa el contexto al template
        context = super().get_context_data(**kwargs)
        context['message'] = 'Listado de recepciones de HV de Yerba Mate'
        context['recepciones_hv_yerba_mate'] = context['movimiento_list']
        context['q_id'] = self.q_id
        context['q_numero'] = self.q_numero
        context['q_productor'] = self.q_productor

        return context

class RecepcionHvYerbaMateBuscarListView(ListView):
    template_name = 'movimientos/buscar_movimiento.html'

    def get_queryset(self):
        from django.db.models import Q
        filters = Q(id_movimiento=self.query())
        return Movimiento.objects.all().filter(filters)
    def query(self):

        return self.request.GET.get('q')

    def get_context_data(self, **kwargs): ## pasa el contexto al template
        context = super().get_context_data(**kwargs)
        context['message'] = 'Listado de recepciones de HV de Yerba Mate'
        context['query'] = self.query()

        return context

class MovimientoDetailView(DetailView):
    model = Movimiento
    template_name = 'movimientos/movimiento.html'

"""def movimiento_update(request,pk):
    # 1. Recuperamos la instancia exacta que queremos actualizar
    objeto_pesaje = get_object_or_404(MovimientoPesaje, pk=pk)
    objeto_yerba_mate = get_object_or_404(MovimientoHvYerbaMate, pk=pk)
    error_al_modificar=False ##si tengo error al modificar me cambia el formato de la fecha, ver forma más elegante de arreglar
    if request.method == "POST":
        if request.POST.get("modificar"):
            form = IngresoHvYerbaMateForm(request.POST,
                                          instance=objeto_yerba_mate)  # request.POST pq si no el isbound da False
            form_pesaje = MovimientoPesajeForm(request.POST, instance=objeto_pesaje)
            # 2. En POST, poblamos el formulario con los datos enviados y la instancia existente
            valido=True
            if form_pesaje.is_bound and form_pesaje.is_valid() and \
                    form.is_bound and form.is_valid():
                form.save()
                form_pesaje.save()
                messages.success(request, f'El movimiento con ID: {form.instance.pk} se ha registrado correctamente')
                return redirect('movimientos:buscar_movimiento_2')
            else:
                error_al_modificar=True
            #return redirect('lista_de_modelos')
        else:
            form = IngresoHvYerbaMateForm(instance=objeto_yerba_mate)  # request.POST pq si no el isbound da False
            form_pesaje = MovimientoPesajeForm(instance=objeto_pesaje)
        form.bruto = objeto_pesaje.bruto
        form.bruto = objeto_pesaje.tara
        form.descuento = objeto_pesaje.descuento
    else:
        # 3. En GET, simplemente pasamos la instancia para que el formulario se rellene
        form = IngresoHvYerbaMateForm(instance=objeto_yerba_mate)
        form_pesaje = MovimientoPesajeForm(instance=objeto_pesaje)

    contexto = {'form_movimiento_hv': form, 'form_pesaje':form_pesaje,'modificar':True,'error_al_modificar':error_al_modificar}
    return render(request, 'movimientos/recepcion_hv_yerba_mate.html', contexto)
"""

class MovimientoSearchListView(ListView):
    template_name = ""
    def get_queryset(self):
        return Movimiento.objects.filter(id=self.query())

    def query(self):
        return self.request.get('id_movimiento')

def mi_vista_buscar(request):
    from .forms import BuscarMovimientoForm
    if request.method == 'POST':
        # Vinculamos los datos del usuario al formulario
        form = BuscarMovimientoForm(request.POST)

        if form.is_valid():
            # Procesar los datos limpios (form.cleaned_data)
            # O guardar directamente si es un ModelForm
            id_movimiento = request.POST.get('id_movimiento')
            numero = request.POST.get('numero')
            contexto = None
            
            if id_movimiento:
                contexto={'movimientos': Movimiento.objects.filter(id_movimiento=id_movimiento)}
            elif numero:
                    contexto = {'movimientos': Movimiento.objects.filter(numero=numero)}
            else:
                producto = request.POST.get('producto')
                movimientos = Movimiento.objects.all()
                if producto:
                    movimientos = movimientos.filter(producto=producto)
                entidad_emisor = request.POST.get('emisor')
                if entidad_emisor:
                    movimientos = movimientos.filter(entidad_emisor__id=entidad_emisor)
                receptor = request.POST.get('receptor')
                if receptor:
                    movimientos = movimientos.filter(entidad_receptor=receptor)
                fecha_desde = request.POST.get('fecha_desde')
                if fecha_desde:
                    movimientos = movimientos.filter(fecha__gte=fecha_desde)#mayor o igual
                fecha_hasta = request.POST.get('fecha_hasta')
                if fecha_hasta:
                    movimientos = movimientos.filter(fecha__lte=fecha_hasta)  # menor o igual
                #agrupar por emisor_inym solo tiene sentido si es por otro valor que no sea ID o Numero
                agrupar_por_emisor_inym=False
                if request.POST.get('agrupar')=='por_emisor_inym' and movimientos:
                    agrupar_por_emisor_inym=True
                    from django.db.models import Sum, F
                    #le pongo alias a los campos
                    #movimientos_agrupados = movimientos.values(entidad_emisor=F'movimientohvyerbamate__inym_operador_origen',producto=F'producto').annotate(total=Sum('total'))
                    movimientos_agrupados = (
                    movimientos
                    .values('movimientohvyerbamate__inym_operador_origen', \
                    'movimientohvyerbamate__inym_operador_origen__entidad__nombre', 'producto__nombre')
                    .annotate(
                        total=Sum('total'),
                        id_inym_emisor = F('movimientohvyerbamate__inym_operador_origen'),
                        entidad_emisor=F('movimientohvyerbamate__inym_operador_origen__entidad__nombre'),
                        producto = F('producto__nombre')
                        )
                    .order_by('movimientohvyerbamate__inym_operador_origen__entidad__nombre')
                        
                    )
                if movimientos and not agrupar_por_emisor_inym:
                    contexto = {'movimientos': movimientos,'form':form}
                elif agrupar_por_emisor_inym:
                    contexto_agrupados = {'movimientos': movimientos_agrupados,'form':form}
            if contexto:
                return render(
                    request,
                    'movimientos/listado_movimiento.html',
                    context=contexto)
            elif agrupar_por_emisor_inym:
                return render(
                        request,
                        'movimientos/listado_movimiento_agrupado.html',
                        context=contexto_agrupados)
        


            #redirect('movimientos:listado_movimiento')
    else:
        print("CONSOLA ENTRA")
        # Formulario vacío inicial
        form = BuscarMovimientoForm()

    producto_id = form['producto'].value()
    producto_texto = str(ProductoDetalle.objects.filter(pk=producto_id).first() or '') if producto_id else ''

    return render(request, 'movimientos/mi_plantilla.html', {'form': form, 'producto_texto': producto_texto})

def vista_buscar_saldo_producto(request):
    from .forms import BuscarSaldoProductoMovimientoComprobanteForm
    if request.method == 'POST':
        # Vinculamos los datos del usuario al formulario
        form = BuscarSaldoProductoMovimientoComprobanteForm(request.POST)

        if form.is_valid():
            # Procesar los datos limpios (form.cleaned_data)
            # O guardar directamente si es un ModelForm 
            contexto = None
            producto = request.POST.get('producto')
            movimientos = Movimiento.objects.all()
            if producto:
                movimientos = movimientos.filter(producto=producto)
            fecha_desde = request.POST.get('fecha_desde')
            if fecha_desde:
                movimientos = movimientos.filter(fecha__gte=fecha_desde)#mayor o igual
            fecha_hasta = request.POST.get('fecha_hasta')
            if fecha_hasta:
                movimientos = movimientos.filter(fecha__lte=fecha_hasta)  # menor o igual
            #agrupar por emisor_inym solo tiene sentido si es por otro valor que no sea ID o Numero
        
            if movimientos:
                contexto = {'movimientos': movimientos,'form':form}
            if contexto:
                tipo_accion = request.POST.get('accion')
                print("despues",print(request.POST.get('accion')))
                if tipo_accion=='Generar Excel':
                    return generar_excel_producto_saldo(request)
                    
            
            #redirect('movimientos:listado_movimiento')
    else:
        print("CONSOLA ENTRA")
        # Formulario vacío inicial
        form = BuscarSaldoProductoMovimientoComprobanteForm()

    producto_id = form['producto'].value()
    producto_texto = str(ProductoDetalle.objects.filter(pk=producto_id).first() or '') if producto_id else ''

    return render(request, 'movimientos/producto_buscar_saldo.html', {'form': form, 'producto_texto': producto_texto})

     
def crear_response_excel():
    from django.http import HttpResponse
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = 'attachment; filename=reporte.xlsx'

    return response

class MovimientoSearchListView(ListView):
    Model= Movimiento
    template_name = 'listado_movimiento.html'
    context_object_name = 'movimientos'
    queryset = Movimiento.objects.all()
    paginate_by = 10

def remove_movimiento(request):
     #si el objeto no existe no se lanza una excepcion, si no error 404
    movimiento = get_object_or_404(Movimiento,pk=request.POST.get('movimiento_id'))
    movimiento.delete()

    return redirect('movimientos:buscar_movimiento_2')


class MovimientoHvYerbaMateUpdateView(UpdateView):
    from django.urls import reverse_lazy
    model = MovimientoHvYerbaMate
    form_class = MovimientoHvYerbaMateForm
    template_name = 'movimientos/movimiento_update_form.html'
    #fields = ['fecha', 'numero', 'inym_operador_origen','inym_operador_destino','total','unidad_de_medida','producto']
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        print("¿EL OBJETO EXISTE EN LA VISTA?:", self.object) 
        print("DATOS DEL FORMULARIO EN CONTEXTO:", context['form'].initial)
        return context
    def get_success_url(self):
        # Vuelve al listado de recepciones (antes iba a la ficha del movimiento).
        return reverse_lazy('movimientos:listado_recepcion_hv_yerba_mate')
    def form_valid(self, form): ##para guardar también en pesaje
        # Guardamos primero el objeto principal para asegurar que self.object esté actualizado
        response = super().form_valid(form)

        bruto = form.cleaned_data.get('bruto')
        tara = form.cleaned_data.get('tara')
        descuento = form.cleaned_data.get('descuento')

        # Evaluamos si realmente vienen datos vacíos/ceros
        hay_pesaje = any([bruto, tara, descuento])

        if not hay_pesaje:
        # Si no hay datos, intentamos borrar el pesaje si existía
            MovimientoPesaje.objects.filter(pk=self.object.pk).delete()
        else:
            # Si hay datos, creamos o actualizamos
            MovimientoPesaje.objects.update_or_create(
                pk=self.object.pk,
                defaults={'bruto': bruto, 'tara': tara, 'descuento': descuento}
            )

        messages.success(
            self.request,
            f"El movimiento número {self.object.numero} con id {self.object.pk} se modificó correctamente"
        )

        return response


def movimiento_hv_yerba_mate_eliminar(request, pk):
    """Confirmación + baja de un movimiento de HV de Yerba Mate, desde el
    listado de 'Modificación' del submenú Ingreso H.V. de Yerba Mate.
    Elimina el Movimiento completo (la fila padre): por el on_delete=CASCADE
    ya declarado en los modelos, arrastra también la fila hija
    (movimiento_hv_yerba_mate) y el pesaje asociado (movimiento_pesaje),
    igual que ya hace remove_movimiento con cualquier otro movimiento."""
    movimiento_hv = get_object_or_404(MovimientoHvYerbaMate, pk=pk)

    if request.method == 'POST':
        numero = movimiento_hv.numero
        movimiento_hv.delete()
        messages.success(
            request,
            f"El movimiento número {numero} con id {pk} se eliminó correctamente"
        )
        return redirect('movimientos:listado_recepcion_hv_yerba_mate')

    return render(
        request,
        'movimientos/movimiento_hv_yerba_mate_eliminar_confirm.html',
        {'movimiento': movimiento_hv},
    )


def _filtrar_reporte_hv_yerba_mate(request):
    """Arma el form y el queryset del Reporte de Ingreso H.V. de Yerba Mate,
    compartido entre la vista de pantalla y las de exportación (Excel/PDF) para
    que ambas apliquen exactamente los mismos filtros.

    Nota: MovimientoPesaje declara su FK a Movimiento con parent_link=True sin
    heredar realmente de Movimiento, y esto rompe la resolución por nombre del
    related_name 'movimiento_pesaje' cuando se consulta desde el modelo hijo
    MovimientoHvYerbaMate (Django tira 'Cannot resolve keyword movimiento_pesaje
    into field'). Por eso el filtro "tiene pesaje cargado" se arma con un
    Exists() apuntando directo a MovimientoPesaje, en vez de
    .filter(movimiento_pesaje__isnull=False)."""
    form = ReporteHvYerbaMateForm(request.GET or None)

    tiene_pesaje = MovimientoPesaje.objects.filter(movimiento_id=OuterRef('pk'))
    movimientos = (
        MovimientoHvYerbaMate.objects.select_related(
            'producto', 'entidad_emisor', 'entidad_receptor', 'unidad_de_medida',
            'inym_operador_origen__entidad', 'inym_operador_destino__entidad',
        )
        .annotate(_tiene_pesaje=Exists(tiene_pesaje))
        .filter(_tiene_pesaje=True)
        .order_by('fecha', 'numero')
    )

    if form.is_valid():
        fecha_desde = form.cleaned_data.get('fecha_desde')
        fecha_hasta = form.cleaned_data.get('fecha_hasta')
        entidad_busqueda = form.cleaned_data.get('entidad_emisor')
        if fecha_desde:
            movimientos = movimientos.filter(fecha__gte=fecha_desde)
        if fecha_hasta:
            movimientos = movimientos.filter(fecha__lte=fecha_hasta)
        if entidad_busqueda:
            entidad_ids = _buscar_entidades_con_inym_operador(entidad_busqueda)
            movimientos = movimientos.filter(entidad_emisor_id__in=entidad_ids)

    return form, movimientos


def _buscar_entidades_con_inym_operador(texto_busqueda):
    """Devuelve los IDs de las Entidades que tienen al menos un Inym Operador
    asociado y que además matchean la búsqueda por ID de entidad, nombre
    (parcial) o ID de Inym Operador."""
    texto = texto_busqueda.strip()
    entidades = Entidad.objects.filter(inym_operador__isnull=False)

    filtro = Q(nombre__icontains=texto)
    if texto.isdigit():
        valor = int(texto)
        filtro |= Q(id=valor) | Q(inym_operador__id=valor)

    return entidades.filter(filtro).values_list('id', flat=True).distinct()


def _lineas_entidad_operador(entidad_id, entidad, operador):
    """Arma el texto combinado de una Entidad con su Inym Operador asociado
    (se usa tanto para el lado emisor como el receptor de un movimiento).
    Devuelve una tupla (linea_entidad, linea_operador); linea_operador queda
    '' si no hay operador cargado."""
    if not entidad_id:
        return ('-', '')
    linea_entidad = f"ID: {entidad_id} Nombre Entidad: {entidad.nombre if entidad and entidad.nombre else ''}"
    linea_operador = ''
    if operador:
        nombre_operador = operador.entidad.nombre if operador.entidad_id and operador.entidad.nombre else ''
        linea_operador = f"Inym Operador ID: {operador.id} ({operador.tipo_operador}) {nombre_operador}"
    return (linea_entidad, linea_operador)


def _enriquecer_movimientos(movimientos):
    """Recibe una lista (ya evaluada/recortada) de MovimientoHvYerbaMate y le
    engancha en memoria, por cada uno: el MovimientoPesaje asociado (atributo
    .pesaje) y el texto combinado de Entidad + Inym Operador para emisor y
    receptor (.emisor_entidad_txt/.emisor_operador_txt y
    .receptor_entidad_txt/.receptor_operador_txt), para no repetir esa lógica
    en cada vista (pantalla, Excel, PDF).

    Nota: el pesaje se trae en una consulta aparte porque no se puede
    resolver 'movimiento_pesaje' por nombre desde este modelo (ver
    _filtrar_reporte_hv_yerba_mate)."""
    movimientos = list(movimientos)
    pesajes = MovimientoPesaje.objects.filter(
        movimiento_id__in=[m.id_movimiento for m in movimientos]
    )
    pesaje_por_id = {p.movimiento_id: p for p in pesajes}
    for m in movimientos:
        m.pesaje = pesaje_por_id.get(m.id_movimiento)
        m.emisor_entidad_txt, m.emisor_operador_txt = _lineas_entidad_operador(
            m.entidad_emisor_id, m.entidad_emisor, m.inym_operador_origen
        )
        m.receptor_entidad_txt, m.receptor_operador_txt = _lineas_entidad_operador(
            m.entidad_receptor_id, m.entidad_receptor, m.inym_operador_destino
        )
    return movimientos


def reporte_hv_yerba_mate(request):
    """Reporte de Ingreso H.V. de Yerba Mate: reemplaza al viejo enlace de
    'Reportes' del submenú (que apuntaba a una plantilla inexistente y no
    llevaba a ningún lado). Filtra por fecha desde/hasta y entidad emisora, y
    permite exportar la misma consulta a Excel o PDF."""
    form, movimientos = _filtrar_reporte_hv_yerba_mate(request)

    totales = movimientos.aggregate(
        total_cantidad=Coalesce(Sum('total'), Value(Decimal('0')), output_field=DecimalField(max_digits=12, decimal_places=2)),
    )

    movimientos = aplicar_orden_queryset(request, movimientos, {
        'id': 'id_movimiento',
        'fecha': 'fecha',
        'numero': 'numero',
        'producto': 'producto__nombre',
        'unidad': 'unidad_de_medida__nombre',
        'total': 'total',
    }, default=('fecha', 'numero'))

    return render(request, 'movimientos/reporte_hv_yerba_mate.html', {
        'form': form,
        'movimientos': _enriquecer_movimientos(movimientos[:500]),
        'totales': totales,
    })


def _filtros_aplicados_reporte(form, movimientos_sin_recortar):
    """Arma la lista de (etiqueta, valor) con los filtros aplicados al
    Reporte de Ingreso H.V. de Yerba Mate, para mostrarlos en la cabecera de
    las exportaciones a Excel y PDF en vez de repetir en cada fila datos que
    en este reporte siempre son los mismos (Producto y Unidad de Medida)."""
    form.is_valid()
    filtros = []

    fecha_desde = form.cleaned_data.get('fecha_desde')
    fecha_hasta = form.cleaned_data.get('fecha_hasta')
    entidad_busqueda = form.cleaned_data.get('entidad_emisor')

    if fecha_desde:
        filtros.append(('Fecha desde', fecha_desde.strftime('%d/%m/%Y')))
    if fecha_hasta:
        filtros.append(('Fecha hasta', fecha_hasta.strftime('%d/%m/%Y')))

    producto_ids = list(movimientos_sin_recortar.values_list('producto_id', flat=True).distinct())
    productos = ProductoDetalle.objects.filter(id__in=producto_ids)
    nombres_producto = sorted({p.nombre or str(p) for p in productos})
    filtros.append(('Producto', ', '.join(nombres_producto) if nombres_producto else '-'))

    unidad_ids = [u for u in movimientos_sin_recortar.values_list('unidad_de_medida_id', flat=True).distinct() if u]
    unidades = ComprobanteUnidadDeMedida.objects.filter(id__in=unidad_ids)
    nombres_unidad = sorted({u.nombre or str(u) for u in unidades})
    filtros.append(('Unidad de Medida', ', '.join(nombres_unidad) if nombres_unidad else '-'))

    if entidad_busqueda:
        filtros.append(('Entidad Emisora', entidad_busqueda))

    return filtros


def reporte_hv_yerba_mate_excel(request):
    import openpyxl
    from django.http import HttpResponse
    from openpyxl.styles import Alignment, Font
    from services.gestorexcel import definir_estilo_general, formatear_celda_fecha, formatear_celda_numero

    form, movimientos = _filtrar_reporte_hv_yerba_mate(request)
    filtros = _filtros_aplicados_reporte(form, movimientos)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Ingreso HV Yerba Mate'

    # Cabecera con los filtros usados (reemplaza a la columna Producto, que
    # antes se repetía igual en todas las filas).
    for etiqueta, valor in filtros:
        ws.append([f'{etiqueta}: {valor}'])
    ws.append([''])  # fila en blanco de separación (openpyxl no actualiza
    # max_row con un append([]) totalmente vacío, por eso la celda vacía)
    fila_encabezado = ws.max_row + 1

    columnas = [
        'ID', 'Fecha', 'Número', 'Entidad Emisora / Inym Operador',
        'Entidad Receptora / Inym Operador', 'Bruto', 'Tara', 'Descuento', 'Total',
    ]
    ws.append(columnas)

    for m in _enriquecer_movimientos(movimientos[:2000]):
        pesaje = m.pesaje
        emisor = m.emisor_entidad_txt + (f"\n{m.emisor_operador_txt}" if m.emisor_operador_txt else '')
        receptor = m.receptor_entidad_txt + (f"\n{m.receptor_operador_txt}" if m.receptor_operador_txt else '')
        ws.append([
            m.id_movimiento,
            m.fecha,
            m.numero,
            emisor,
            receptor,
            pesaje.bruto if pesaje else None,
            pesaje.tara if pesaje else None,
            pesaje.descuento if pesaje else None,
            float(m.total) if m.total is not None else None,
        ])

    formatear_celda_fecha(wb, ws, 'B')
    for columna in ('F', 'G', 'H', 'I'):
        formatear_celda_numero(ws, columna)
    definir_estilo_general(ws)

    for columna in ('D', 'E'):
        ws.column_dimensions[columna].width = 45
        for cell in ws[columna]:
            cell.alignment = Alignment(wrap_text=True, vertical='top')

    # Negrita para las líneas de filtros de la cabecera (se aplica después de
    # definir_estilo_general para que no se pise con la fuente general).
    fuente_filtro = Font(name='Arial', size=8, bold=True)
    for fila in range(1, fila_encabezado - 1):
        for cell in ws[fila]:
            cell.font = fuente_filtro

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=ingreso_hv_yerba_mate.xlsx'
    wb.save(response)
    return response


def reporte_hv_yerba_mate_excel_yerba(request):
    """Excel simplificado 'Reporte Yerba': Fecha, Nombre de la entidad
    emisora, Bruto, Tara (= tara - descuento), Neto (= bruto - tara),
    Declarado (igual al neto) y Número. Usa los mismos filtros que el
    reporte detallado (fecha desde/hasta y entidad emisora)."""
    import openpyxl
    from django.http import HttpResponse
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from services.gestorexcel import definir_estilo_general, formatear_celda_fecha

    _, movimientos = _filtrar_reporte_hv_yerba_mate(request)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Reporte Yerba'
    columnas = ['Fecha', 'Nombre', 'Bruto', 'Tara', 'Neto', 'Declarado', 'Número']
    ws.append(columnas)

    bruto_total = 0
    tara_total = 0
    neto_total = 0

    for m in _enriquecer_movimientos(movimientos[:2000]):
        pesaje = m.pesaje
        bruto = (pesaje.bruto or 0) if pesaje else 0
        tara = ((pesaje.tara or 0) - (pesaje.descuento or 0)) if pesaje else 0
        neto = bruto - tara
        bruto_total += bruto
        tara_total += tara
        neto_total += neto
        ws.append([
            m.fecha,
            m.entidad_emisor.nombre if m.entidad_emisor_id and m.entidad_emisor.nombre else '',
            bruto,
            tara,
            neto,
            neto,
            m.numero,
        ])

    fila_totales = ws.max_row + 1
    ws.append(['', 'Total', bruto_total, tara_total, neto_total, neto_total, ''])

    formatear_celda_fecha(wb, ws, 'A')
    for columna in ('C', 'D', 'E', 'F'):
        for cell in ws[columna]:
            if cell.row > 1:
                cell.number_format = '#,##0'

    definir_estilo_general(ws)

    # Resalta Neto/Declarado en verde y Número con fondo oscuro (solo en las
    # filas de datos), como en el formato de referencia que se usa para este
    # reporte.
    fuente_verde = Font(name='Arial', size=8, color='1B7A1B', bold=True)
    relleno_numero = PatternFill(start_color='1F3864', end_color='1F3864', fill_type='solid')
    fuente_numero = Font(name='Arial', size=8, color='FFFFFF', bold=True)
    for columna in ('E', 'F'):
        for cell in ws[columna]:
            if 1 < cell.row < fila_totales:
                cell.font = fuente_verde
    for cell in ws['G']:
        if 1 < cell.row < fila_totales:
            cell.fill = relleno_numero
            cell.font = fuente_numero
            cell.alignment = Alignment(horizontal='center')

    # Fila de totales: negrita y separada con un borde superior.
    fuente_total = Font(name='Arial', size=8, bold=True)
    borde_superior = Border(top=Side(style='thin', color='000000'))
    for celda in ws[fila_totales]:
        celda.font = fuente_total
        celda.border = borde_superior

    ws.column_dimensions['B'].width = 28

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=reporte_yerba.xlsx'
    wb.save(response)
    return response


def reporte_hv_yerba_mate_pdf(request):
    from django.http import HttpResponse
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import legal, landscape
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from .templatetags.movimientos_extras import separador_miles

    form, movimientos = _filtrar_reporte_hv_yerba_mate(request)
    filtros = _filtros_aplicados_reporte(form, movimientos)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename=ingreso_hv_yerba_mate.pdf'

    # Hoja tamaño legal apaisada (más ancha que A4 apaisado) y márgenes chicos,
    # para que las columnas entren sin desbordar el margen del papel.
    doc = SimpleDocTemplate(
        response, pagesize=landscape(legal),
        topMargin=0.6 * cm, bottomMargin=0.6 * cm, leftMargin=0.6 * cm, rightMargin=0.6 * cm,
    )
    estilos = getSampleStyleSheet()
    estilo_celda = ParagraphStyle('celda_reporte', parent=estilos['Normal'], fontSize=6, leading=7)
    estilo_filtro = ParagraphStyle('filtro_reporte', parent=estilos['Normal'], fontSize=9, leading=12)
    elementos = [
        Paragraph('Reporte de Ingreso H.V. de Yerba Mate', estilos['Title']),
    ]
    # Filtros aplicados: reemplazan a la columna Producto (que antes se
    # repetía igual en todas las filas, ya que este reporte siempre filtra
    # por el mismo producto).
    for etiqueta, valor in filtros:
        elementos.append(Paragraph(f"<b>{etiqueta}:</b> {valor}", estilo_filtro))
    elementos.append(Spacer(1, 0.4 * cm))

    encabezados = [
        'ID', 'Fecha', 'Número', 'Entidad Emisora / Inym Operador',
        'Entidad Receptora / Inym Operador',
        'Bruto', 'Tara', 'Descuento', 'Total',
    ]
    filas = [encabezados]
    total_general = Decimal('0')

    def _celda_entidad(entidad_txt, operador_txt):
        texto = entidad_txt if not operador_txt else f"{entidad_txt}<br/>{operador_txt}"
        return Paragraph(texto, estilo_celda)

    for m in _enriquecer_movimientos(movimientos[:2000]):
        pesaje = m.pesaje
        filas.append([
            str(m.id_movimiento),
            m.fecha.strftime('%d/%m/%Y') if m.fecha else '-',
            str(m.numero) if m.numero else '-',
            _celda_entidad(m.emisor_entidad_txt, m.emisor_operador_txt),
            _celda_entidad(m.receptor_entidad_txt, m.receptor_operador_txt),
            separador_miles(pesaje.bruto) if pesaje and pesaje.bruto is not None else '-',
            separador_miles(pesaje.tara) if pesaje and pesaje.tara is not None else '-',
            separador_miles(pesaje.descuento) if pesaje and pesaje.descuento is not None else '-',
            separador_miles(m.total) if m.total is not None else '-',
        ])
        if m.total:
            total_general += m.total

    filas.append(['', '', '', '', '', '', '', 'Total', separador_miles(total_general)])

    columnas_cm = [1.0, 1.8, 1.6, 10.4, 10.4, 2.0, 2.0, 2.0, 2.2]
    tabla = Table(filas, repeatRows=1, colWidths=[c * cm for c in columnas_cm])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#343a40')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, 0), 7),
        ('FONTSIZE', (0, 1), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#f2f2f2')]),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('ALIGN', (0, 0), (2, -1), 'CENTER'),
        ('ALIGN', (5, 0), (-1, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elementos.append(tabla)
    doc.build(elementos)
    return response


class MovimientoHvYerbaMateDetailView(DetailView):
    model = MovimientoHvYerbaMate
    template_name = 'movimientos/movimiento_hv_yerba_mate.html'
    context_object_name='movimiento'
    
    def get_context_data(self, **kwargs):
        # 1. Llama al contexto base primero
        context = super().get_context_data(**kwargs)
        print("esta aqui en detail mov yerba mate",self.object.pk)
        # 2. Agrega tus valores o clases relacionadas al contexto
        pesaje = MovimientoPesaje.objects.get(pk=self.object.pk)
        if pesaje:
        # Pongo en bruto tara y descuento lo que hay en pesaje si hay
            context['movimiento'].bruto = pesaje.bruto
            context['movimiento'].tara = pesaje.tara
            context['movimiento'].descuento = pesaje.descuento
            
        return context
def salida(request):
    form = SalidaForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        movimiento = form.save(commit=False)
        # El emisor de una salida es siempre la propia empresa; se usa la
        # misma entidad "propia" (Inym_Operador id=181) que ya se usa como
        # constante en IngresoHvYerbaMateForm y SalidaYerbaMateCanchadaForm.
        try:
            movimiento.entidad_emisor = Inym_Operador.objects.get(id=181).entidad
        except Inym_Operador.DoesNotExist:
            movimiento.entidad_emisor = None
        try:
            movimiento.save()
            messages.success(
                request,
                f'La salida {movimiento.numero or movimiento.id_movimiento} se ha registrado correctamente'
            )
            return redirect('movimientos:salida')
        except IntegrityError:
            form.add_error(None, "Ya existe un registro con estos datos o viola una restricción de integridad.")

    return render(request, 'movimientos/salida.html', {'form': form})
def accion_listado_movimiento(request):
     if request.POST.get('boton')=='completo' or request.POST.get('boton')=='vinculo':
         return generar_excel_completo(request)
     elif request.POST.get('boton')=='eliminar':
         return remove_movimiento(request)

def generar_excel_inym_emisor_agrupado(request):
    print("PRO DONDE VA")
    import openpyxl
    from django.http import HttpResponse
    from openpyxl.styles import NamedStyle
    print("qara")
     # 1. Crear un libro de trabajo y una hoja
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte de Datos"
    # 2. Agregar encabezados
    columnas = ['Producto', 'OP.Inym Emisor','Emisor','Total']
    ws.append(columnas)
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    
    response['Content-Disposition'] = 'attachment; filename=reporte.xlsx'

    # 5. Guardar el libro en la respuesta
    wb.save(response)
    return response
             
def generar_excel_completo(request):
    import openpyxl
    from django.http import HttpResponse
    from openpyxl.styles import NamedStyle

     # 1. Crear un libro de trabajo y una hoja
    wb = openpyxl.Workbook()
    
    if request.POST.get('boton')=='completo':
        ws = wb.active
        ws.title = "Reporte de Datos"
        # 2. Agregar encabezados
        columnas = ['ID', 'Numero', 'Fecha','Producto','ID Emisor','Op Inym Emisor','Emisor','ID Receptor','Op Inym Receptor','Receptor','Bruto','Tara','Descuento','Total']
        ws.append(columnas)

        # 3. Obtener los datos (por ejemplo, desde un modelo)
        ids = request.POST.getlist('movimientos_ids')
        datos = Movimiento.objects.filter(id_movimiento__in=ids).select_related('movimiento_pesaje').values_list('id_movimiento','numero','fecha','producto__nombre',
                                                                                        'entidad_emisor__id','movimientohvyerbamate__inym_operador_origen__id','entidad_emisor__nombre',
                                                                                        'entidad_receptor__id','movimientohvyerbamate__inym_operador_destino__id','entidad_receptor__nombre','movimiento_pesaje__bruto','movimiento_pesaje__tara','movimiento_pesaje__descuento','total')
        for fila in datos:
            ws.append(fila)
    elif request.POST.get('boton')=='vinculo':
        ws0 = wb.active
        ws0.title = "Productos en renglones y movimientos"
        columnas = ['Producto','En renglones','En movimientos']
        ws0.append(columnas)
        ids = request.POST.getlist('movimientos_ids')
        id_emisor = Movimiento.objects.filter(id_movimiento__in=ids).select_related('movimiento_pesaje').values_list('entidad_emisor__id', flat=True).first()
        productos_distintos_en_movimientos =  Movimiento.objects.filter(id_movimiento__in=ids).values_list("producto__nombre", flat=True).distinct()
        productos_distintos_en_renglones =  comprobantes_emisor = ComprobanteRenglon.objects.filter(comprobante__entidad_emisor__id=id_emisor).select_related('comprobante_renglon_detalle').values_list("producto__nombre", flat=True).distinct()
        productos_dicc = {}
        for producto in productos_distintos_en_movimientos:
            if producto not in productos_dicc.keys():
                productos_dicc[producto]=['No','Si']
        for producto in productos_distintos_en_renglones:
            if producto not in productos_dicc.keys():
                productos_dicc[producto]=['Si','No']
            else:
                productos_dicc[producto]='Si','Si'
        for k, v in productos_dicc.items():
                ws0.append([k,v[0],v[1]])
                #por cada producto voy a crear una hoja
                ws = wb.create_sheet(title=f"Vinculos comprobantes y movimientos {k}")
                columnas = ['Tipo', 'Id', 'Fecha','P.de Venta','Número','Producto','U.de Medida','Debe','Haber','Saldo']
                ws.append(columnas)
                movimientos_emisor = Movimiento.objects.filter(entidad_emisor__id=id_emisor,producto__nombre=k).values_list('id_movimiento','fecha','numero','producto__nombre','unidad_de_medida__nombre','total').order_by('fecha','id_movimiento')  
                comprobantes_emisor = ComprobanteRenglon.objects.filter(comprobante__entidad_emisor__id=id_emisor,producto__nombre=k).select_related('comprobante_renglon_detalle').values_list('comprobante_id','comprobante__fecha','comprobante__punto_de_venta',
                                                                                                                                                                     'comprobante__numero','producto__nombre','renglon_detalle_comprobante__unidad_de_medida__nombre',
                                                                                                                                     'renglon_detalle_comprobante__cantidad').order_by('comprobante__fecha','id').order_by('comprobante__fecha','id')
                lista_completa = []
                for fila in movimientos_emisor:
                    #convierto en lista para agregar datos con indices, pq tengo una tupla que es inmutable
                    fila = list(fila)
                    fila.insert(0,'Movimiento')
                    fila.insert(3,'-')
                    fila.insert(7,float(0))
                    lista_completa.append(fila)
                for fila in comprobantes_emisor:
                    fila = list(fila)
                    fila.insert(0,'Renglón de comprobante')
                    fila.insert(8,float(0))
                    lista_completa.append(fila)
                #ordeno por fecha y ahi recién hago saldo
                lista_completa = sorted(lista_completa, key=lambda x: x[2]) #indice de fecha es el 2
                saldo=0
                for renglon in lista_completa:
                    print("ID",renglon[0],renglon[1])
                    # si es IVa no tiene cantidad entonces comprobar para que no sea None
                    if renglon[7]!=None:
                        saldo+=Decimal(renglon[8])-Decimal(renglon[7])
                    else:
                        saldo+=Decimal(renglon[8])
                    renglon.insert(9,saldo)
                    ws.append(renglon)
    elif request.POST.get('boton')=='saldo_producto':
        generar_saldo_por_producto(request,wb)
        # 3. Obtener los datos (por ejemplo, desde un modelo)
        
        
    if request.POST.get('boton')!='saldo_producto': #saldo producto no aplica este formato fecha            
        #le doy formato de fecha a la celda del date
        # Create the style
        date_style = NamedStyle(name="custom_date", number_format="DD/MM/YYYY")
        wb.add_named_style(date_style)

    # Apply it to the column cells en este caso la columna es la C
        
        for cell in ws["C"]:
            if cell.row == 1:
                continue  # Skip header
            cell.style = date_style

     # 4. Preparar la respuesta HTTP para descargar el archivo
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    
    response['Content-Disposition'] = 'attachment; filename=reporte.xlsx'

    # 5. Guardar el libro en la respuesta
    wb.save(response)
    return response

def crear_excel():
    import openpyxl
        # 1. Crear un libro de trabajo y una hoja
    return openpyxl.Workbook()

def generar_excel_producto_saldo(request):
    wb = crear_excel()
    from django.http import HttpResponse
    generar_saldo_por_producto(request,wb)
    response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        
    response['Content-Disposition'] = 'attachment; filename=reporte.xlsx'

    # 5. Guardar el libro en la respuesta
    wb.save(response)
    return response

def generar_saldo_por_producto(request,wb):
    ws0 = wb.active
    ws0.title = "Saldo por Productos"
    columnas = ['Producto','En renglones','En movimientos']
    ws0.append(columnas)
    producto = request.POST.get('producto')
    if producto:
        movimientos = Movimiento.objects.filter(producto=producto)
        comprobantes_renglon = ComprobanteRenglon.objects.filter(producto=producto)
    else:
        movimientos = Movimiento.objects.all()
        comprobantes_renglon = ComprobanteRenglon.objects.all()
    fecha_desde = request.POST.get('fecha_desde')
    print("fecha_desde",fecha_desde)
    if fecha_desde:
        movimientos = movimientos.filter(fecha__gte=fecha_desde)#mayor o igual
        comprobantes_renglon = comprobantes_renglon.filter(comprobante__fecha__gte=fecha_desde)#mayor o igual
    fecha_hasta = request.POST.get('fecha_hasta')
    print("fecha_hasta",fecha_hasta)
    if fecha_hasta:
        movimientos = movimientos.filter(fecha__lte=fecha_hasta)  # menor o igual
        comprobantes_renglon = comprobantes_renglon.filter(comprobante__fecha__lte=fecha_hasta)#menor o igual
    print("comprobantes_emisor",comprobantes_renglon)
    #productos = Movimiento.objects.filter(id_movimiento__in=ids).select_related('producto').distinct()
    productos_distintos_en_movimientos = movimientos.values_list("producto__nombre", flat=True).distinct()
    #solo voy a trabajar con productos que tengas movimientos
    #productos_distintos_en_renglones = comprobantes_renglon.select_related('comprobante_renglon_detalle').values_list("producto__nombre", flat=True).distinct()
    productos_dicc = {}
    for producto in productos_distintos_en_movimientos:
        if producto not in productos_dicc.keys():
            productos_dicc[producto]={'debe':0,'haber':0,'filas':[]}
        """    
        for producto in productos_distintos_en_renglones:
            if producto not in productos_dicc.keys():
                productos_dicc[producto]={'debe':0,'haber':0,'filas':[]}
        """
    for k in productos_dicc.keys():
        from django.db.models import Sum
        ws0.append([k])
        #por cada producto voy a crear una hoja
        ws = wb.create_sheet(title=f"{k}")
        columnas = ['Id Entidad', 'Entidad','U.de Medida','Debe','Haber','Saldo']
        ws.append(columnas)
        #movimientos_emisor = movimientos.filter(producto__nombre=k).values('entidad_emisor').annotate(total=Sum('total')).values_list('entidad_emisor__id','entidad_emisor__nombre','unidad_de_medida__nombre','total').order_by('entidad_emisor__id,')  
        #agrupor por id emisor y despues algo el sum HABER
        movimientos_emisor = movimientos.values('entidad_emisor').annotate(total_haber=Sum('total')).values_list('entidad_emisor__id','entidad_emisor__nombre','unidad_de_medida__nombre','total_haber')
        
        comprobantes_emisor = comprobantes_renglon.values('comprobante__entidad_emisor').annotate(total_debe=Sum('renglon_detalle_comprobante__cantidad')).values_list('comprobante__entidad_emisor_id','comprobante__entidad_emisor__nombre','renglon_detalle_comprobante__unidad_de_medida__nombre','total_debe').order_by('comprobante__entidad_emisor__id')  
        fila_por_producto = {}
        
        for fila in movimientos_emisor:
            #convierto en lista para agregar datos con indices, pq tengo una tupla que es inmutable
           
            fila_por_producto[fila[0]] = {'nombre':fila[1],'unidad':fila[2],'debe':0,'haber':fila[3],'saldo':fila[3]}

        for fila in comprobantes_emisor:
            if fila[0] in fila_por_producto.keys():
                print("FILA #",fila)
                valor_haber =Decimal(fila_por_producto[fila[0]]['saldo'])
                fila_por_producto[fila[0]]['debe'] = fila[3]
                fila_por_producto[fila[0]]['saldo']=valor_haber-fila[3]
            else:
                 #el saldo va en negativo ya que no hay saldo para restar a favor
                 fila_por_producto[fila[0]] = {'nombre':fila[1],'unidad':fila[2],'haber':0,'debe':fila[3],'saldo':-fila[3]}        
        #ordeno por nombre y ahi recién hago saldo
        fila_por_producto = dict(sorted(fila_por_producto.items(), key=lambda item: item[1]['nombre'].lower())
)
        for k,renglon in fila_por_producto.items():
            print("rengloncito",renglon)
            fila = [k,renglon['nombre'],renglon['unidad'],renglon['debe'],renglon['haber'],renglon['saldo']]
            ws.append(fila)

def generar_excel_vinculo_con_movimiento(request):
    
    from openpyxl.styles import NamedStyle
    wb = crear_excel()
    ws = wb.active
    ws.title = "Reporte de Datos"

    # 2. Agregar encabezados
    columnas = ['Tipo', 'Id', 'Fecha','P.de Venta','Número','Debe','Haber','Saldo']
    ws.append(columnas)

    # 3. Obtener los datos (por ejemplo, desde un modelo)
    ids = request.POST.getlist('movimientos_ids')
    id_emisor = Movimiento.objects.filter(id_movimiento__in=ids).select_related('movimiento_pesaje').values_list('entidad_emisor__id', flat=True).first().order_by('fecha','id')
    movimientos_emisor = Movimiento.objects.filter(entidad_emisor__id=id_emisor)       
    comprobantes_emisor = Movimiento.filter(entidad_emisor__id=id_emisor).select_related('comprobante_renglon_detalle').values_list('id_comprobante','comprobante__fecha','comprobante__punto_de_venta','comprobante__numero',
                                                                                                                                    'total').order_by('comprobante__fecha','id')
    movimientos_emisor = [['movimiento'] + fila for fila in movimientos_emisor]
    comprobantes_emisor = [['comprobante'] + fila for fila in comprobantes_emisor]
    #le doy formato de fecha a la celda del date
    # Create the style
    date_style = NamedStyle(name="custom_date", number_format="DD/MM/YYYY")
    wb.add_named_style(date_style)

# Apply it to the column cells en este caso la columna es la C
    for cell in ws["C"]:
        if cell.row == 1:
            continue  # Skip header
        cell.style = date_style

     # 4. Preparar la respuesta HTTP para descargar el archivo
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    
    response['Content-Disposition'] = 'attachment; filename=reporte.xlsx'

    # 5. Guardar el libro en la respuesta
    wb.save(response)
    return response

def filtrar_comprobantes_renglones(request):
    producto = request.POST.get('producto')
    if producto:
        comprobantes_renglon = ComprobanteRenglon.objects.filter(producto=producto)
    else:
        comprobantes_renglon = ComprobanteRenglon.objects.all()
    fecha_desde = request.POST.get('fecha_desde')
    if fecha_desde:
        comprobantes_renglon = comprobantes_renglon.filter(comprobante__fecha__gte=fecha_desde)#mayor o igual
    fecha_hasta = request.POST.get('fecha_hasta')
    if fecha_hasta:
        comprobantes_renglon = comprobantes_renglon.filter(comprobante__fecha__lte=fecha_hasta)#menor o igual
    entidad_emisor = request.POST.get('emisor')
    if entidad_emisor:
       comprobantes_renglon = comprobantes_renglon.filter(comprobante__entidad_emisor=entidad_emisor)
    return comprobantes_renglon


def filtrar_movimientos(request):
    producto = request.POST.get('producto')
    if producto:
        movimientos = Movimiento.objects.filter(producto=producto)
    else:
        movimientos = Movimiento.objects.all()
    fecha_desde = request.POST.get('fecha_desde')
    if fecha_desde:
        movimientos = movimientos.filter(fecha__gte=fecha_desde)#mayor o igual
    fecha_hasta = request.POST.get('fecha_hasta')
    if fecha_hasta:
        movimientos = movimientos.filter(fecha__lte=fecha_hasta)  # menor o igual
    entidad_emisor = request.POST.get('emisor')
    if entidad_emisor:
        movimientos = movimientos.filter(entidad_emisor__id=entidad_emisor)
    return movimientos

def vista_buscar_saldo_producto_entidad(request):
    from .forms import BuscarSaldoProductoEntidadForm
    if request.method == 'POST':
        # Vinculamos los datos del usuario al formulario
        form = BuscarSaldoProductoEntidadForm(request.POST)
        if form.is_valid():
            # Procesar los datos limpios (form.cleaned_data)
            # O guardar directamente si es un ModelForm 
            contexto = None
            movimientos = filtrar_movimientos(request)
            comprobantes_renglon = filtrar_comprobantes_renglones(request)
            nombre_hoja = f"{request.POST.get('emisor')} - {request.POST.get('producto')}"
            if movimientos or comprobantes_renglon:
                tipo_accion = request.POST.get('accion')
                if tipo_accion=='Generar Excel':
                    return generar_vinculo_movimientos_comprobantes(nombre_hoja,movimientos,comprobantes_renglon)
                    
            
            #redirect('movimientos:listado_movimiento')
    else:
        print("CONSOLA ENTRA")
        # Formulario vacío inicial
        form = BuscarSaldoProductoEntidadForm()

    producto_id = form['producto'].value()
    producto_texto = str(ProductoDetalle.objects.filter(pk=producto_id).first() or '') if producto_id else ''

    return render(request, 'movimientos/producto_por_entidad_buscar_saldo.html', {'form': form, 'producto_texto': producto_texto})

def generar_vinculo_movimientos_comprobantes(nombre_hoja,movimientos,comprobantes_renglones):
    from openpyxl.styles import NamedStyle
    wb = crear_excel()
    ws = wb.active
    ws.title = "Reporte de Datos"

    wb = generar_relacion_producto_entidad_saldo(wb,nombre_hoja,movimientos,comprobantes_renglones)
    #le doy formato de fecha a la celda del date
    # Create the style
    date_style = NamedStyle(name="custom_date", number_format="DD/MM/YYYY")
    wb.add_named_style(date_style)

    # Apply it to the column cells en este caso la columna es la C
    for cell in ws["C"]:
        if cell.row == 1:
            continue  # Skip header
        cell.style = date_style

        # 4. Preparar la respuesta HTTP para descargar el archivo
    response = crear_response_excel()
    # 5. Guardar el libro en la respuesta
    wb.save(response)
    return response
  
    

def generar_relacion_producto_entidad_saldo(wb,nombre_hoja,movimientos,comprobantes_renglones):
    ws = wb.create_sheet(title=f"Vinculos comprobantes y movimientos {nombre_hoja}")
    columnas = ['Tipo', 'Id', 'Fecha','P.de Venta','Número','Producto','U.de Medida','Debe','Haber','Saldo']
    ws.append(columnas)
    movimientos = movimientos.values_list('id_movimiento','fecha','numero','producto__nombre','unidad_de_medida__nombre','total').order_by('fecha','id_movimiento')  
    comprobantes_renglones = comprobantes_renglones.select_related('comprobante_renglon_detalle').values_list('comprobante_id','comprobante__fecha','comprobante__punto_de_venta',
                                                                                                                                                        'comprobante__numero','producto__nombre','renglon_detalle_comprobante__unidad_de_medida__nombre',
                                                                                                                        'renglon_detalle_comprobante__cantidad').order_by('comprobante__fecha','id').order_by('comprobante__fecha','id')
    lista_completa = []
    for fila in movimientos:
        #convierto en lista para agregar datos con indices, pq tengo una tupla que es inmutable
        fila = list(fila)
        fila.insert(0,'Movimiento')
        fila.insert(3,'-')
        fila.insert(7,float(0))
        lista_completa.append(fila)
    for fila in comprobantes_renglones:
        fila = list(fila)
        fila.insert(0,'Renglón de comprobante')
        fila.insert(8,float(0))
        lista_completa.append(fila)
    #ordeno por fecha y ahi recién hago saldo
    lista_completa = sorted(lista_completa, key=lambda x: x[2]) #indice de fecha es el 2
    saldo=0
    for renglon in lista_completa:
        print("ID",renglon[0],renglon[1])
        # si es IVa no tiene cantidad entonces comprobar para que no sea None
        if renglon[7]!=None:
            saldo+=Decimal(renglon[8])-Decimal(renglon[7])
        else:
            saldo+=Decimal(renglon[8])
        renglon.insert(9,saldo)
        ws.append(renglon)
    print("llegando 6")


# ---------------------------------------------------------------------------
# Alta / Modificación / Reportes de Movimiento (gestión genérica de productos)
# ---------------------------------------------------------------------------

@transaction.atomic
def movimiento_form(request, pk=None):
    """Alta y modificación de un Movimiento de producto (misma vista, pk=None para alta).

    Además de los datos básicos del movimiento, permite cargar el pesaje
    (bruto/tara/descuento, con el total calculado en el cliente) y, cuando
    corresponde, los operadores INYM emisor/receptor de un movimiento de
    H.V. de Yerba Mate. En ese caso, el emisor/receptor (Entidad) del
    movimiento se deriva de los operadores elegidos.
    """
    movimiento = get_object_or_404(Movimiento, pk=pk) if pk else None
    movimiento_hv_existente = MovimientoHvYerbaMate.objects.filter(pk=movimiento.pk).first() if movimiento else None
    pesaje_existente = MovimientoPesaje.objects.filter(pk=movimiento.pk).first() if movimiento else None

    if request.method == 'POST':
        form = MovimientoForm(request.POST, instance=movimiento)
        pesaje_form = PesajeInlineForm(request.POST, instance=pesaje_existente)
        es_yerba_mate = request.POST.get('es_yerba_mate') == 'on'
        operadores_form = MovimientoInymOperadoresForm(request.POST)

        operador_origen = operador_destino = None
        operadores_ok = True
        if es_yerba_mate:
            operadores_ok = operadores_form.is_valid()
            if operadores_ok:
                operador_origen = operadores_form.cleaned_data.get('inym_operador_origen')
                operador_destino = operadores_form.cleaned_data.get('inym_operador_destino')
                if not (operador_origen and operador_destino):
                    operadores_form.add_error(None, 'Seleccioná el operador INYM emisor y el receptor.')
                    operadores_ok = False

        if form.is_valid() and pesaje_form.is_valid() and operadores_ok:
            nuevo = form.save(commit=False)
            if es_yerba_mate:
                nuevo.entidad_emisor = operador_origen.entidad
                nuevo.entidad_receptor = operador_destino.entidad
            nuevo.save()

            bruto = pesaje_form.cleaned_data.get('bruto')
            tara = pesaje_form.cleaned_data.get('tara')
            descuento = pesaje_form.cleaned_data.get('descuento')
            if any(valor not in (None, '') for valor in (bruto, tara, descuento)):
                MovimientoPesaje.objects.update_or_create(
                    pk=nuevo.pk,
                    defaults={'bruto': bruto, 'tara': tara, 'descuento': descuento},
                )

            if es_yerba_mate:
                if movimiento_hv_existente:
                    # Ya existe la fila hija (movimiento_hv_yerba_mate): actualizo solo
                    # sus columnas propias con un UPDATE directo, sin pasar por
                    # instance.save(). Si se guardara la instancia completa acá, la
                    # herencia multi-tabla de Django volvería a grabar también la fila
                    # de 'movimiento' (producto, fecha, total, etc.), y el método
                    # Movimiento.save() sobrescribe entidad_emisor/receptor leyendo de
                    # nuevo la fila (todavía vieja en ese punto) -> mejor evitarlo.
                    MovimientoHvYerbaMate.objects.filter(pk=nuevo.pk).update(
                        inym_operador_origen=operador_origen,
                        inym_operador_destino=operador_destino,
                    )
                else:
                    # No existe todavía: creo la fila hija completando TODOS los campos
                    # heredados de Movimiento con los valores ya guardados en 'nuevo'.
                    # Si se omiten (como pasaba antes), Django los guarda como None al
                    # volver a grabar la fila padre por la herencia multi-tabla, lo que
                    # provoca 'Column id_producto cannot be null'.
                    MovimientoHvYerbaMate(
                        movimiento=nuevo,
                        fecha=nuevo.fecha,
                        numero=nuevo.numero,
                        producto=nuevo.producto,
                        total=nuevo.total,
                        unidad_de_medida=nuevo.unidad_de_medida,
                        entidad_emisor=operador_origen.entidad,
                        entidad_receptor=operador_destino.entidad,
                        inym_operador_origen=operador_origen,
                        inym_operador_destino=operador_destino,
                    ).save()

            messages.success(request, f'Movimiento {nuevo.id_movimiento} guardado correctamente.')
            return redirect('movimientos:movimiento_modificar')
    else:
        form = MovimientoForm(instance=movimiento)
        pesaje_form = PesajeInlineForm(instance=pesaje_existente)
        es_yerba_mate = movimiento_hv_existente is not None
        operadores_initial = {}
        if movimiento_hv_existente:
            operadores_initial = {
                'inym_operador_origen': movimiento_hv_existente.inym_operador_origen_id,
                'inym_operador_destino': movimiento_hv_existente.inym_operador_destino_id,
            }
        operadores_form = MovimientoInymOperadoresForm(initial=operadores_initial)

    operador_entidad_map = {operador.id: operador.entidad_id for operador in Inym_Operador.objects.all()}

    # Defaults para el producto 'HOJA VERDE DE YERBA MATE PUESTA EN SECADERO' (id=2):
    # el receptor siempre es el operador INYM de Fontana como Secadero (id=181, la
    # misma constante que ya se usa en IngresoHvYerbaMateForm, SalidaYerbaMateCanchadaForm
    # y en la vista salida()); el emisor por defecto es el operador de tipo 'PRODUCTORES'
    # de la entidad emisora que ya esté seleccionada. Siempre se pueden cambiar antes de
    # guardar; el JS del template solo los precarga si el campo todavía está vacío.
    PRODUCTO_HOJA_VERDE_SECADERO_ID = 2
    OPERADOR_FONTANA_SECADERO_ID = 181
    operadores_productores_por_entidad = {
        operador.entidad_id: operador.id
        for operador in Inym_Operador.objects.filter(tipo_operador__nombre__iexact='PRODUCTORES')
    }

    producto_id = form['producto'].value()
    producto_texto = str(ProductoDetalle.objects.filter(pk=producto_id).first() or '') if producto_id else ''

    return render(request, 'movimientos/movimiento_gestion_form.html', {
        'form': form,
        'pesaje_form': pesaje_form,
        'operadores_form': operadores_form,
        'movimiento': movimiento,
        'es_yerba_mate': es_yerba_mate,
        'operador_entidad_map_json': json.dumps(operador_entidad_map),
        'producto_hoja_verde_secadero_id': PRODUCTO_HOJA_VERDE_SECADERO_ID,
        'operador_fontana_secadero_id': OPERADOR_FONTANA_SECADERO_ID,
        'operadores_productores_por_entidad_json': json.dumps(operadores_productores_por_entidad),
        'producto_texto': producto_texto,
    })


def movimiento_listado(request):
    """Listado/búsqueda de movimientos de producto; es la puerta de entrada de 'Modificación'."""
    movimientos = (
        Movimiento.objects.select_related('producto', 'entidad_emisor', 'entidad_receptor', 'unidad_de_medida')
        .order_by('-fecha', '-id_movimiento')
    )

    q_receptor = request.GET.get('receptor', '').strip()
    q_id = request.GET.get('id', '').strip()
    q_fecha = request.GET.get('fecha', '').strip()
    q_numero = request.GET.get('numero', '').strip()

    if q_receptor:
        movimientos = movimientos.filter(
            Q(entidad_receptor__nombre__icontains=q_receptor) | Q(entidad_receptor__cuit__icontains=q_receptor)
        )
    if q_id:
        if q_id.isdigit():
            movimientos = movimientos.filter(id_movimiento=int(q_id))
        else:
            movimientos = movimientos.none()
    if q_fecha:
        movimientos = movimientos.filter(fecha=q_fecha)
    if q_numero:
        if q_numero.isdigit():
            movimientos = movimientos.filter(numero=int(q_numero))
        else:
            movimientos = movimientos.none()

    movimientos = aplicar_orden_queryset(request, movimientos, {
        'id': 'id_movimiento',
        'fecha': 'fecha',
        'numero': 'numero',
        'producto': 'producto__nombre',
        'total': 'total',
        'emisor': 'entidad_emisor__nombre',
        'receptor': 'entidad_receptor__nombre',
        'unidad': 'unidad_de_medida__nombre',
    })

    return render(request, 'movimientos/movimiento_gestion_listado.html', {
        'movimientos': movimientos[:200],
        'q_receptor': q_receptor,
        'q_id': q_id,
        'q_fecha': q_fecha,
        'q_numero': q_numero,
    })


def movimiento_eliminar(request, pk):
    """Confirmación + baja de un Movimiento de producto, desde el listado de
    'Modificación' de la gestión genérica. Por el on_delete=CASCADE ya
    declarado en los modelos, arrastra también el pesaje asociado
    (movimiento_pesaje) y, si era de H.V. de Yerba Mate, esa fila hija
    también (mismo criterio que movimiento_hv_yerba_mate_eliminar).
    """
    from django.db import IntegrityError
    from django.db.models import ProtectedError

    movimiento = get_object_or_404(Movimiento, pk=pk)

    if request.method == 'POST':
        numero = movimiento.numero
        try:
            movimiento.delete()
        except (ProtectedError, IntegrityError):
            messages.error(
                request,
                f'El movimiento {pk} no se puede eliminar porque está siendo usado en otro registro.'
            )
        else:
            messages.success(request, f'El movimiento número {numero} con id {pk} se eliminó correctamente.')
        return redirect('movimientos:movimiento_modificar')

    return render(request, 'movimientos/movimiento_gestion_eliminar_confirm.html', {
        'movimiento': movimiento,
    })


def movimiento_reporte(request):
    form = MovimientoReporteForm(request.GET or None)
    movimientos = (
        Movimiento.objects.select_related('producto', 'entidad_emisor', 'entidad_receptor', 'unidad_de_medida')
        .order_by('-fecha', '-id_movimiento')
    )

    producto = None
    entidad_emisor = None
    entidad_receptor = None
    if form.is_valid():
        producto = form.cleaned_data.get('producto')
        entidad_emisor = form.cleaned_data.get('entidad_emisor')
        entidad_receptor = form.cleaned_data.get('entidad_receptor')
        fecha_desde = form.cleaned_data.get('fecha_desde')
        fecha_hasta = form.cleaned_data.get('fecha_hasta')
        if producto:
            movimientos = movimientos.filter(producto=producto)
        if entidad_emisor:
            movimientos = movimientos.filter(entidad_emisor=entidad_emisor)
        if entidad_receptor:
            movimientos = movimientos.filter(entidad_receptor=entidad_receptor)
        if fecha_desde:
            movimientos = movimientos.filter(fecha__gte=fecha_desde)
        if fecha_hasta:
            movimientos = movimientos.filter(fecha__lte=fecha_hasta)

    totales = movimientos.aggregate(
        total_cantidad=Coalesce(Sum('total'), Value(Decimal('0')), output_field=DecimalField(max_digits=12, decimal_places=2)),
    )

    movimientos = aplicar_orden_queryset(request, movimientos, {
        'id': 'id_movimiento',
        'fecha': 'fecha',
        'numero': 'numero',
        'producto': 'producto__nombre',
        'total': 'total',
        'emisor': 'entidad_emisor__nombre',
        'receptor': 'entidad_receptor__nombre',
    })

    return render(request, 'movimientos/movimiento_gestion_reporte.html', {
        'form': form,
        'movimientos': movimientos[:500],
        'totales': totales,
        'entidad_emisor_texto': texto_entidad_buscador(entidad_emisor),
        'entidad_receptor_texto': texto_entidad_buscador(entidad_receptor),
        'producto_texto': str(producto) if producto else '',
    })


# --- Ranking de productores (por total entregado, filtrando por fecha y producto) ---

def _movimientos_ranking_productores_filtrados(request):
    """Aplica a Movimiento los filtros de RankingProductoresForm (fecha de
    emisión y, opcionalmente, producto). Devuelve (form, queryset,
    filtros_activos), centralizado para que la pantalla y las exportaciones
    (Excel / PDF) usen siempre los mismos criterios."""
    form = RankingProductoresForm(request.GET or None)
    movimientos = Movimiento.objects.all()

    filtros_activos = False

    if form.is_valid():
        fecha_desde = form.cleaned_data.get('fecha_desde')
        fecha_hasta = form.cleaned_data.get('fecha_hasta')
        producto = form.cleaned_data.get('producto')
        if fecha_desde:
            movimientos = movimientos.filter(fecha__gte=fecha_desde)
        if fecha_hasta:
            movimientos = movimientos.filter(fecha__lte=fecha_hasta)
        if producto:
            movimientos = movimientos.filter(producto=producto)
        filtros_activos = bool(fecha_desde or fecha_hasta or producto)

    return form, movimientos, filtros_activos


def _calcular_ranking_productores(movimientos):
    """A partir de un queryset de Movimiento, arma el ranking de entidades
    emisoras (productores) por total entregado (de mayor a menor) y el total
    general. Devuelve (ranking, total_general)."""
    ranking = list(
        movimientos.values('entidad_emisor_id', 'entidad_emisor__nombre')
        .annotate(total_entregado=Sum('total'), cantidad=Count('id_movimiento'))
        .order_by('-total_entregado')
    )

    total_general = sum((fila['total_entregado'] for fila in ranking), Decimal('0'))
    for posicion, fila in enumerate(ranking, start=1):
        fila['posicion'] = posicion
        fila['porcentaje'] = (fila['total_entregado'] / total_general * 100) if total_general else Decimal('0')

    return ranking, total_general


@requiere_grupo('Rankings')
def movimiento_ranking_productores(request):
    """Ranking de productores (entidad emisora) según la suma del total de
    sus entregas, de mayor a menor, filtrando opcionalmente por un rango de
    fecha de emisión y por producto."""
    form, movimientos, filtros_activos = _movimientos_ranking_productores_filtrados(request)
    ranking, total_general = _calcular_ranking_productores(movimientos)
    ranking = aplicar_orden_lista(request, ranking, {
        'posicion': lambda f: f['posicion'],
        'productor': lambda f: (f['entidad_emisor__nombre'] or '').lower(),
        'cantidad': lambda f: f['cantidad'],
        'total': lambda f: f['total_entregado'],
        'porcentaje': lambda f: f['porcentaje'],
    })

    # Texto a mostrar en el buscador de producto: el ya elegido (viene de
    # ?producto=<id> en la URL), o vacío si no hay filtro de producto.
    producto_id = form['producto'].value()
    producto_texto = str(ProductoDetalle.objects.filter(pk=producto_id).first() or '') if producto_id else ''

    return render(request, 'movimientos/movimiento_ranking_productores.html', {
        'form': form,
        'ranking': ranking,
        'total_general': total_general,
        'filtros_activos': filtros_activos,
        'producto_texto': producto_texto,
    })


def _filas_ranking_productores(ranking):
    columnas = ['#', 'Productor', 'Entregas', 'Total entregado', 'Participación %']
    filas = [
        [
            fila['posicion'],
            fila['entidad_emisor__nombre'] or 'Sin nombre',
            fila['cantidad'],
            float(fila['total_entregado']) if fila['total_entregado'] is not None else None,
            float(fila['porcentaje']) if fila['porcentaje'] is not None else None,
        ]
        for fila in ranking
    ]
    return {
        'columnas': columnas,
        'filas': filas,
        'columnas_numericas': {3, 4},  # Total entregado, Participación %
        'anchos': [0.4, 2.2, 1.0, 1.2, 1.2],
    }


@requiere_grupo('Rankings')
def movimiento_ranking_productores_excel(request):
    _form, movimientos, _filtros_activos = _movimientos_ranking_productores_filtrados(request)
    ranking, _total_general = _calcular_ranking_productores(movimientos)
    resultado = _filas_ranking_productores(ranking)
    return excel_response('ranking_productores', resultado)


@requiere_grupo('Rankings')
def movimiento_ranking_productores_pdf(request):
    _form, movimientos, _filtros_activos = _movimientos_ranking_productores_filtrados(request)
    ranking, _total_general = _calcular_ranking_productores(movimientos)
    resultado = _filas_ranking_productores(ranking)
    return pdf_response('ranking_productores', 'Ranking de productores por total entregado', resultado)
    return wb