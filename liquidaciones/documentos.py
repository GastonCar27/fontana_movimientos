"""
Generación de PDF y Excel para una liquidación individual.

El formato del PDF replica el que genera el sistema de escritorio
"Fontana 1.0" (Mixin/Generador_Pdf.py::crear_reporte_liquidacion /
armar_tabla_pdf / armar_tabla_renglones), para que el documento impreso
desde acá sea igual al que ya conocen los usuarios. El Excel es un formato
"similar en espíritu": el sistema de escritorio no generaba un Excel para
una liquidación individual (solo un reporte masivo por rango de fechas con
pandas), así que acá se armó uno nuevo, ordenado en hojas Resumen/Debe/
Haber (+ Renglones si es completo), con el mismo criterio de formato
numérico que el resto de esta aplicación (separador de miles con punto,
decimales con coma).

Diferencia deliberada respecto del PDF: en el sistema de escritorio, el
límite de 5 renglones + "Otros" existía por una limitación de espacio en la
hoja del PDF. En el Excel "completo" no hay ese límite: se listan todos los
renglones de todos los comprobantes.
"""
import os
from decimal import Decimal

from django.core.exceptions import ObjectDoesNotExist
from django.http import HttpResponse

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from comprobantes.models import ComprobanteRenglonDetalle
from movimientos.templatetags.movimientos_extras import separador_miles
from movimientos_caja.models import MovimientoCajaBancoCuentaEntidad
from services.numero_a_letras import numero_a_moneda

# Máximo de renglones (líneas de producto) que se listan en el PDF antes de
# agregar el marcador "Otros" — mismo límite que usaba el sistema de
# escritorio (ahí era por espacio de página).
MAX_RENGLONES_PDF = 5

FORMATO_MILES_EXCEL = '#,##0.00'

DEBE_HEADERS = ['Id', 'Fecha', 'Comprobante', 'Entidad', 'Total']
HABER_HEADERS = ['Id', 'Fecha', 'F.diferido', 'Detalle', 'Entidad', 'Total']
RENGLONES_HEADERS = ['Id', 'Producto', 'Cantidad', 'U. De Medida', 'Precio Unitario', 'Iva Tipo', 'Sector']

# static/images/logo_fsa.jpg, relativo a la raíz del proyecto (este archivo
# vive en <raiz>/liquidaciones/documentos.py).
LOGO_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    'static', 'images', 'logo_fsa.jpg',
)


# ---------------------------------------------------------------------------
# Helpers de datos
# ---------------------------------------------------------------------------

def _monto_comprobante(comprobante):
    """Total del comprobante convertido a pesos (mismo criterio que
    liquidaciones.views._monto_item / Liquidacion.recalcular_totales): si el
    comprobante está en moneda distinta a pesos, se multiplica por el tipo
    de cambio guardado en comprobante_tipo_de_cambio."""
    total = comprobante.total or Decimal('0')
    try:
        total = (total * comprobante.tipo_de_cambio.tipo_de_cambio).quantize(Decimal('0.01'))
    except ObjectDoesNotExist:
        pass
    return total


def _cuit_con_guiones(cuit):
    digitos = ''.join(ch for ch in str(cuit) if ch.isdigit())
    if len(digitos) == 11:
        return f'{digitos[0:2]}-{digitos[2:10]}-{digitos[10:]}'
    return str(cuit)


def _texto_comprobante(comprobante):
    if comprobante.comprobante_string:
        return comprobante.comprobante_string
    partes = []
    tipo = comprobante.tipo_comprobante
    if tipo and tipo.abreviatura:
        partes.append(tipo.abreviatura)
    if comprobante.punto_de_venta is not None and comprobante.numero is not None:
        partes.append(f'{comprobante.punto_de_venta:04d}-{comprobante.numero:08d}')
    return ' '.join(partes) if partes else f'Comprobante {comprobante.id}'


def _texto_movimiento(movimiento_caja):
    """Replica item.get_tipo_movimiento()+" "+item.get_numero(), con el caso
    especial de "Pago Proveedores": si el movimiento tiene una cuenta
    bancaria de destino cargada, se muestra ese número en vez del número
    propio del movimiento."""
    tipo_nombre = movimiento_caja.tipo.nombre if movimiento_caja.tipo_id else ''
    numero_a_mostrar = movimiento_caja.numero
    if tipo_nombre == 'Pago Proveedores':
        try:
            cuenta_destino = movimiento_caja.movimientocajabancocuentaentidad
        except MovimientoCajaBancoCuentaEntidad.DoesNotExist:
            cuenta_destino = None
        if cuenta_destino and cuenta_destino.numero_cuenta_entidad_destino:
            numero_a_mostrar = cuenta_destino.numero_cuenta_entidad_destino
    return f'{tipo_nombre} {numero_a_mostrar}'.strip()


def _diferido_movimiento(movimiento_caja):
    try:
        return movimiento_caja.movimientocajadiferido.diferido
    except Exception:
        return None


def _texto_retencion(retencion):
    """Número identificatorio de la retención, con año y número separados
    por guion (ej. "2026-15"), igual que lo mostraba el sistema de
    escritorio y que el __str__ de este mismo modelo. 'comprobante_string'
    no se usa acá porque en las retenciones ese campo suele venir vacío."""
    if retencion.año is not None and retencion.numero is not None:
        return f'Ret. {retencion.año}-{retencion.numero}'
    if retencion.comprobante_string:
        return f'Ret. {retencion.comprobante_string}'
    return f'Ret. {retencion.id}'


def _filas_debe(liquidacion):
    """Arma las filas del lado Debe ("Comprobantes a Pagar"), mezclando
    comprobantes, movimientos de caja, retenciones y retenciones INYM en el
    mismo orden que usaba el sistema de escritorio. Devuelve (filas, total)."""
    filas = []
    total = Decimal('0')
    entidad_texto = liquidacion.entidad.nombre if liquidacion.entidad else ''

    for lc in liquidacion.comprobantes.filter(tipo='debe').select_related(
        'comprobante__tipo_de_cambio', 'comprobante__tipo_comprobante'
    ):
        c = lc.comprobante
        monto = _monto_comprobante(c)
        total += monto
        filas.append([c.id, c.fecha, _texto_comprobante(c), entidad_texto, monto])

    for lm in liquidacion.movimientos.filter(tipo='debe').select_related(
        'movimiento_caja__tipo', 'movimiento_caja__movimientocajabancocuentaentidad'
    ):
        m = lm.movimiento_caja
        monto = m.monto or Decimal('0')
        total += monto
        filas.append([m.id, m.emision, _texto_movimiento(m), entidad_texto, monto])

    for lr in liquidacion.retenciones.filter(tipo='debe').select_related('retencion'):
        r = lr.retencion
        monto = r.total or Decimal('0')
        total += monto
        filas.append([r.id, r.fecha, _texto_retencion(r), entidad_texto, monto])

    for lri in liquidacion.retenciones_inym.filter(tipo='debe').select_related('retencion_inym'):
        ri = lri.retencion_inym
        monto = ri.total or Decimal('0')
        total += monto
        filas.append([ri.id, ri.fecha, f'Ret.Inym {ri.id}', entidad_texto, monto])

    return filas, total


def _filas_haber(liquidacion):
    """Arma las filas del lado Haber ("Pago"). Devuelve (filas, total)."""
    filas = []
    total = Decimal('0')
    entidad_texto = liquidacion.entidad.nombre if liquidacion.entidad else ''

    for lc in liquidacion.comprobantes.filter(tipo='haber').select_related(
        'comprobante__tipo_de_cambio', 'comprobante__tipo_comprobante'
    ):
        c = lc.comprobante
        monto = _monto_comprobante(c)
        total += monto
        filas.append([c.id, c.fecha, '-', _texto_comprobante(c), entidad_texto, monto])

    for lm in liquidacion.movimientos.filter(tipo='haber').select_related(
        'movimiento_caja__tipo',
        'movimiento_caja__movimientocajabancocuentaentidad',
        'movimiento_caja__movimientocajadiferido',
    ):
        m = lm.movimiento_caja
        monto = m.monto or Decimal('0')
        total += monto
        filas.append([m.id, m.emision, _diferido_movimiento(m), _texto_movimiento(m), entidad_texto, monto])

    for lr in liquidacion.retenciones.filter(tipo='haber').select_related('retencion'):
        r = lr.retencion
        monto = r.total or Decimal('0')
        total += monto
        filas.append([r.id, r.fecha, '-', _texto_retencion(r), entidad_texto, monto])

    for lri in liquidacion.retenciones_inym.filter(tipo='haber').select_related('retencion_inym'):
        ri = lri.retencion_inym
        monto = ri.total or Decimal('0')
        total += monto
        filas.append([ri.id, ri.fecha, '', f'Ret.Inym {ri.id}', entidad_texto, monto])

    return filas, total


def _renglones_comprobantes(comprobante_ids, limitar=None):
    """Renglones (líneas de producto) de los comprobantes dados, con las
    mismas columnas que usaba el sistema de escritorio (Repositorio/
    Comprobante_Renglon.py::obtener_campos_para_presentacion_renglones_liquidacion):
    Id, Producto, Cantidad, U. de Medida, Precio Unitario, Iva Tipo, Sector.
    Solo se incluyen renglones cuyo producto es de tipo "Producto"
    (id_item_tipo=1), igual que el filtro original.

    Si 'limitar' tiene valor, corta la lista a esa cantidad y devuelve
    hay_mas=True si había más renglones (para agregar el marcador "Otros").
    Sin 'limitar' (Excel completo) se devuelven todos.
    """
    if not comprobante_ids:
        return [], False

    detalles = (
        ComprobanteRenglonDetalle.objects.filter(
            comprobante_renglon__comprobante_id__in=comprobante_ids,
            comprobante_renglon__producto__item_tipo_id=1,
        )
        .select_related(
            'comprobante_renglon__producto',
            'unidad_de_medida',
            'iva_tipo',
            'sector_tipo',
        )
        .order_by('comprobante_renglon__comprobante_id', 'comprobante_renglon_id')
    )

    filas = []
    hay_mas = False
    for indice, detalle in enumerate(detalles):
        if limitar is not None and indice >= limitar:
            hay_mas = True
            break
        renglon = detalle.comprobante_renglon
        filas.append([
            renglon.id,
            renglon.producto.nombre if renglon.producto else '',
            detalle.cantidad,
            detalle.unidad_de_medida.nombre if detalle.unidad_de_medida else '',
            detalle.precio_unitario,
            detalle.iva_tipo.nombre if detalle.iva_tipo else '',
            detalle.sector_tipo.nombre if detalle.sector_tipo else '',
        ])
    return filas, hay_mas


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _tabla_pdf(headers, filas):
    datos = [headers]
    ultimo_indice = len(headers) - 1
    for fila in filas:
        fila_fmt = []
        for indice, valor in enumerate(fila):
            if indice == ultimo_indice:
                fila_fmt.append(separador_miles(valor))
            elif hasattr(valor, 'strftime'):
                fila_fmt.append(valor.strftime('%d/%m/%Y'))
            else:
                fila_fmt.append('' if valor is None else str(valor))
        datos.append(fila_fmt)

    tabla = Table(data=datos)
    tabla.setStyle(TableStyle([
        ('TEXTCOLOR', (0, 1), (1, -1), colors.black),
        ('FONTSIZE', (0, 0), (-1, -1), 6),
        ('TEXTCOLOR', (1, 2), (3, -1), colors.black),
        ('BACKGROUND', (1, 1), (-1, -1), colors.white),
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('BOX', (0, 0), (-1, -1), 1.25, colors.grey),
        ('INNERGRID', (0, 0), (-1, -1), 1, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    return tabla


def _tabla_renglones_pdf(filas, hay_mas):
    datos = [RENGLONES_HEADERS]
    for fila in filas:
        datos.append([
            fila[0],
            fila[1],
            separador_miles(fila[2]) if fila[2] is not None else '',
            fila[3],
            separador_miles(fila[4]) if fila[4] is not None else '',
            fila[5],
            fila[6],
        ])
    if hay_mas:
        datos.append(['-', 'Otros', '-', '-', '-', '-', '-'])

    tabla = Table(data=datos)
    tabla.setStyle(TableStyle([
        ('TEXTCOLOR', (0, 1), (1, -1), colors.black),
        ('FONTSIZE', (0, 0), (-1, -1), 4),
        ('TEXTCOLOR', (1, 2), (3, -1), colors.black),
        ('BACKGROUND', (1, 1), (-1, -1), colors.white),
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('BOX', (0, 0), (-1, -1), 1.25, colors.grey),
        ('INNERGRID', (0, 0), (-1, -1), 1, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
    ]))
    return tabla


def generar_pdf_liquidacion(liquidacion, completa=False):
    """Genera el PDF de la liquidación, con el mismo formato que el sistema
    de escritorio. Si completa=True, agrega debajo de cada tabla (Debe y
    Haber) el detalle de renglones de los comprobantes de ese lado (hasta 5
    renglones en total, con un "Otros" si hay más)."""
    response = HttpResponse(content_type='application/pdf')
    sufijo = 'completa' if completa else 'simple'
    response['Content-Disposition'] = f'attachment; filename="liquidacion_{liquidacion.id}_{sufijo}.pdf"'

    doc = SimpleDocTemplate(response, pagesize=A4, topMargin=0.1)

    estilo_hoja = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        'Cabecera', fontName='Helvetica-Bold', fontSize=10, parent=estilo_hoja['Heading1'],
        alignment=TA_CENTER, spaceAfter=10,
    )
    estilo_no_valido = ParagraphStyle(
        'NoValidoComoFactura', fontName='Helvetica', fontSize=10, parent=estilo_hoja['Heading1'],
        alignment=TA_CENTER, spaceAfter=10,
    )
    estilo_cabecera = ParagraphStyle(
        'Comprobante', fontName='Helvetica-Bold', fontSize=10, parent=estilo_hoja['Heading2'],
        alignment=TA_CENTER, spaceAfter=10,
    )
    estilo_pie = ParagraphStyle(
        'Pie', fontName='Helvetica-Bold', fontSize=10, parent=estilo_hoja['Normal'],
        alignment=TA_RIGHT, spaceAfter=10,
    )
    estilo_pie_recibo = ParagraphStyle(
        'PieRecibo', fontName='Helvetica-Bold', fontSize=8, parent=estilo_hoja['Normal'],
        alignment=TA_LEFT, spaceAfter=10,
    )

    filas_debe, total_debe = _filas_debe(liquidacion)
    filas_haber, total_haber = _filas_haber(liquidacion)

    tabla_debe = _tabla_pdf(DEBE_HEADERS, filas_debe)
    tabla_haber = _tabla_pdf(HABER_HEADERS, filas_haber)

    entidad = liquidacion.entidad

    story = []
    if os.path.exists(LOGO_PATH):
        story.append(Image(LOGO_PATH, width=192, height=92))
    story.append(Spacer(0, 5))
    story.append(Paragraph('DOCUMENTO NO VÁLIDO COMO FACTURA', estilo_no_valido))
    story.append(Spacer(0, 5))

    fecha_texto = liquidacion.fecha.strftime('%d/%m/%Y') if liquidacion.fecha else ''
    story.append(Paragraph(f'Liquidación N°: {liquidacion.id}      ({fecha_texto})', estilo_titulo))

    if entidad:
        cuit_texto = _cuit_con_guiones(entidad.cuit) if entidad.cuit else ''
        story.append(Paragraph(
            f'Proveedor:  ({entidad.id}) Cuit: {cuit_texto} {entidad.nombre}', estilo_titulo,
        ))

    story.append(Paragraph('Comprobantes a Pagar', estilo_cabecera))
    story.append(tabla_debe)
    story.append(Spacer(0, 5))
    story.append(Paragraph(f'Total a Pagar:   {separador_miles(total_debe)}', estilo_pie))
    story.append(Spacer(0, 10))

    if completa:
        comprobante_ids_debe = list(
            liquidacion.comprobantes.filter(tipo='debe').values_list('comprobante_id', flat=True)
        )
        filas_renglones_debe, hay_mas_debe = _renglones_comprobantes(comprobante_ids_debe, limitar=MAX_RENGLONES_PDF)
        if filas_renglones_debe or hay_mas_debe:
            story.append(_tabla_renglones_pdf(filas_renglones_debe, hay_mas_debe))
            story.append(Spacer(0, 10))

    story.append(Paragraph('Pago', estilo_cabecera))
    story.append(tabla_haber)
    story.append(Spacer(0, 5))
    story.append(Paragraph(f'Total de Pago:  {separador_miles(total_haber)}', estilo_pie))
    story.append(Spacer(0, 10))

    if completa:
        comprobante_ids_haber = list(
            liquidacion.comprobantes.filter(tipo='haber').values_list('comprobante_id', flat=True)
        )
        filas_renglones_haber, hay_mas_haber = _renglones_comprobantes(comprobante_ids_haber, limitar=MAX_RENGLONES_PDF)
        if filas_renglones_haber or hay_mas_haber:
            story.append(_tabla_renglones_pdf(filas_renglones_haber, hay_mas_haber))
            story.append(Spacer(0, 10))

    numero_en_letras = numero_a_moneda(total_haber)
    story.append(Paragraph(f'Recibí conforme en Pesos:  {numero_en_letras}', estilo_pie_recibo))
    story.append(Spacer(0, 10))
    story.append(Paragraph('Firma:  ', estilo_pie_recibo))
    story.append(Spacer(0, 5))
    story.append(Paragraph('Aclaración:  ', estilo_pie_recibo))
    story.append(Spacer(0, 5))
    story.append(Paragraph('DNI:  ', estilo_pie_recibo))
    story.append(Spacer(0, 10))

    doc.build(story)
    return response


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _numero_o_none(valor):
    if valor is None:
        return None
    return float(valor)


def _formatear_columna_numerica(ws, columna_1_based):
    from openpyxl.utils import get_column_letter
    letra = get_column_letter(columna_1_based)
    for celda in ws[letra][1:]:
        if celda.value is not None:
            celda.number_format = FORMATO_MILES_EXCEL


def _autoajustar_columnas(ws):
    for columna in ws.columns:
        letra = columna[0].column_letter
        largo_max = max((len(str(c.value)) for c in columna if c.value is not None), default=10)
        ws.column_dimensions[letra].width = min(max(largo_max + 2, 10), 40)


def generar_excel_liquidacion(liquidacion, completa=False):
    """Genera un Excel con hojas Resumen / Debe / Haber (+ Renglones si es
    completo). No es una réplica de un formato existente (el sistema de
    escritorio no generaba Excel para una liquidación individual): se armó
    con el mismo criterio de formato numérico/estilo que el resto de los
    reportes de esta aplicación."""
    import openpyxl
    from services.gestorexcel import definir_estilo_general

    wb = openpyxl.Workbook()

    entidad = liquidacion.entidad
    filas_debe, total_debe = _filas_debe(liquidacion)
    filas_haber, total_haber = _filas_haber(liquidacion)

    ws_resumen = wb.active
    ws_resumen.title = 'Resumen'
    ws_resumen.append(['Liquidación N°', liquidacion.numero or liquidacion.id])
    ws_resumen.append(['Fecha', liquidacion.fecha])
    ws_resumen.append(['Proveedor', f'({entidad.id}) {entidad.nombre}' if entidad else ''])
    ws_resumen.append(['Cuit', _cuit_con_guiones(entidad.cuit) if entidad and entidad.cuit else ''])
    ws_resumen.append(['Total a Pagar (Debe)', _numero_o_none(total_debe)])
    ws_resumen.append(['Total de Pago (Haber)', _numero_o_none(total_haber)])
    ws_resumen.append(['Diferencia', _numero_o_none(total_debe - total_haber)])
    for fila_idx in (5, 6, 7):
        ws_resumen.cell(row=fila_idx, column=2).number_format = FORMATO_MILES_EXCEL
    definir_estilo_general(ws_resumen)
    _autoajustar_columnas(ws_resumen)

    ws_debe = wb.create_sheet('Debe')
    ws_debe.append(DEBE_HEADERS)
    for fila in filas_debe:
        ws_debe.append([fila[0], fila[1], fila[2], fila[3], _numero_o_none(fila[4])])
    _formatear_columna_numerica(ws_debe, len(DEBE_HEADERS))
    definir_estilo_general(ws_debe)
    _autoajustar_columnas(ws_debe)

    ws_haber = wb.create_sheet('Haber')
    ws_haber.append(HABER_HEADERS)
    for fila in filas_haber:
        ws_haber.append([fila[0], fila[1], fila[2], fila[3], fila[4], _numero_o_none(fila[5])])
    _formatear_columna_numerica(ws_haber, len(HABER_HEADERS))
    definir_estilo_general(ws_haber)
    _autoajustar_columnas(ws_haber)

    if completa:
        comprobante_ids_debe = list(
            liquidacion.comprobantes.filter(tipo='debe').values_list('comprobante_id', flat=True)
        )
        comprobante_ids_haber = list(
            liquidacion.comprobantes.filter(tipo='haber').values_list('comprobante_id', flat=True)
        )
        renglones_debe, _ = _renglones_comprobantes(comprobante_ids_debe)
        renglones_haber, _ = _renglones_comprobantes(comprobante_ids_haber)

        ws_renglones = wb.create_sheet('Renglones')
        ws_renglones.append(['Lado'] + RENGLONES_HEADERS)
        for fila in renglones_debe:
            ws_renglones.append(['Debe', fila[0], fila[1], _numero_o_none(fila[2]), fila[3], _numero_o_none(fila[4]), fila[5], fila[6]])
        for fila in renglones_haber:
            ws_renglones.append(['Haber', fila[0], fila[1], _numero_o_none(fila[2]), fila[3], _numero_o_none(fila[4]), fila[5], fila[6]])
        _formatear_columna_numerica(ws_renglones, 4)  # Cantidad
        _formatear_columna_numerica(ws_renglones, 6)  # Precio Unitario
        definir_estilo_general(ws_renglones)
        _autoajustar_columnas(ws_renglones)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    sufijo = 'completa' if completa else 'simple'
    response['Content-Disposition'] = f'attachment; filename="liquidacion_{liquidacion.id}_{sufijo}.xlsx"'
    wb.save(response)
    return response
