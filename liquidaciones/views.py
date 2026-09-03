from decimal import Decimal

from django.contrib import messages
from django.core.exceptions import ObjectDoesNotExist
from django.db import transaction
from django.db.models import Case, DecimalField, F, Max, OuterRef, Q, Subquery, Sum, Value, When
from django.db.models.functions import Coalesce, Cast
from django.http import JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse

from comprobantes.models import Comprobante
from entidades.models import Entidad
from services.ordenamiento import aplicar_orden_lista, aplicar_orden_queryset
from movimientos_caja.models import MovimientoCaja
from retenciones.models import Retencion
from retenciones_inym.models import RetencionInym

from . import documentos
from .forms import LiquidacionSeleccionForm, LiquidacionReporteForm, SinLiquidarFiltroForm
from .models import (
    Liquidacion,
    LiquidacionComprobante,
    LiquidacionRetencion,
    LiquidacionRetencionInym,
    LiquidacionMovimiento,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _monto_item(item_obj):
    """
    MovimientoCaja usa 'monto'; Comprobante/Retencion/RetencionInym usan
    'total'. Si el objeto es un Comprobante en moneda distinta a pesos
    (tiene registro en comprobante_tipo_de_cambio), el total se multiplica
    por ese tipo de cambio.
    """
    if hasattr(item_obj, 'monto'):
        return item_obj.monto or Decimal('0')
    total = getattr(item_obj, 'total', None) or Decimal('0')
    if isinstance(item_obj, Comprobante):
        try:
            total = (total * item_obj.tipo_de_cambio.tipo_de_cambio).quantize(Decimal('0.01'))
        except ObjectDoesNotExist:
            pass
    return total


def _armar_items(entidad, liquidacion_actual=None):
    """
    Devuelve, para la entidad dada, los movimientos/comprobantes/retenciones/
    retenciones_inym que no tienen liquidación asociada. Si se pasa
    liquidacion_actual (modo edición), también incluye los ítems que ya
    están vinculados a ESA liquidación, marcados como seleccionados.

    Además, en modo edición, devuelve por separado ('movimientos_otros' /
    'comprobantes_otros') los movimientos_caja / comprobantes que ya están
    vinculados a esa liquidación pero que pertenecen a OTRA entidad (se
    agregaron a mano desde el buscador de "otros movimientos / comprobantes",
    por ejemplo para aplicar un cheque_recibido de un tercero) — si no se
    devolvieran acá, desaparecerían de la pantalla al editar la liquidación.
    """

    def excluidos(modelo_intermedio, campo_fk):
        qs = modelo_intermedio.objects.all()
        if liquidacion_actual:
            qs = qs.exclude(liquidacion=liquidacion_actual)
        return qs.values_list(campo_fk, flat=True)

    def tipo_actual(modelo_intermedio, campo_fk):
        if not liquidacion_actual:
            return {}
        return dict(
            modelo_intermedio.objects.filter(liquidacion=liquidacion_actual)
            .values_list(campo_fk, 'tipo')
        )

    # --- Movimientos de caja ---
    mov_excl = excluidos(LiquidacionMovimiento, 'movimiento_caja_id')
    mov_tipo = tipo_actual(LiquidacionMovimiento, 'movimiento_caja_id')
    movimientos = list(
        MovimientoCaja.objects.filter(receptor=entidad)
        .exclude(id__in=mov_excl)
        .select_related('tipo', 'rel_numero')
        .order_by('-emision')
    )
    for m in movimientos:
        m.monto_mostrar = _monto_item(m)
        m.tipo_actual = mov_tipo.get(m.id)
        m.seleccionado = m.id in mov_tipo

    # --- Comprobantes ---
    comp_excl = excluidos(LiquidacionComprobante, 'comprobante_id')
    comp_tipo = tipo_actual(LiquidacionComprobante, 'comprobante_id')
    comprobantes = list(
        Comprobante.objects.filter(entidad_emisor=entidad)
        .exclude(id__in=comp_excl)
        .select_related('tipo_de_cambio', 'tipo_comprobante')
        .order_by('-fecha')
    )
    for c in comprobantes:
        c.monto_mostrar = _monto_item(c)
        c.tipo_actual = comp_tipo.get(c.id)
        c.seleccionado = c.id in comp_tipo

    # --- Retenciones ---
    ret_excl = excluidos(LiquidacionRetencion, 'retencion_id')
    ret_tipo = tipo_actual(LiquidacionRetencion, 'retencion_id')
    retenciones = list(
        Retencion.objects.filter(entidad=entidad)
        .exclude(id__in=ret_excl)
        .select_related('id_regimen', 'id_impuesto')
        .order_by('-fecha')
    )
    for r in retenciones:
        r.monto_mostrar = _monto_item(r)
        r.tipo_actual = ret_tipo.get(r.id)
        r.seleccionado = r.id in ret_tipo

    # --- Retenciones INYM (la entidad se llega vía operador_retenido) ---
    retinym_excl = excluidos(LiquidacionRetencionInym, 'retencion_inym_id')
    retinym_tipo = tipo_actual(LiquidacionRetencionInym, 'retencion_inym_id')
    retenciones_inym = list(
        RetencionInym.objects.filter(operador_retenido__entidad=entidad)
        .exclude(id__in=retinym_excl)
        .order_by('-fecha')
    )
    for ri in retenciones_inym:
        ri.monto_mostrar = _monto_item(ri)
        ri.tipo_actual = retinym_tipo.get(ri.id)
        ri.seleccionado = ri.id in retinym_tipo

    # --- Movimientos / comprobantes de OTRA entidad, ya vinculados a esta liquidación ---
    movimientos_otros = []
    comprobantes_otros = []
    if liquidacion_actual:
        mov_otros_tipo = dict(
            LiquidacionMovimiento.objects.filter(liquidacion=liquidacion_actual)
            .exclude(movimiento_caja__receptor=entidad)
            .values_list('movimiento_caja_id', 'tipo')
        )
        if mov_otros_tipo:
            movimientos_otros = list(
                MovimientoCaja.objects.filter(id__in=mov_otros_tipo.keys())
                .select_related('tipo', 'rel_numero', 'receptor')
                .order_by('-emision')
            )
            for m in movimientos_otros:
                m.monto_mostrar = _monto_item(m)
                m.tipo_actual = mov_otros_tipo.get(m.id)
                m.seleccionado = True

        comp_otros_tipo = dict(
            LiquidacionComprobante.objects.filter(liquidacion=liquidacion_actual)
            .exclude(comprobante__entidad_emisor=entidad)
            .values_list('comprobante_id', 'tipo')
        )
        if comp_otros_tipo:
            comprobantes_otros = list(
                Comprobante.objects.filter(id__in=comp_otros_tipo.keys())
                .select_related('tipo_de_cambio', 'tipo_comprobante', 'entidad_emisor')
                .order_by('-fecha')
            )
            for c in comprobantes_otros:
                c.monto_mostrar = _monto_item(c)
                c.tipo_actual = comp_otros_tipo.get(c.id)
                c.seleccionado = True

    return {
        'movimientos': movimientos,
        'comprobantes': comprobantes,
        'retenciones': retenciones,
        'retenciones_inym': retenciones_inym,
        'movimientos_otros': movimientos_otros,
        'comprobantes_otros': comprobantes_otros,
    }


def _siguiente_id_liquidacion():
    ultimo = Liquidacion.objects.aggregate(Max('id'))['id__max'] or 0
    return ultimo + 1


# ---------------------------------------------------------------------------
# Búsqueda de entidad (usada por el autocompletado en el form de alta/edición)
# ---------------------------------------------------------------------------

def entidad_buscar(request):
    """
    Devuelve, en JSON, hasta 20 entidades cuyo ID coincida exactamente (si lo
    buscado es numérico) o cuyo nombre o CUIT contengan el texto buscado. Se
    usa desde el JS de las pantallas de liquidaciones (alta/edición, reporte
    y los 4 listados "sin liquidar") para armar un buscador con
    autocompletado en vez de un <select> con todas las entidades.
    """
    q = request.GET.get('q', '').strip()
    resultados = []

    if len(q) >= 2:
        filtro = Q(nombre__icontains=q) | Q(cuit__icontains=q)
        if q.isdigit():
            filtro |= Q(id=int(q))
        entidades = Entidad.objects.filter(filtro).order_by('nombre')[:20]
        resultados = [
            {
                'id': ent.id,
                'text': f'{ent.nombre} (CUIT {ent.cuit})' if ent.cuit else ent.nombre,
            }
            for ent in entidades
        ]

    return JsonResponse({'resultados': resultados})


def _texto_entidad(entidad):
    """Texto a mostrar/precargar en el input visible del buscador de entidad
    (nombre + CUIT si tiene), usado en reporte.html y en los 4 listados "sin
    liquidar" para que el filtro ya aplicado se vea, no solo su id oculto."""
    if not entidad:
        return ''
    return f'{entidad.nombre} (CUIT {entidad.cuit})' if entidad.cuit else entidad.nombre


# ---------------------------------------------------------------------------
# Búsqueda de movimientos_caja / comprobantes de CUALQUIER entidad, sin
# liquidar, para el apartado "Otros movimientos / comprobantes" del form de
# alta/edición de liquidación. Permite, por ejemplo, agregar a la liquidación
# de un proveedor un cheque_recibido que pertenece a otra entidad y todavía
# no se usó en ninguna liquidación.
# ---------------------------------------------------------------------------

def item_sin_liquidar_buscar(request):
    q = request.GET.get('q', '').strip()
    entidad_excluir_id = request.GET.get('entidad_excluir', '').strip()
    resultados = []

    if len(q) < 2:
        return JsonResponse({'resultados': resultados})

    es_numero = q.isdigit()

    # --- Movimientos de caja ---
    movimientos = MovimientoCaja.objects.filter(liquidaciones__isnull=True)
    if entidad_excluir_id.isdigit():
        movimientos = movimientos.exclude(receptor_id=int(entidad_excluir_id))

    filtro_mov = (
        Q(receptor__nombre__icontains=q)
        | Q(receptor__cuit__icontains=q)
        | Q(emisor_relacion__id_entidad__nombre__icontains=q)
        | Q(rel_concepto__concepto_tipo__nombre__icontains=q)
    )
    if es_numero:
        filtro_mov |= Q(id=int(q)) | Q(rel_numero__numero=int(q))

    movimientos = (
        movimientos.filter(filtro_mov)
        .select_related('receptor', 'tipo', 'rel_numero')
        .order_by('-emision')[:20]
    )
    for m in movimientos:
        resultados.append({
            'tipo': 'mov',
            'id': m.id,
            'entidad': str(m.receptor) if m.receptor else 'Sin entidad',
            'fecha': m.emision.strftime('%d/%m/%Y') if m.emision else '',
            'detalle': ' - '.join(filter(None, [
                m.tipo.nombre if m.tipo_id else None,
                f'Nº {m.numero}' if m.numero else None,
            ])),
            'monto': float(_monto_item(m)),
        })

    # --- Comprobantes ---
    comprobantes = Comprobante.objects.filter(liquidaciones__isnull=True)
    if entidad_excluir_id.isdigit():
        comprobantes = comprobantes.exclude(entidad_emisor_id=int(entidad_excluir_id))

    filtro_comp = (
        Q(entidad_emisor__nombre__icontains=q)
        | Q(entidad_emisor__cuit__icontains=q)
        | Q(entidad_nombre__icontains=q)
        | Q(tipo_comprobante__nombre__icontains=q)
        | Q(comprobante_string__icontains=q)
    )
    if es_numero:
        filtro_comp |= Q(id=int(q)) | Q(numero=int(q))

    comprobantes = (
        comprobantes.filter(filtro_comp)
        .select_related('entidad_emisor', 'tipo_comprobante', 'tipo_de_cambio')
        .order_by('-fecha')[:20]
    )
    for c in comprobantes:
        resultados.append({
            'tipo': 'comp',
            'id': c.id,
            'entidad': str(c.entidad_emisor) if c.entidad_emisor else (c.entidad_nombre or 'Sin entidad'),
            'fecha': c.fecha.strftime('%d/%m/%Y') if c.fecha else '',
            'detalle': ' - '.join(filter(None, [
                c.tipo_comprobante.nombre if c.tipo_comprobante_id else None,
                c.comprobante_string,
            ])),
            'monto': float(_monto_item(c)),
        })

    return JsonResponse({'resultados': resultados[:30]})


# ---------------------------------------------------------------------------
# Listado
# ---------------------------------------------------------------------------

def liquidacion_list(request):
    liquidaciones = (
        Liquidacion.objects.select_related('entidad')
        .order_by('-fecha', '-id')
    )

    q_entidad = request.GET.get('entidad', '').strip()
    q_id = request.GET.get('id', '').strip()
    q_fecha = request.GET.get('fecha', '').strip()

    if q_entidad:
        liquidaciones = liquidaciones.filter(
            Q(entidad__nombre__icontains=q_entidad) | Q(entidad__cuit__icontains=q_entidad)
        )
    if q_id:
        if q_id.isdigit():
            liquidaciones = liquidaciones.filter(id=int(q_id))
        else:
            liquidaciones = liquidaciones.none()
    if q_fecha:
        liquidaciones = liquidaciones.filter(fecha=q_fecha)

    liquidaciones = aplicar_orden_queryset(request, liquidaciones, {
        'id': 'id',
        'numero': 'numero',
        'fecha': 'fecha',
        'entidad': 'entidad__nombre',
        'debe': 'debe',
        'haber': 'haber',
        'diferencia': F('debe') - F('haber'),
    })

    return render(request, 'liquidaciones/list.html', {
        'liquidaciones': liquidaciones,
        'q_entidad': q_entidad,
        'q_id': q_id,
        'q_fecha': q_fecha,
    })

# ---------------------------------------------------------------------------
# Alta / Edición (misma vista, pk=None para alta)
# ---------------------------------------------------------------------------

CONFIG_ITEMS = {
    'mov': ('mov_selected', 'mov_tipo', MovimientoCaja, LiquidacionMovimiento, 'movimiento_caja'),
    'comp': ('comp_selected', 'comp_tipo', Comprobante, LiquidacionComprobante, 'comprobante'),
    'ret': ('ret_selected', 'ret_tipo', Retencion, LiquidacionRetencion, 'retencion'),
    'retinym': ('retinym_selected', 'retinym_tipo', RetencionInym, LiquidacionRetencionInym, 'retencion_inym'),
}


@transaction.atomic
def liquidacion_form(request, pk=None):
    liquidacion = get_object_or_404(Liquidacion, pk=pk) if pk else None

    entidad = None
    fecha = None
    items = None

    if request.method == 'POST':
        fecha = request.POST.get('fecha')
        entidad_id = request.POST.get('entidad')
        entidad = get_object_or_404(Entidad, pk=entidad_id) if entidad_id else None

        if not fecha or not entidad:
            messages.error(request, 'Debe indicar fecha y entidad.')
        else:
            a_crear = []  # (modelo_intermedio, fk_name, item_id, tipo)

            for campo_sel, campo_tipo, modelo_item, modelo_intermedio, fk_name in CONFIG_ITEMS.values():
                for item_id in request.POST.getlist(campo_sel):
                    tipo = request.POST.get(f'{campo_tipo}_{item_id}')
                    if tipo not in ('debe', 'haber'):
                        continue
                    a_crear.append((modelo_intermedio, fk_name, item_id, tipo))

            if not a_crear:
                messages.error(request, 'Debe seleccionar al menos un ítem para la liquidación.')
            else:
                if liquidacion is None:
                    liquidacion = Liquidacion(id=_siguiente_id_liquidacion())
                else:
                    LiquidacionMovimiento.objects.filter(liquidacion=liquidacion).delete()
                    LiquidacionComprobante.objects.filter(liquidacion=liquidacion).delete()
                    LiquidacionRetencion.objects.filter(liquidacion=liquidacion).delete()
                    LiquidacionRetencionInym.objects.filter(liquidacion=liquidacion).delete()

                liquidacion.fecha = fecha
                liquidacion.entidad = entidad
                if not liquidacion.numero:
                    liquidacion.numero = f'LIQ-{liquidacion.id}'
                liquidacion.save()

                for modelo_intermedio, fk_name, item_id, tipo in a_crear:
                    modelo_intermedio.objects.create(
                        liquidacion=liquidacion,
                        tipo=tipo,
                        **{f'{fk_name}_id': int(item_id)},
                    )

                # Única fuente de verdad para debe/haber: se recalcula desde la DB
                liquidacion.recalcular_totales()

                messages.success(request, f'Liquidación {liquidacion.numero} guardada correctamente.')
                return redirect(f"{reverse('liquidaciones:listado')}?id={liquidacion.id}")

    else:
        if liquidacion:
            fecha = liquidacion.fecha
            entidad = liquidacion.entidad

        entidad_id = request.GET.get('entidad')
        if entidad_id:
            entidad = get_object_or_404(Entidad, pk=entidad_id)

        fecha_get = request.GET.get('fecha')
        if fecha_get:
            fecha = fecha_get

    if entidad:
        items = _armar_items(entidad, liquidacion)

    entidad_texto = ''
    if entidad:
        entidad_texto = f'{entidad.nombre} (CUIT {entidad.cuit})' if entidad.cuit else entidad.nombre

    form = LiquidacionSeleccionForm(initial={'fecha': fecha, 'entidad': entidad})

    return render(request, 'liquidaciones/form.html', {
        'form': form,
        'liquidacion': liquidacion,
        'entidad': entidad,
        'entidad_texto': entidad_texto,
        'fecha': fecha,
        'items': items,
    })


# ---------------------------------------------------------------------------
# Eliminar
# ---------------------------------------------------------------------------

def liquidacion_eliminar(request, pk):
    liquidacion = get_object_or_404(Liquidacion, pk=pk)

    if request.method == 'POST':
        LiquidacionMovimiento.objects.filter(liquidacion=liquidacion).delete()
        LiquidacionComprobante.objects.filter(liquidacion=liquidacion).delete()
        LiquidacionRetencion.objects.filter(liquidacion=liquidacion).delete()
        LiquidacionRetencionInym.objects.filter(liquidacion=liquidacion).delete()
        numero = liquidacion.numero
        liquidacion.delete()
        messages.success(request, f'Liquidación {numero} eliminada.')
        return redirect('liquidaciones:listado')

    return render(request, 'liquidaciones/eliminar_confirm.html', {'liquidacion': liquidacion})


# ---------------------------------------------------------------------------
# Impresión de una liquidación (PDF / Excel, simple o completa)
# ---------------------------------------------------------------------------

def liquidacion_pdf(request, pk):
    liquidacion = get_object_or_404(Liquidacion.objects.select_related('entidad'), pk=pk)
    completa = request.GET.get('completa') == '1'
    return documentos.generar_pdf_liquidacion(liquidacion, completa=completa)


def liquidacion_excel(request, pk):
    liquidacion = get_object_or_404(Liquidacion.objects.select_related('entidad'), pk=pk)
    completa = request.GET.get('completa') == '1'
    return documentos.generar_excel_liquidacion(liquidacion, completa=completa)


# ---------------------------------------------------------------------------
# Reportes
# ---------------------------------------------------------------------------

def liquidacion_reporte(request):
    form = LiquidacionReporteForm(request.GET or None)
    liquidaciones = Liquidacion.objects.select_related('entidad').order_by('-fecha')
    entidad_texto = ''

    if form.is_valid():
        entidad = form.cleaned_data.get('entidad')
        fecha_desde = form.cleaned_data.get('fecha_desde')
        fecha_hasta = form.cleaned_data.get('fecha_hasta')
        if entidad:
            liquidaciones = liquidaciones.filter(entidad=entidad)
            entidad_texto = _texto_entidad(entidad)
        if fecha_desde:
            liquidaciones = liquidaciones.filter(fecha__gte=fecha_desde)
        if fecha_hasta:
            liquidaciones = liquidaciones.filter(fecha__lte=fecha_hasta)

    totales = liquidaciones.aggregate(
        total_debe=Coalesce(Sum('debe'), Decimal('0'), output_field=DecimalField()),
        total_haber=Coalesce(Sum('haber'), Decimal('0'), output_field=DecimalField()),
    )

    liquidaciones = aplicar_orden_queryset(request, liquidaciones, {
        'numero': 'numero',
        'fecha': 'fecha',
        'entidad': 'entidad__nombre',
        'debe': 'debe',
        'haber': 'haber',
    })

    return render(request, 'liquidaciones/reporte.html', {
        'form': form,
        'liquidaciones': liquidaciones,
        'totales': totales,
        'entidad_texto': entidad_texto,
    })


# ---------------------------------------------------------------------------
# Diferencias (debe/haber guardado vs. calculado desde las relaciones)
# ---------------------------------------------------------------------------

def liquidacion_diferencias(request):
    """
    Compara, para cada liquidación, el debe/haber guardado contra lo que
    da sumar los ítems relacionados. Todo se resuelve en UNA sola consulta
    SQL: cada suma se arma como una subquery correlacionada (Subquery +
    OuterRef), y el filtro de "solo las que difieren" se aplica en la
    base de datos, no en Python.

    Los comprobantes que son Notas de Crédito se RESTAN en vez de sumarse
    (ver suma_comprobantes): el campo 'total' de Comprobante siempre está
    guardado en positivo en la base, el signo se aplica acá según el
    nombre del tipo de comprobante.
    """

    def suma_relacionada(modelo_intermedio, campo_monto, tipo):
        return Coalesce(
            Subquery(
                modelo_intermedio.objects
                .filter(liquidacion=OuterRef('pk'), tipo=tipo)
                .order_by()
                .values('liquidacion')
                .annotate(total=Sum(campo_monto))
                .values('total'),
                output_field=DecimalField(max_digits=20, decimal_places=2),
            ),
            Value(Decimal('0')),
        )

    def suma_comprobantes(tipo):
        """
        Igual que suma_relacionada, pero además:
        - invierte el signo del total cuando el comprobante es una Nota de
          Crédito: el campo 'total' de Comprobante siempre se guarda en
          positivo en la base, así que acá se le aplica el signo negativo
          cuando el nombre del tipo de comprobante contiene "nota de
          credito".
        - convierte el total multiplicándolo por el tipo de cambio, si el
          comprobante tiene un registro en comprobante_tipo_de_cambio
          (moneda distinta a pesos). Si no tiene, el factor es 1.
        """
        factor_cambio = Coalesce(
            F('comprobante__tipo_de_cambio__tipo_de_cambio'),
            Value(Decimal('1')),
            output_field=DecimalField(max_digits=10, decimal_places=2),
        )
        return Coalesce(
            Subquery(
                LiquidacionComprobante.objects
                .filter(liquidacion=OuterRef('pk'), tipo=tipo)
                .annotate(
                    # Cast(): multiplicar dos DecimalField en SQL da más
                    # decimales de los que declara output_field (Django no
                    # lo redondea solo) — sin este Cast, monto_signado queda
                    # con 4+ decimales y ya no coincide con lo guardado.
                    monto_signado=Cast(
                        Case(
                            When(
                                comprobante__tipo_comprobante__nombre__icontains='nota de credito',
                                then=-F('comprobante__total') * factor_cambio,
                            ),
                            default=F('comprobante__total') * factor_cambio,
                            output_field=DecimalField(max_digits=20, decimal_places=4),
                        ),
                        output_field=DecimalField(max_digits=20, decimal_places=2),
                    )
                )
                .order_by()
                .values('liquidacion')
                .annotate(total=Sum('monto_signado'))
                .values('total'),
                output_field=DecimalField(max_digits=20, decimal_places=2),
            ),
            Value(Decimal('0')),
        )

    liquidaciones = (
        Liquidacion.objects
        .select_related('entidad')
        .annotate(
            mov_debe=suma_relacionada(LiquidacionMovimiento, 'movimiento_caja__monto', 'debe'),
            mov_haber=suma_relacionada(LiquidacionMovimiento, 'movimiento_caja__monto', 'haber'),
            comp_debe=suma_comprobantes('debe'),
            comp_haber=suma_comprobantes('haber'),
            ret_debe=suma_relacionada(LiquidacionRetencion, 'retencion__total', 'debe'),
            ret_haber=suma_relacionada(LiquidacionRetencion, 'retencion__total', 'haber'),
            retinym_debe=suma_relacionada(LiquidacionRetencionInym, 'retencion_inym__total', 'debe'),
            retinym_haber=suma_relacionada(LiquidacionRetencionInym, 'retencion_inym__total', 'haber'),
        )
        .annotate(
            debe_calculado=F('mov_debe') + F('comp_debe') + F('ret_debe') + F('retinym_debe'),
            haber_calculado=F('mov_haber') + F('comp_haber') + F('ret_haber') + F('retinym_haber'),
            debe_guardado=Coalesce('debe', Value(Decimal('0')), output_field=DecimalField(max_digits=20, decimal_places=2)),
            haber_guardado=Coalesce('haber', Value(Decimal('0')), output_field=DecimalField(max_digits=20, decimal_places=2)),
        )
        .exclude(debe_guardado=F('debe_calculado'), haber_guardado=F('haber_calculado'))
        .order_by('-fecha', '-id')
    )

    filas = [
        {
            'liquidacion': liq,
            'debe_guardado': liq.debe_guardado,
            'haber_guardado': liq.haber_guardado,
            'debe_calculado': liq.debe_calculado,
            'haber_calculado': liq.haber_calculado,
            'diferencia_debe': liq.debe_guardado - liq.debe_calculado,
            'diferencia_haber': liq.haber_guardado - liq.haber_calculado,
        }
        for liq in liquidaciones
    ]

    filas = aplicar_orden_lista(request, filas, {
        'id': lambda f: f['liquidacion'].id,
        'numero': lambda f: f['liquidacion'].numero or 0,
        'fecha': lambda f: f['liquidacion'].fecha,
        'entidad': lambda f: str(f['liquidacion'].entidad or '').lower(),
        'debe_guardado': lambda f: f['debe_guardado'],
        'debe_calculado': lambda f: f['debe_calculado'],
        'diferencia_debe': lambda f: f['diferencia_debe'],
        'haber_guardado': lambda f: f['haber_guardado'],
        'haber_calculado': lambda f: f['haber_calculado'],
        'diferencia_haber': lambda f: f['diferencia_haber'],
    })

    return render(request, 'liquidaciones/diferencias.html', {'filas': filas})


def liquidacion_recalcular(request, pk):
    """Recalcula y guarda debe/haber de UNA liquidación puntual, desde la vista de diferencias."""
    liquidacion = get_object_or_404(Liquidacion, pk=pk)
    if request.method == 'POST':
        liquidacion.recalcular_totales()
        messages.success(request, f'Liquidación {liquidacion.numero} recalculada y actualizada.')
    return redirect('liquidaciones:diferencias')


# ---------------------------------------------------------------------------
# "Sin liquidar": listados filtrables (Entidad / fecha desde-hasta) de los
# comprobantes, movimientos de caja, retenciones y retenciones INYM que
# todavía NO están vinculados a ninguna liquidación (mismo criterio que ya
# usa SinLiquidacionFilterBase en los admin: '<related_name>__isnull=True').
# Desde cada fila se accede al botón LIQUIDAR, que lleva directo al alta de
# liquidación con la entidad de esa fila ya seleccionada y sus pendientes ya
# cargados (liquidacion_form ya soporta '?entidad=<id>' por GET).
# ---------------------------------------------------------------------------

# Entidad que representa a la propia empresa (Fontana). Se usa, en
# Comprobantes, como contraparte "implícita" del lado que no tiene entidad
# propia (ver es_emisor más abajo) — mismo criterio que ya usa la property
# MovimientoCaja.emisor para su entidad por defecto.
ENTIDAD_PROPIA_ID = 100


def _entidad_propia():
    return Entidad.objects.filter(id=ENTIDAD_PROPIA_ID).first()


# --- Comprobantes -----------------------------------------------------------

def _comprobantes_sin_liquidar(request):
    form = SinLiquidarFiltroForm(request.GET or None)
    comprobantes = (
        Comprobante.objects.filter(liquidaciones__isnull=True)
        .select_related('entidad_emisor', 'tipo_comprobante')
        .order_by('-fecha', '-id')
    )
    if form.is_valid():
        entidad = form.cleaned_data.get('entidad')
        fecha_desde = form.cleaned_data.get('fecha_desde')
        fecha_hasta = form.cleaned_data.get('fecha_hasta')
        if entidad:
            comprobantes = comprobantes.filter(entidad_emisor=entidad)
        if fecha_desde:
            comprobantes = comprobantes.filter(fecha__gte=fecha_desde)
        if fecha_hasta:
            comprobantes = comprobantes.filter(fecha__lte=fecha_hasta)
    return form, comprobantes


def sin_liquidar_comprobantes(request):
    form, comprobantes = _comprobantes_sin_liquidar(request)
    comprobantes = list(comprobantes[:500])

    entidad_propia = _entidad_propia()
    for c in comprobantes:
        # es_emisor = 0 -> la entidad del comprobante es la RECEPTORA (nosotros emitimos).
        # es_emisor = 1 o vacío (valor por defecto) -> la entidad es la EMISORA (nosotros recibimos).
        if c.es_emisor == 0:
            c.entidad_emisor_mostrar = entidad_propia
            c.entidad_receptor_mostrar = c.entidad_emisor
        else:
            c.entidad_emisor_mostrar = c.entidad_emisor
            c.entidad_receptor_mostrar = entidad_propia
        # La entidad contra la que se liquida siempre es la contraparte real
        # del comprobante (nunca la entidad propia).
        c.entidad_liquidar_id = c.entidad_emisor_id

    return render(request, 'liquidaciones/sin_liquidar_comprobantes.html', {
        'form': form,
        'comprobantes': comprobantes,
        'abrir_modal': not request.GET,
        'entidad_texto': _texto_entidad(getattr(form, 'cleaned_data', {}).get('entidad')),
    })


# --- Movimientos de caja -----------------------------------------------------

def _movimientos_caja_sin_liquidar(request):
    form = SinLiquidarFiltroForm(request.GET or None)
    movimientos = (
        MovimientoCaja.objects.filter(liquidaciones__isnull=True)
        .select_related(
            'caja', 'receptor',
            'emisor_relacion__id_entidad', 'rel_numero',
            'asiento_libro__libro', 'movimientocajadiferido',
        )
        .order_by('-emision', '-id')
    )
    if form.is_valid():
        entidad = form.cleaned_data.get('entidad')
        fecha_desde = form.cleaned_data.get('fecha_desde')
        fecha_hasta = form.cleaned_data.get('fecha_hasta')
        if entidad:
            movimientos = movimientos.filter(receptor=entidad)
        if fecha_desde:
            movimientos = movimientos.filter(emision__gte=fecha_desde)
        if fecha_hasta:
            movimientos = movimientos.filter(emision__lte=fecha_hasta)
    return form, movimientos


def sin_liquidar_movimientos_caja(request):
    form, movimientos = _movimientos_caja_sin_liquidar(request)
    movimientos = list(movimientos[:500])
    for m in movimientos:
        # La entidad contra la que se liquida es el receptor del movimiento
        # (mismo campo que usa liquidacion_form._armar_items).
        m.entidad_liquidar_id = m.receptor_id

    return render(request, 'liquidaciones/sin_liquidar_movimientos_caja.html', {
        'form': form,
        'movimientos': movimientos,
        'abrir_modal': not request.GET,
        'entidad_texto': _texto_entidad(getattr(form, 'cleaned_data', {}).get('entidad')),
    })


# --- Retenciones --------------------------------------------------------------

def _retenciones_sin_liquidar(request):
    form = SinLiquidarFiltroForm(request.GET or None)
    retenciones = (
        Retencion.objects.filter(liquidaciones__isnull=True)
        .select_related('entidad', 'id_regimen', 'id_impuesto')
        .order_by('-fecha', '-id')
    )
    if form.is_valid():
        entidad = form.cleaned_data.get('entidad')
        fecha_desde = form.cleaned_data.get('fecha_desde')
        fecha_hasta = form.cleaned_data.get('fecha_hasta')
        if entidad:
            retenciones = retenciones.filter(entidad=entidad)
        if fecha_desde:
            retenciones = retenciones.filter(fecha__gte=fecha_desde)
        if fecha_hasta:
            retenciones = retenciones.filter(fecha__lte=fecha_hasta)
    return form, retenciones


def sin_liquidar_retenciones(request):
    form, retenciones = _retenciones_sin_liquidar(request)
    retenciones = list(retenciones[:500])
    for r in retenciones:
        r.entidad_liquidar_id = r.entidad_id

    return render(request, 'liquidaciones/sin_liquidar_retenciones.html', {
        'form': form,
        'retenciones': retenciones,
        'abrir_modal': not request.GET,
        'entidad_texto': _texto_entidad(getattr(form, 'cleaned_data', {}).get('entidad')),
    })


# --- Retenciones INYM ---------------------------------------------------------

def _retenciones_inym_sin_liquidar(request):
    form = SinLiquidarFiltroForm(request.GET or None)
    retenciones_inym = (
        RetencionInym.objects.filter(liquidaciones__isnull=True)
        .select_related(
            'operador_emisor__entidad', 'operador_emisor__tipo_operador',
            'operador_retenido__entidad', 'operador_retenido__tipo_operador',
        )
        .order_by('-fecha', '-id')
    )
    if form.is_valid():
        entidad = form.cleaned_data.get('entidad')
        fecha_desde = form.cleaned_data.get('fecha_desde')
        fecha_hasta = form.cleaned_data.get('fecha_hasta')
        if entidad:
            # La entidad puede aparecer como parte emisora o como retenida:
            # se busca en cualquiera de las dos.
            retenciones_inym = retenciones_inym.filter(
                Q(operador_emisor__entidad=entidad) | Q(operador_retenido__entidad=entidad)
            )
        if fecha_desde:
            retenciones_inym = retenciones_inym.filter(fecha__gte=fecha_desde)
        if fecha_hasta:
            retenciones_inym = retenciones_inym.filter(fecha__lte=fecha_hasta)
    return form, retenciones_inym


def sin_liquidar_retenciones_inym(request):
    form, retenciones_inym = _retenciones_inym_sin_liquidar(request)
    retenciones_inym = list(retenciones_inym[:500])
    for ri in retenciones_inym:
        # La entidad contra la que se liquida es la del operador retenido
        # (mismo criterio que liquidacion_form._armar_items).
        ri.entidad_liquidar_id = ri.operador_retenido.entidad_id if ri.operador_retenido_id else None

    return render(request, 'liquidaciones/sin_liquidar_retenciones_inym.html', {
        'form': form,
        'retenciones_inym': retenciones_inym,
        'abrir_modal': not request.GET,
        'entidad_texto': _texto_entidad(getattr(form, 'cleaned_data', {}).get('entidad')),
    })


# --- Exportación a Excel / PDF de cada listado "sin liquidar" -----------------
#
# Cada función _filas_*_sin_liquidar devuelve un dict con:
#   columnas            : nombres de columna (encabezado)
#   filas               : lista de filas (una lista de valores por fila)
#   columnas_numericas  : índices (0-based) de las columnas que son montos/
#                         cantidades, para darles formato de miles con punto
#                         y decimales con coma (tanto en Excel como en PDF)
#   anchos              : ancho relativo de cada columna en el PDF (no hace
#                         falta que sumen 1, se normalizan solos); controla
#                         que las columnas con texto más largo (entidades,
#                         tipos, etc.) tengan más lugar y así no se corten ni
#                         se salgan del margen de la hoja.

FORMATO_MILES_EXCEL = '#,##0.00'  # con Excel en español/Argentina se ve como 1.234,56


def _numero_o_none(valor):
    """Convierte Decimal/float a float para que openpyxl no se queje; deja
    pasar None (celda vacía) tal cual."""
    if valor is None:
        return None
    return float(valor)


def _excel_response(nombre_archivo, resultado):
    import openpyxl
    from django.http import HttpResponse
    from openpyxl.utils import get_column_letter
    from services.gestorexcel import definir_estilo_general

    columnas = resultado['columnas']
    filas = resultado['filas']
    columnas_numericas = resultado.get('columnas_numericas', set())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sin liquidar'
    ws.append(columnas)
    fila_encabezado = ws.max_row
    for fila in filas:
        ws.append(fila)
    definir_estilo_general(ws)

    # Formato de miles con punto y decimales con coma (mismo criterio que
    # services/gestorexcel.py: el código de formato usa "," y "." como en
    # Excel en inglés, pero con Excel configurado en español/Argentina se
    # muestra como 1.234,56).
    for indice in columnas_numericas:
        letra_columna = get_column_letter(indice + 1)
        for celda in ws[letra_columna]:
            if celda.row > fila_encabezado:
                celda.number_format = FORMATO_MILES_EXCEL

    for columna in ws.columns:
        letra = columna[0].column_letter
        largo_max = max((len(str(c.value)) for c in columna if c.value is not None), default=10)
        ws.column_dimensions[letra].width = min(max(largo_max + 2, 10), 40)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename={nombre_archivo}.xlsx'
    wb.save(response)
    return response


def _pdf_response(nombre_archivo, titulo, resultado):
    from django.http import HttpResponse
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    from movimientos.templatetags.movimientos_extras import separador_miles

    columnas = resultado['columnas']
    filas = resultado['filas']
    columnas_numericas = resultado.get('columnas_numericas', set())
    anchos_relativos = resultado.get('anchos') or [1] * len(columnas)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename={nombre_archivo}.pdf'

    margen = 1 * cm
    doc = SimpleDocTemplate(
        response, pagesize=landscape(A4),
        topMargin=margen, bottomMargin=margen, leftMargin=margen, rightMargin=margen,
    )
    ancho_disponible = landscape(A4)[0] - doc.leftMargin - doc.rightMargin

    estilos = getSampleStyleSheet()
    estilo_celda = ParagraphStyle('celda_sin_liquidar', parent=estilos['Normal'], fontSize=7, leading=8.5)
    estilo_encabezado = ParagraphStyle(
        'encabezado_sin_liquidar', parent=estilo_celda, textColor=colors.white, fontName='Helvetica-Bold',
    )

    elementos = [Paragraph(titulo, estilos['Title']), Spacer(1, 0.4 * cm)]

    def formatear_valor(indice, valor):
        if valor is None or valor == '':
            return ''
        if indice in columnas_numericas:
            return separador_miles(valor)
        return str(valor)

    # Todo el contenido pasa por Paragraph (no texto plano) para que, si no
    # entra en el ancho de columna, el texto haga salto de línea en vez de
    # forzar la columna a ensancharse y terminar saliéndose de la hoja.
    fila_encabezado = [Paragraph(str(col), estilo_encabezado) for col in columnas]
    datos = [fila_encabezado]
    for fila in filas:
        datos.append([
            Paragraph(formatear_valor(indice, valor), estilo_celda)
            for indice, valor in enumerate(fila)
        ])

    total_relativo = sum(anchos_relativos) or 1
    col_widths = [ancho_disponible * (peso / total_relativo) for peso in anchos_relativos]

    tabla = Table(datos, colWidths=col_widths, repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#343a40')),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]))
    elementos.append(tabla)
    doc.build(elementos)
    return response


def _filas_comprobantes_sin_liquidar(request):
    _, comprobantes = _comprobantes_sin_liquidar(request)
    entidad_propia = _entidad_propia()
    columnas = ['ID', 'Fecha', 'Entidad emisora', 'Entidad receptora', 'Tipo de comprobante', 'Punto de venta', 'Número', 'Total']
    filas = []
    for c in comprobantes:
        if c.es_emisor == 0:
            emisora, receptora = entidad_propia, c.entidad_emisor
        else:
            emisora, receptora = c.entidad_emisor, entidad_propia
        filas.append([
            c.id, c.fecha, str(emisora) if emisora else '', str(receptora) if receptora else '',
            str(c.tipo_comprobante) if c.tipo_comprobante else '', c.punto_de_venta, c.numero,
            _numero_o_none(c.total),
        ])
    return {
        'columnas': columnas,
        'filas': filas,
        'columnas_numericas': {7},  # Total
        'anchos': [0.6, 0.9, 2.1, 2.1, 1.7, 1.0, 0.9, 1.2],
    }


def _filas_movimientos_caja_sin_liquidar(request):
    _, movimientos = _movimientos_caja_sin_liquidar(request)
    columnas = ['ID', 'Fecha emisión', 'Entidad emisora', 'Entidad receptora', 'Número', 'Caja', 'Libro', 'Libro - Hoja', 'Libro - Renglón', 'Diferido', 'Efectivización', 'Monto']
    filas = []
    for m in movimientos:
        try:
            asiento = m.asiento_libro
        except Exception:
            asiento = None
        try:
            diferido = m.movimientocajadiferido.diferido
        except Exception:
            diferido = None
        filas.append([
            m.id, m.emision, str(m.emisor) if m.emisor else '', str(m.receptor) if m.receptor else '',
            m.numero, str(m.caja) if m.caja else '',
            str(asiento.libro) if asiento and asiento.libro else '',
            asiento.hoja if asiento else None,
            asiento.renglon if asiento else None,
            diferido, m.efectivizacion, _numero_o_none(m.monto),
        ])
    return {
        'columnas': columnas,
        'filas': filas,
        'columnas_numericas': {11},  # Monto
        'anchos': [0.5, 0.8, 1.7, 1.7, 0.7, 0.9, 1.2, 0.7, 0.8, 0.8, 0.8, 1.0],
    }


def _filas_retenciones_sin_liquidar(request):
    _, retenciones = _retenciones_sin_liquidar(request)
    columnas = ['ID', 'Fecha', 'Entidad', 'Régimen', 'Impuesto', 'Año', 'Número', 'Total']
    filas = []
    for r in retenciones:
        filas.append([
            r.id, r.fecha, str(r.entidad) if r.entidad else '',
            str(r.id_regimen) if r.id_regimen else '', str(r.id_impuesto) if r.id_impuesto else '',
            r.año, r.numero, _numero_o_none(r.total),
        ])
    return {
        'columnas': columnas,
        'filas': filas,
        'columnas_numericas': {7},  # Total
        'anchos': [0.6, 0.9, 2.1, 1.9, 1.7, 0.7, 0.8, 1.1],
    }


def _filas_retenciones_inym_sin_liquidar(request):
    _, retenciones_inym = _retenciones_inym_sin_liquidar(request)
    columnas = ['ID', 'Fecha', 'Entidad emisora', 'Operador INYM emisor', 'Entidad retenida', 'Operador INYM retenido', 'Kgs', 'Tarifa', 'Eliminación', 'Total']
    filas = []
    for ri in retenciones_inym:
        emisora = ri.operador_emisor.entidad if ri.operador_emisor_id else None
        retenido = ri.operador_retenido.entidad if ri.operador_retenido_id else None
        filas.append([
            ri.id, ri.fecha,
            str(emisora) if emisora else '', str(ri.operador_emisor) if ri.operador_emisor_id else '',
            str(retenido) if retenido else '', str(ri.operador_retenido) if ri.operador_retenido_id else '',
            _numero_o_none(ri.kgs), _numero_o_none(ri.tarifa), ri.eliminacion, _numero_o_none(ri.total),
        ])
    return {
        'columnas': columnas,
        'filas': filas,
        'columnas_numericas': {6, 7, 9},  # Kgs, Tarifa, Total
        'anchos': [0.5, 0.8, 1.6, 1.6, 1.6, 1.6, 0.8, 0.8, 0.8, 1.0],
    }


_SIN_LIQUIDAR_FILAS = {
    'comprobantes': ('comprobantes_sin_liquidar', 'Comprobantes sin liquidar', _filas_comprobantes_sin_liquidar),
    'movimientos_caja': ('movimientos_caja_sin_liquidar', 'Movimientos de caja sin liquidar', _filas_movimientos_caja_sin_liquidar),
    'retenciones': ('retenciones_sin_liquidar', 'Retenciones sin liquidar', _filas_retenciones_sin_liquidar),
    'retenciones_inym': ('retenciones_inym_sin_liquidar', 'Retenciones INYM sin liquidar', _filas_retenciones_inym_sin_liquidar),
}


def sin_liquidar_excel(request, tipo):
    config = _SIN_LIQUIDAR_FILAS.get(tipo)
    if not config:
        messages.error(request, 'Listado de "sin liquidar" desconocido.')
        return redirect('liquidaciones:listado')
    nombre_archivo, _titulo, obtener_filas = config
    resultado = obtener_filas(request)
    return _excel_response(nombre_archivo, resultado)


def sin_liquidar_pdf(request, tipo):
    config = _SIN_LIQUIDAR_FILAS.get(tipo)
    if not config:
        messages.error(request, 'Listado de "sin liquidar" desconocido.')
        return redirect('liquidaciones:listado')
    nombre_archivo, titulo, obtener_filas = config
    resultado = obtener_filas(request)
    return _pdf_response(nombre_archivo, titulo, resultado)