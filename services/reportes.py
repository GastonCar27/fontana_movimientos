"""
Helpers genéricos para exportar un reporte tabular (columnas + filas ya
armadas por la vista) a Excel o PDF.

Es la extracción, sin cambios de comportamiento, de las funciones
'_excel_response' / '_pdf_response' que ya existían duplicadas dentro de
movimientos_caja/views.py (reporte de movimientos de caja y ranking de
entidades). Se deja acá para que cualquier app nueva (como el ranking de
productores de movimientos/views.py) las reutilice en vez de volver a
copiar el mismo código de openpyxl/reportlab por tercera vez.

'resultado' es siempre un dict con:
    columnas:           lista de encabezados (str)
    filas:              lista de listas, un valor por columna
    columnas_numericas: set() con los índices (0-based) de columnas que son
                         montos/números (se formatean con separador de miles)
    anchos:             (solo para PDF) lista de anchos relativos de columna,
                         mismo largo que 'columnas'; si se omite, todas iguales
"""


def excel_response(nombre_archivo, resultado):
    import openpyxl
    from openpyxl.utils import get_column_letter
    from django.http import HttpResponse

    from services.gestorexcel import definir_estilo_general

    columnas = resultado['columnas']
    filas = resultado['filas']
    columnas_numericas = resultado.get('columnas_numericas', set())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Reporte'
    ws.append(columnas)
    fila_encabezado = ws.max_row
    for fila in filas:
        ws.append(fila)
    definir_estilo_general(ws)

    for indice in columnas_numericas:
        letra_columna = get_column_letter(indice + 1)
        for celda in ws[letra_columna]:
            if celda.row > fila_encabezado:
                celda.number_format = '#,##0.00'

    for columna in ws.columns:
        letra = columna[0].column_letter
        largo_max = max((len(str(c.value)) for c in columna if c.value is not None), default=10)
        ws.column_dimensions[letra].width = min(max(largo_max + 2, 10), 40)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename={nombre_archivo}.xlsx'
    wb.save(response)
    return response


def pdf_response(nombre_archivo, titulo, resultado):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from django.http import HttpResponse

    from movimientos.templatetags.movimientos_extras import separador_miles

    columnas = resultado['columnas']
    filas = resultado['filas']
    columnas_numericas = resultado.get('columnas_numericas', set())
    anchos_relativos = resultado.get('anchos') or [1] * len(columnas)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename={nombre_archivo}.pdf'

    margen = 1 * cm
    doc = SimpleDocTemplate(
        response, pagesize=landscape(A4),
        topMargin=margen, bottomMargin=margen, leftMargin=margen, rightMargin=margen,
    )
    ancho_disponible = landscape(A4)[0] - doc.leftMargin - doc.rightMargin

    estilos = getSampleStyleSheet()
    estilo_celda = ParagraphStyle('celda_reporte', parent=estilos['Normal'], fontSize=7, leading=8.5)
    estilo_encabezado = ParagraphStyle(
        'encabezado_reporte', parent=estilo_celda, textColor=colors.white, fontName='Helvetica-Bold',
    )

    elementos = [Paragraph(titulo, estilos['Title']), Spacer(1, 0.4 * cm)]

    def formatear_valor(indice, valor):
        if valor is None or valor == '':
            return ''
        if indice in columnas_numericas:
            return separador_miles(valor)
        return str(valor)

    fila_encabezado = [Paragraph(str(col), estilo_encabezado) for col in columnas]
    datos = [fila_encabezado]
    for fila in filas:
        datos.append([
            Paragraph(formatear_valor(indice, valor), estilo_celda)
            for indice, valor in enumerate(fila)
        ])

    total_relativo = sum(anchos_relativos) or 1
    col_widths = [ancho_disponible * (peso / total_relativo) for peso in anchos_relativos]

    tabla = Table(datos, colWidths=col_widths, repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#343a40')),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]))
    elementos.append(tabla)
    doc.build(elementos)
    return response
