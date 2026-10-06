"""
Exporta a Excel los comprobantes CARGADOS en el sistema en una fecha dada,
según comprobante.fecha_agregado (pedido de Gastón, 06/10/2026: para revisar
lo que entró del CSV de AFIP).

fecha_agregado se guarda en UTC; la fecha que se pide acá es la fecha de
Argentina (America/Argentina/Buenos_Aires), así que se toma el día completo
local (00:00 a 23:59 hora argentina).

Uso:
    python manage.py exportar_comprobantes_agregados                    # hoy
    python manage.py exportar_comprobantes_agregados --fecha 2026-10-06
    python manage.py exportar_comprobantes_agregados --desde 2026-10-01 --hasta 2026-10-06
    python manage.py exportar_comprobantes_agregados --salida "documentacion\\comprobantes_hoy.xlsx"
"""
import datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from django.core.management.base import BaseCommand, CommandError
from django.utils import timezone

from comprobantes.models import Comprobante

TZ_AR = ZoneInfo('America/Argentina/Buenos_Aires')


def _fecha(texto, nombre):
    try:
        return datetime.date.fromisoformat(texto)
    except ValueError:
        raise CommandError(f'{nombre} inválida: "{texto}" (usar AAAA-MM-DD).')


def _num(valor):
    return float(valor) if valor is not None else None


class Command(BaseCommand):
    help = 'Exporta a Excel los comprobantes cargados (fecha_agregado) en una fecha o rango (hora argentina).'

    def add_arguments(self, parser):
        parser.add_argument('--fecha', help='Día de carga (AAAA-MM-DD). Default: hoy.')
        parser.add_argument('--desde', help='Desde (AAAA-MM-DD), en vez de --fecha.')
        parser.add_argument('--hasta', help='Hasta (AAAA-MM-DD), en vez de --fecha.')
        parser.add_argument('--salida', help='Archivo .xlsx a generar. Default: comprobantes_agregados_<fecha>.xlsx')

    def handle(self, *args, **opts):
        hoy = timezone.now().astimezone(TZ_AR).date()
        if opts.get('desde') or opts.get('hasta'):
            desde = _fecha(opts['desde'], '--desde') if opts.get('desde') else hoy
            hasta = _fecha(opts['hasta'], '--hasta') if opts.get('hasta') else desde
        else:
            desde = hasta = _fecha(opts['fecha'], '--fecha') if opts.get('fecha') else hoy
        if hasta < desde:
            raise CommandError('--hasta no puede ser anterior a --desde.')

        inicio = datetime.datetime.combine(desde, datetime.time.min, tzinfo=TZ_AR)
        fin = datetime.datetime.combine(hasta + datetime.timedelta(days=1), datetime.time.min, tzinfo=TZ_AR)

        comprobantes = list(
            Comprobante.objects
            .filter(fecha_agregado__gte=inicio, fecha_agregado__lt=fin)
            .order_by('fecha_agregado', 'id')
        )
        # Sin select_related('entidad_emisor'): ese FK no está declarado
        # null=True en el modelo, así que Django haría un INNER JOIN y se
        # perderían los comprobantes con id_entidad NULL (típico del CSV de
        # AFIP cuando no se pudo matchear la entidad). Se resuelven aparte.
        from entidades.models import Entidad
        from comprobantes.models import ComprobanteTipo
        entidades = {e.id: e for e in Entidad.objects.filter(
            id__in={c.entidad_emisor_id for c in comprobantes if c.entidad_emisor_id})}
        tipos = {t.id: t for t in ComprobanteTipo.objects.filter(
            id__in={c.tipo_comprobante_id for c in comprobantes if c.tipo_comprobante_id})}

        rango = desde.strftime('%d/%m/%Y') + ('' if hasta == desde else ' al ' + hasta.strftime('%d/%m/%Y'))
        self.stdout.write(f'Comprobantes cargados el {rango} (hora argentina): {len(comprobantes)}')
        if not comprobantes:
            sin_fecha = Comprobante.objects.filter(fecha_agregado__isnull=True).count()
            self.stdout.write(self.style.WARNING(
                f'No hay comprobantes con fecha de carga en ese rango. (Comprobantes sin fecha de carga: {sin_fecha}. '
                'Si los cargaste hoy y aparecen acá, la columna no tomó el valor por defecto en MySQL.)'
            ))
            return

        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Comprobantes cargados'
        ws.append([f'Comprobantes cargados el {rango} (hora argentina) -- {len(comprobantes)} comprobantes'])
        ws['A1'].font = Font(bold=True, size=12)
        ws.append([])

        columnas = [
            'ID', 'Fecha comprobante', 'Tipo', 'Punto de venta', 'Número', 'Número hasta', 'Entidad', 'CUIT',
            'Fontana es', 'Neto gravado', 'Neto no gravado', 'Exento', 'IVA', 'Otros tributos', 'Total',
            'Moneda', 'Cód. autorización', 'Agregado desde', 'Fecha y hora de carga',
        ]
        ws.append(columnas)
        fila_enc = ws.max_row
        for celda in ws[fila_enc]:
            celda.font = Font(bold=True, color='FFFFFF')
            celda.fill = PatternFill('solid', fgColor='2F5496')
            celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        numericas = {10, 11, 12, 13, 14, 15}  # 1-based: Neto gravado .. Total
        totales = {c: Decimal('0') for c in numericas}
        for c in comprobantes:
            entidad = entidades.get(c.entidad_emisor_id)
            nombre = (entidad.nombre if entidad else None) or c.entidad_nombre or ''
            cuit = (entidad.cuit if entidad else '') or ''
            if c.es_emisor == 0:
                rol = 'Emisora'
            elif c.es_emisor == 1:
                rol = 'Receptora'
            else:
                rol = ''
            cargado = c.fecha_agregado.astimezone(TZ_AR).replace(tzinfo=None) if c.fecha_agregado else None
            fila = [
                c.id, c.fecha, str(tipos.get(c.tipo_comprobante_id) or ''), c.punto_de_venta, c.numero,
                c.numero_hasta, nombre, cuit, rol,
                _num(c.neto_gravado), _num(c.neto_no_gravado), _num(c.exento), _num(c.iva), _num(c.otros_tributos),
                _num(c.total), c.moneda or '',
                ('%.0f' % c.codigo_autorizacion) if c.codigo_autorizacion else '',
                c.agregado_desde or '', cargado,
            ]
            ws.append(fila)
            for col in numericas:
                valor = fila[col - 1]
                if valor is not None:
                    totales[col] += Decimal(str(valor))

        ultima = ws.max_row
        ws.append(['Total'] + [''] * 8 + [float(totales[c]) for c in sorted(numericas)] + [''] * 4)
        for celda in ws[ws.max_row]:
            celda.font = Font(bold=True)
            celda.fill = PatternFill('solid', fgColor='E9ECEF')

        for fila in ws.iter_rows(min_row=fila_enc + 1, max_row=ws.max_row):
            for celda in fila:
                if celda.column in numericas:
                    celda.number_format = '#,##0.00'
                elif isinstance(celda.value, datetime.datetime):
                    celda.number_format = 'DD/MM/YYYY HH:MM'
                elif isinstance(celda.value, datetime.date):
                    celda.number_format = 'DD/MM/YYYY'

        anchos = [8, 12, 22, 9, 10, 10, 38, 14, 11, 14, 14, 12, 13, 13, 15, 8, 18, 16, 17]
        for i, ancho in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
        ws.row_dimensions[fila_enc].height = 30
        ws.freeze_panes = ws.cell(row=fila_enc + 1, column=1)
        ws.auto_filter.ref = f'A{fila_enc}:{get_column_letter(len(columnas))}{ultima}'

        salida = opts.get('salida') or (
            f'comprobantes_agregados_{desde:%Y-%m-%d}'
            + ('' if hasta == desde else f'_al_{hasta:%Y-%m-%d}') + '.xlsx'
        )
        wb.save(salida)
        self.stdout.write(self.style.SUCCESS(f'Excel generado: {salida}'))
