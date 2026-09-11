from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.db import IntegrityError, transaction
from django.db.models import Case, Count, DecimalField, F, IntegerField, ProtectedError, Q, Sum, Value, When
from django.db.models.functions import Coalesce
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme

from entidades.models import Entidad
from services.buscadores import texto_entidad_buscador
from services.ordenamiento import aplicar_orden_lista, aplicar_orden_queryset
from services.permisos import requiere_grupo

from .forms import (
    AsignarLibroMovimientoForm,
    EstadoCajaForm,
    MovimientoCajaForm,
    MovimientoCajaRelacionadosForm,
    MovimientoCajaReporteForm,
    RankingEntidadesForm,
)
from .models import (
    BancoCuentaEntidad,
    Caja,
    LibroCaja,
    LibroMovim,
    MovimientoCaja,
    MovimientoCajaBancoCuentaEntidad,
    MovimientoCajaConcepto,
    MovimientoCajaDiferido,
    MovimientoCajaEmisor,
    MovimientoCajaNumero,
)

# Tope de renglones por hoja del libro de caja: al llegar al 25, el
# siguiente movimiento de un alta encadenada pasa a renglón 1 de la hoja
# siguiente (ver _proximo_renglon_y_hoja).
RENGLON_MAXIMO_POR_HOJA = 25


def _url_next_segura(request, next_url):
    """Devuelve next_url si es una URL interna válida (para volver, tras
    editar/eliminar, a la página que corresponda: el listado general de
    'Modificación' o el de 'Modificar en libro'); si no, devuelve None."""
    if next_url and url_has_allowed_host_and_scheme(
        next_url, allowed_hosts={request.get_host()}, require_https=request.is_secure()
    ):
        return next_url
    return None


# ---------------------------------------------------------------------------
# Alta / Modificación (misma vista, pk=None para alta)
# ---------------------------------------------------------------------------

def _proximo_renglon_y_hoja(ultimo_renglon, ultima_hoja):
    """Dado el renglón/hoja del movimiento recién guardado en un alta, calcula
    los valores a precargar para el próximo: el renglón +1 manteniendo la
    misma hoja, salvo que el renglón ya guardado haya llegado al tope
    (RENGLON_MAXIMO_POR_HOJA = 25), en cuyo caso el próximo pasa a ser
    renglón 1 de la hoja siguiente (+1)."""
    if ultimo_renglon is None:
        return None, ultima_hoja
    if ultimo_renglon >= RENGLON_MAXIMO_POR_HOJA:
        proxima_hoja = (ultima_hoja + 1) if ultima_hoja is not None else None
        return 1, proxima_hoja
    return ultimo_renglon + 1, ultima_hoja


@transaction.atomic
def movimiento_caja_form(request, pk=None):
    """Alta y modificación de un Movimiento de Caja, junto con los datos de
    sus tablas relacionadas: libro/hoja/renglón/número (LibroMovim /
    MovimientoCajaNumero), fecha de diferido (MovimientoCajaDiferido),
    concepto (MovimientoCajaConcepto) y cuenta bancaria del receptor
    (MovimientoCajaBancoCuentaEntidad); antes sólo se podían cargar desde el
    admin.

    El siguiente ID y la validación (full_clean) del movimiento en sí ya los
    resuelve MovimientoCaja.save(). Los registros relacionados son opcionales:
    si se dejan vacíos no se guardan (o se borran, si ya existían).

    En el alta (pk=None), al guardar no se vuelve al listado: se vuelve a
    mostrar esta misma pantalla de alta, con el banco (caja), el libro y la
    fecha de emisión iguales a los del movimiento recién guardado, el
    renglón y el número +1 (renglón con tope 25 -> pasa a 1 y la hoja +1), y
    el resto de los campos vacíos, para poder cargar rápido varios
    movimientos seguidos del mismo lote. En la modificación de un movimiento
    existente el comportamiento no cambia: se vuelve al listado.
    """
    movimiento = get_object_or_404(MovimientoCaja, pk=pk) if pk else None
    es_alta = movimiento is None

    receptor_id = None
    if request.method == 'POST':
        receptor_id = request.POST.get('receptor') or None
    elif movimiento:
        receptor_id = movimiento.receptor_id

    if request.method == 'POST':
        form_rel = MovimientoCajaRelacionadosForm(request.POST, receptor_id=receptor_id)
        form_rel_valido = form_rel.is_valid()

        if movimiento is not None and form_rel_valido:
            # Le "adelantamos" al clean() de MovimientoCaja (que compara la
            # caja del libro asignado con la caja del movimiento) el libro
            # que se va a guardar a continuación, no el que tenía antes, para
            # que no rechace un cambio de caja hecho junto con un cambio de
            # libro en esta misma pantalla.
            movimiento.asiento_libro = LibroMovim(
                movimiento_caja=movimiento,
                libro=form_rel.cleaned_data.get('libro'),
                hoja=form_rel.cleaned_data.get('hoja'),
                renglon=form_rel.cleaned_data.get('renglon'),
            )

        form = MovimientoCajaForm(request.POST, instance=movimiento)
        form_valido = form.is_valid()

        # Validaciones cruzadas entre 'form' (emisión/efectivización) y
        # 'form_rel' (diferido): no pueden resolverse en el clean() de cada
        # form por separado porque son dos formularios distintos.
        if form_valido and form_rel_valido:
            emision = form.cleaned_data.get('emision')
            efectivizacion = form.cleaned_data.get('efectivizacion')
            diferido = form_rel.cleaned_data.get('diferido')
            if emision and diferido and diferido < emision:
                form_rel.add_error('diferido', 'La fecha de diferido no puede ser anterior a la fecha de emisión.')
                form_rel_valido = False
            if efectivizacion and diferido and diferido > efectivizacion:
                form_rel.add_error('diferido', 'La fecha de diferido no puede ser posterior a la fecha de efectivización.')
                form_rel_valido = False

        if form_valido and form_rel_valido:
            nuevo = form.save()

            libro = form_rel.cleaned_data.get('libro')
            hoja = form_rel.cleaned_data.get('hoja')
            renglon = form_rel.cleaned_data.get('renglon')
            if libro:
                LibroMovim.objects.update_or_create(
                    movimiento_caja=nuevo,
                    defaults={
                        'libro': libro,
                        'hoja': hoja,
                        'renglon': renglon,
                    },
                )
            else:
                LibroMovim.objects.filter(movimiento_caja=nuevo).delete()

            diferido = form_rel.cleaned_data.get('diferido')
            if diferido:
                MovimientoCajaDiferido.objects.update_or_create(id=nuevo, defaults={'diferido': diferido})
            else:
                MovimientoCajaDiferido.objects.filter(id=nuevo).delete()

            concepto_tipo = form_rel.cleaned_data.get('concepto_tipo')
            if concepto_tipo:
                MovimientoCajaConcepto.objects.update_or_create(
                    movimiento_caja=nuevo, defaults={'concepto_tipo': concepto_tipo},
                )
            else:
                MovimientoCajaConcepto.objects.filter(movimiento_caja=nuevo).delete()

            numero = form_rel.cleaned_data.get('numero')
            if numero is not None:
                MovimientoCajaNumero.objects.update_or_create(
                    movimiento_caja=nuevo, defaults={'numero': numero},
                )
            else:
                MovimientoCajaNumero.objects.filter(movimiento_caja=nuevo).delete()

            emisor = form_rel.cleaned_data.get('emisor')
            if emisor:
                MovimientoCajaEmisor.objects.update_or_create(id=nuevo, defaults={'id_entidad': emisor})
            else:
                MovimientoCajaEmisor.objects.filter(id=nuevo).delete()

            cuenta_bancaria = form_rel.cleaned_data.get('cuenta_bancaria_entidad')
            if cuenta_bancaria:
                numero_destino = (cuenta_bancaria.numero or cuenta_bancaria.cbu or '')[:30]
                MovimientoCajaBancoCuentaEntidad.objects.update_or_create(
                    id=nuevo, defaults={'numero_cuenta_entidad_destino': numero_destino},
                )
            else:
                MovimientoCajaBancoCuentaEntidad.objects.filter(id=nuevo).delete()

            messages.success(request, f'Movimiento de caja {nuevo.id} guardado correctamente.')

            if es_alta:
                proximo_renglon, proxima_hoja = _proximo_renglon_y_hoja(renglon, hoja)
                proximo_numero = (numero + 1) if numero is not None else None
                request.session['movimiento_caja_prefill'] = {
                    'caja': nuevo.caja_id,
                    'emision': nuevo.emision.isoformat() if nuevo.emision else None,
                    'libro': libro.id if libro else None,
                    'hoja': proxima_hoja,
                    'renglon': proximo_renglon,
                    'numero': proximo_numero,
                }
                return redirect('movimientos_caja:movimiento_caja_alta')

            return redirect('movimientos_caja:movimiento_caja_modificar')
    else:
        prefill = request.session.pop('movimiento_caja_prefill', None) if es_alta else None

        initial_form = {}
        if prefill:
            if prefill.get('caja') is not None:
                initial_form['caja'] = prefill['caja']
            if prefill.get('emision') is not None:
                initial_form['emision'] = prefill['emision']
        form = MovimientoCajaForm(instance=movimiento, initial=initial_form)

        initial_rel = {}
        if movimiento:
            asiento = LibroMovim.objects.filter(movimiento_caja=movimiento).first()
            diferido_obj = MovimientoCajaDiferido.objects.filter(id=movimiento).first()
            concepto_obj = MovimientoCajaConcepto.objects.filter(movimiento_caja=movimiento).first()
            emisor_obj = MovimientoCajaEmisor.objects.filter(id=movimiento).first()
            cuenta_obj = MovimientoCajaBancoCuentaEntidad.objects.filter(id=movimiento).first()

            if asiento:
                initial_rel.update({'libro': asiento.libro_id, 'hoja': asiento.hoja, 'renglon': asiento.renglon})
            if diferido_obj:
                initial_rel['diferido'] = diferido_obj.diferido
            if concepto_obj:
                initial_rel['concepto_tipo'] = concepto_obj.concepto_tipo_id
            if movimiento.numero is not None:
                initial_rel['numero'] = movimiento.numero
            if emisor_obj and emisor_obj.id_entidad_id:
                initial_rel['emisor'] = emisor_obj.id_entidad_id
            if cuenta_obj and cuenta_obj.numero_cuenta_entidad_destino and movimiento.receptor_id:
                # numero_cuenta_entidad_destino es un texto libre (no una FK):
                # tratamos de reencontrar, a partir de él, la cuenta bancaria
                # del receptor que lo originó, para dejarla preseleccionada.
                cuenta_match = BancoCuentaEntidad.objects.filter(entidad_id=movimiento.receptor_id).filter(
                    Q(numero=cuenta_obj.numero_cuenta_entidad_destino)
                    | Q(cbu=cuenta_obj.numero_cuenta_entidad_destino)
                ).first()
                if cuenta_match:
                    initial_rel['cuenta_bancaria_entidad'] = cuenta_match.id
        elif prefill:
            if prefill.get('libro') is not None:
                initial_rel['libro'] = prefill['libro']
            if prefill.get('hoja') is not None:
                initial_rel['hoja'] = prefill['hoja']
            if prefill.get('renglon') is not None:
                initial_rel['renglon'] = prefill['renglon']
            if prefill.get('numero') is not None:
                initial_rel['numero'] = prefill['numero']

        form_rel = MovimientoCajaRelacionadosForm(initial=initial_rel, receptor_id=receptor_id)

    return render(request, 'movimientos_caja/movimiento_caja_form.html', {
        'form': form,
        'form_rel': form_rel,
        'movimiento': movimiento,
        'libros': LibroCaja.objects.select_related('caja').order_by('nombre'),
    })


def movimiento_caja_entidad_cuentas_bancarias(request, entidad_id):
    """JSON con las cuentas bancarias registradas de una entidad, para poblar
    el selector 'Cuenta bancaria del receptor' del alta/modificación de un
    movimiento de caja al elegir (o cambiar) el receptor, sin recargar la
    página.
    """
    cuentas = (
        BancoCuentaEntidad.objects.filter(entidad_id=entidad_id)
        .select_related('cuenta_tipo', 'producto_tipo')
        .order_by('id')
    )

    def _etiqueta(c):
        partes = [c.cuenta_tipo.nombre if c.cuenta_tipo_id and c.cuenta_tipo else 'Cuenta']
        if c.producto_tipo_id and c.producto_tipo:
            partes.append(c.producto_tipo.nombre)
        partes.append(c.numero or c.cbu or 's/n')
        return ' · '.join(p for p in partes if p)

    data = [{'id': c.id, 'label': _etiqueta(c)} for c in cuentas]
    return JsonResponse({'cuentas': data})


# ---------------------------------------------------------------------------
# Modificación: listado/búsqueda (puerta de entrada para editar)
# ---------------------------------------------------------------------------

def movimiento_caja_listado(request):
    """Listado/búsqueda de movimientos de caja; es la puerta de entrada de 'Modificación'.

    Trae junto con cada movimiento sus liquidaciones asociadas
    (liquidacion_movimiento -> liquidacion) con prefetch_related, para
    poder mostrar en la columna "Liquidación" si el movimiento ya está
    liquidado y con qué id, sin pegarle a la base una vez por fila.
    """
    movimientos = (
        MovimientoCaja.objects.select_related('caja', 'tipo', 'receptor', 'rel_numero', 'emisor_relacion__id_entidad')
        .prefetch_related('liquidaciones__liquidacion')
        .order_by('-emision', '-id')
    )

    q_receptor = request.GET.get('receptor', '').strip()
    q_id = request.GET.get('id', '').strip()
    q_fecha = request.GET.get('fecha', '').strip()
    q_caja = request.GET.get('caja', '').strip()
    q_numero = request.GET.get('numero', '').strip()
    q_monto = request.GET.get('monto', '').strip()

    if q_receptor:
        # Busca tanto en el receptor como en el emisor (si tiene uno cargado
        # explícitamente): antes sólo miraba el receptor, así que un
        # movimiento donde la entidad buscada es la emisora no aparecía.
        movimientos = movimientos.filter(
            Q(receptor__nombre__icontains=q_receptor) | Q(receptor__cuit__icontains=q_receptor)
            | Q(emisor_relacion__id_entidad__nombre__icontains=q_receptor)
            | Q(emisor_relacion__id_entidad__cuit__icontains=q_receptor)
        )
    if q_id:
        if q_id.isdigit():
            movimientos = movimientos.filter(id=int(q_id))
        else:
            movimientos = movimientos.none()
    if q_fecha:
        movimientos = movimientos.filter(emision=q_fecha)
    if q_caja.isdigit():
        movimientos = movimientos.filter(caja_id=int(q_caja))
    if q_numero:
        if q_numero.isdigit():
            movimientos = movimientos.filter(rel_numero__numero=int(q_numero))
        else:
            movimientos = movimientos.none()
    if q_monto:
        try:
            monto_valor = Decimal(q_monto.replace('.', '').replace(',', '.')) if ',' in q_monto else Decimal(q_monto)
        except InvalidOperation:
            movimientos = movimientos.none()
        else:
            movimientos = movimientos.filter(monto=monto_valor)

    movimientos = aplicar_orden_queryset(request, movimientos, {
        'id': 'id',
        'caja': 'caja__nombre',
        'tipo': 'tipo__nombre',
        'emision': 'emision',
        'numero': 'rel_numero__numero',
        'monto': 'monto',
        'emisor': 'emisor_relacion__id_entidad__nombre',
        'receptor': 'receptor__nombre',
        'efectivizacion': 'efectivizacion',
    })

    return render(request, 'movimientos_caja/movimiento_caja_listado.html', {
        'movimientos': movimientos[:200],
        'q_receptor': q_receptor,
        'q_id': q_id,
        'q_fecha': q_fecha,
        'q_caja': q_caja,
        'q_numero': q_numero,
        'q_monto': q_monto,
    })


def movimiento_caja_eliminar(request, pk):
    """Confirmación + baja de un Movimiento de Caja. Se accede tanto desde el
    listado de 'Modificación' como desde el de 'Modificar en libro'; el
    parámetro GET/POST 'next' (si es una URL interna válida) indica a cuál de
    los dos volver al terminar, y por defecto vuelve al de 'Modificación'.
    Si el movimiento ya está incluido en una liquidación no se permite
    borrarlo desde acá (mismo criterio que retenciones.views.
    retencion_eliminar), para no dejar una liquidación apuntando a un
    movimiento inexistente.
    """
    movimiento = get_object_or_404(MovimientoCaja, pk=pk)
    next_url = _url_next_segura(request, request.POST.get('next') or request.GET.get('next'))
    destino = next_url or 'movimientos_caja:movimiento_caja_modificar'

    if movimiento.liquidaciones.exists():
        messages.error(
            request,
            f'El movimiento de caja {pk} ya está incluido en una liquidación y no se puede '
            'eliminar desde acá.'
        )
        return redirect(destino)

    if request.method == 'POST':
        try:
            movimiento.delete()
        except (ProtectedError, IntegrityError):
            messages.error(
                request,
                f'El movimiento de caja {pk} no se puede eliminar porque está siendo usado '
                'en otro registro.'
            )
        else:
            messages.success(request, f'El movimiento de caja {pk} se eliminó correctamente.')
        return redirect(destino)

    return render(request, 'movimientos_caja/movimiento_caja_eliminar_confirm.html', {
        'movimiento': movimiento,
        'next': next_url or '',
    })


# ---------------------------------------------------------------------------
# Modificar en libro: asignar/editar libro, hoja y renglón sin pasar por el
# admin de Django, mostrando primero los movimientos que todavía no tienen
# libro u hoja asignados (o que tienen -1 cargado).
# ---------------------------------------------------------------------------

def movimiento_caja_libro_listado(request):
    """Listado de movimientos de caja pensado para asignarles (o corregirles)
    el libro, la hoja y el renglón: el equivalente, dentro de la app, al
    listado de LibroMovim del admin. Permite filtrar por cuenta de banco
    (caja), por libro y por hoja, y ordena primero los movimientos que no
    tienen libro u hoja asignado (o tienen -1 cargado), para poder
    completarlos.
    """
    movimientos = (
        MovimientoCaja.objects.select_related('caja', 'receptor', 'rel_numero')
        .select_related('asiento_libro', 'asiento_libro__libro', 'asiento_libro__libro__caja')
        .annotate(
            sin_asignar=Case(
                When(asiento_libro__isnull=True, then=Value(0)),
                When(asiento_libro__libro__isnull=True, then=Value(0)),
                When(asiento_libro__libro_id=-1, then=Value(0)),
                When(asiento_libro__hoja__isnull=True, then=Value(0)),
                When(asiento_libro__hoja=-1, then=Value(0)),
                default=Value(1),
                output_field=IntegerField(),
            )
        )
    )

    q_caja = request.GET.get('caja', '').strip()
    q_libro = request.GET.get('libro', '').strip()
    q_hoja = request.GET.get('hoja', '').strip()
    q_solo_sin_asignar = request.GET.get('sin_asignar') == '1'

    if q_caja.isdigit():
        movimientos = movimientos.filter(caja_id=int(q_caja))
    if q_libro.isdigit():
        movimientos = movimientos.filter(asiento_libro__libro_id=int(q_libro))
    if q_hoja:
        try:
            movimientos = movimientos.filter(asiento_libro__hoja=int(q_hoja))
        except ValueError:
            pass
    if q_solo_sin_asignar:
        movimientos = movimientos.filter(sin_asignar=0)

    movimientos = aplicar_orden_queryset(
        request, movimientos,
        {
            'id': 'id',
            'caja': 'caja__nombre',
            'receptor': 'receptor__nombre',
            'monto': 'monto',
            'emision': 'emision',
            'numero': 'rel_numero__numero',
            'libro': 'asiento_libro__libro__nombre',
            'hoja': 'asiento_libro__hoja',
            'renglon': 'asiento_libro__renglon',
        },
        default=('sin_asignar', '-emision', '-id'),
    )

    total_count = movimientos.count()
    movimientos = list(movimientos[:300])
    for m in movimientos:
        asiento = getattr(m, 'asiento_libro', None)
        m.libro_mostrar = asiento.libro if asiento else None
        m.hoja_mostrar = asiento.hoja if (asiento and asiento.hoja not in (None, -1)) else None
        m.renglon_mostrar = asiento.renglon if (asiento and asiento.renglon not in (None, -1)) else None

    return render(request, 'movimientos_caja/movimiento_caja_libro_listado.html', {
        'movimientos': movimientos,
        'total_count': total_count,
        'cajas': Caja.objects.all().order_by('nombre'),
        'libros': LibroCaja.objects.select_related('caja').order_by('nombre'),
        'q_caja': q_caja,
        'q_libro': q_libro,
        'q_hoja': q_hoja,
        'q_solo_sin_asignar': q_solo_sin_asignar,
    })


@transaction.atomic
def movimiento_caja_libro_editar(request, pk):
    """Alta/edición del asiento de libro (libro, hoja y renglón) de un
    Movimiento de Caja puntual: la versión, dentro de la app, del admin de
    LibroMovim (equivalente a /admin/movimientos_caja/libromovim/<id>/change/).
    Si el movimiento todavía no tiene asiento de libro, se crea uno al
    guardar.
    """
    movimiento = get_object_or_404(
        MovimientoCaja.objects.select_related('caja', 'receptor', 'asiento_libro', 'asiento_libro__libro'),
        pk=pk,
    )
    asiento = getattr(movimiento, 'asiento_libro', None)
    next_url = _url_next_segura(request, request.POST.get('next') or request.GET.get('next'))

    if request.method == 'POST':
        form = AsignarLibroMovimientoForm(caja_id=movimiento.caja_id, data=request.POST)
        if form.is_valid():
            LibroMovim.objects.update_or_create(
                movimiento_caja=movimiento,
                defaults={
                    'libro': form.cleaned_data['libro'],
                    'hoja': form.cleaned_data['hoja'],
                    'renglon': form.cleaned_data['renglon'],
                },
            )
            messages.success(request, f'Libro del movimiento de caja {movimiento.id} guardado correctamente.')
            return redirect(next_url or 'movimientos_caja:movimiento_caja_libro_modificar')
    else:
        initial = {}
        if asiento:
            initial = {'libro': asiento.libro_id, 'hoja': asiento.hoja, 'renglon': asiento.renglon}
        form = AsignarLibroMovimientoForm(caja_id=movimiento.caja_id, initial=initial)

    return render(request, 'movimientos_caja/movimiento_caja_libro_form.html', {
        'form': form,
        'movimiento': movimiento,
        'asiento': asiento,
        'next': next_url or '',
    })


# ---------------------------------------------------------------------------
# Reportes
# ---------------------------------------------------------------------------

def _movimientos_reporte_filtrados(request):
    """Aplica a MovimientoCaja los filtros de MovimientoCajaReporteForm.

    Centraliza la lógica de filtrado para que la vista de pantalla y las de
    exportación (Excel / PDF) usen siempre exactamente los mismos criterios.
    Devuelve (form, queryset, filtros_activos).
    """
    form = MovimientoCajaReporteForm(request.GET or None)
    movimientos = (
        MovimientoCaja.objects.select_related(
            'caja', 'tipo', 'receptor', 'rel_numero', 'movimientocajadiferido', 'emisor_relacion__id_entidad',
        )
        .prefetch_related('liquidaciones__liquidacion')
        .order_by('-emision', '-id')
    )

    filtros_activos = False

    if form.is_valid():
        caja = form.cleaned_data.get('caja')
        tipo = form.cleaned_data.get('tipo')
        receptor = form.cleaned_data.get('receptor')
        emisor = form.cleaned_data.get('emisor')
        fecha_desde = form.cleaned_data.get('fecha_desde')
        fecha_hasta = form.cleaned_data.get('fecha_hasta')
        efectivizacion_desde = form.cleaned_data.get('efectivizacion_desde')
        efectivizacion_hasta = form.cleaned_data.get('efectivizacion_hasta')
        diferido_desde = form.cleaned_data.get('diferido_desde')
        diferido_hasta = form.cleaned_data.get('diferido_hasta')
        sin_efectivizar = form.cleaned_data.get('sin_efectivizar')

        if caja:
            movimientos = movimientos.filter(caja=caja)
        if tipo:
            movimientos = movimientos.filter(tipo=tipo)
        if receptor:
            movimientos = movimientos.filter(receptor=receptor)
        if emisor:
            # Sólo encuentra movimientos con este emisor cargado explícitamente
            # (tabla movimiento_caja_emisor); los que no tienen carga son
            # Fontana por defecto pero no matchean un filtro por otra entidad.
            movimientos = movimientos.filter(emisor_relacion__id_entidad=emisor)
        if fecha_desde:
            movimientos = movimientos.filter(emision__gte=fecha_desde)
        if fecha_hasta:
            movimientos = movimientos.filter(emision__lte=fecha_hasta)
        if efectivizacion_desde:
            movimientos = movimientos.filter(efectivizacion__gte=efectivizacion_desde)
        if efectivizacion_hasta:
            movimientos = movimientos.filter(efectivizacion__lte=efectivizacion_hasta)
        if diferido_desde:
            movimientos = movimientos.filter(movimientocajadiferido__diferido__gte=diferido_desde)
        if diferido_hasta:
            movimientos = movimientos.filter(movimientocajadiferido__diferido__lte=diferido_hasta)
        if sin_efectivizar:
            movimientos = movimientos.filter(efectivizacion__isnull=True)

        filtros_activos = any([
            caja, tipo, receptor, emisor, fecha_desde, fecha_hasta,
            efectivizacion_desde, efectivizacion_hasta,
            diferido_desde, diferido_hasta, sin_efectivizar,
        ])

    return form, movimientos, filtros_activos


def movimiento_caja_reporte(request):
    form, movimientos, filtros_activos = _movimientos_reporte_filtrados(request)

    totales = movimientos.aggregate(
        total_monto=Coalesce(Sum('monto'), Value(Decimal('0')), output_field=DecimalField(max_digits=20, decimal_places=2)),
    )

    cantidad_total = movimientos.count()
    movimientos = aplicar_orden_queryset(request, movimientos, {
        'id': 'id',
        'caja': 'caja__nombre',
        'tipo': 'tipo__nombre',
        'numero': 'rel_numero__numero',
        'emision': 'emision',
        'emisor': 'emisor_relacion__id_entidad__nombre',
        'receptor': 'receptor__nombre',
        'monto': 'monto',
        'diferido': 'movimientocajadiferido__diferido',
        'efectivizacion': 'efectivizacion',
    })
    movimientos = movimientos[:500]

    receptor_id = form['receptor'].value()
    receptor_texto = texto_entidad_buscador(Entidad.objects.filter(pk=receptor_id).first()) if receptor_id else ''
    emisor_id = form['emisor'].value()
    emisor_texto = texto_entidad_buscador(Entidad.objects.filter(pk=emisor_id).first()) if emisor_id else ''

    return render(request, 'movimientos_caja/movimiento_caja_reporte.html', {
        'form': form,
        'movimientos': movimientos,
        'totales': totales,
        'cantidad_total': cantidad_total,
        'filtros_activos': filtros_activos,
        'receptor_texto': receptor_texto,
        'emisor_texto': emisor_texto,
    })


# --- Exportación a Excel / PDF del reporte -----------------------------------

def _numero_o_none(valor):
    """Convierte Decimal/float a float para que openpyxl no se queje; deja
    pasar None (celda vacía) tal cual."""
    if valor is None:
        return None
    return float(valor)


def _texto_liquidacion(movimiento):
    numeros = [str(lm.liquidacion.id) for lm in movimiento.liquidaciones.all() if lm.liquidacion]
    return ', '.join(numeros) if numeros else 'Sin liquidar'


def _filas_movimiento_caja_reporte(movimientos):
    columnas = ['ID', 'Caja', 'Tipo', 'Número', 'Emisión', 'Diferido', 'Emisor', 'Receptor', 'Monto', 'Efectivización', 'Liquidación']
    filas = []
    for m in movimientos:
        try:
            diferido = m.movimientocajadiferido.diferido
        except Exception:
            diferido = None
        filas.append([
            m.id,
            str(m.caja) if m.caja else '',
            str(m.tipo) if m.tipo else '',
            m.numero if m.numero is not None else '',
            m.emision,
            diferido,
            str(m.emisor) if m.emisor else '',
            str(m.receptor) if m.receptor else '',
            _numero_o_none(m.monto),
            m.efectivizacion,
            _texto_liquidacion(m),
        ])
    return {
        'columnas': columnas,
        'filas': filas,
        'columnas_numericas': {8},  # Monto
        'columnas_fecha': {4, 5, 9},  # Emisión, Diferido, Efectivización
        'anchos': [0.5, 1.1, 1.3, 0.8, 0.9, 0.9, 1.6, 1.8, 1.0, 1.0, 1.3],
    }


def _excel_response(nombre_archivo, resultado):
    import openpyxl
    from openpyxl.utils import get_column_letter

    from services.gestorexcel import definir_estilo_general

    columnas = resultado['columnas']
    filas = resultado['filas']
    columnas_numericas = resultado.get('columnas_numericas', set())
    columnas_fecha = resultado.get('columnas_fecha', set())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Reporte'
    ws.append(columnas)
    fila_encabezado = ws.max_row
    for fila in filas:
        ws.append(fila)
    definir_estilo_general(ws)

    for indice in columnas_numericas:
        letra_columna = get_column_letter(indice + 1)
        for celda in ws[letra_columna]:
            if celda.row > fila_encabezado:
                celda.number_format = '#,##0.00'

    for indice in columnas_fecha:
        letra_columna = get_column_letter(indice + 1)
        for celda in ws[letra_columna]:
            if celda.row > fila_encabezado:
                celda.number_format = 'DD-MM-YYYY'

    for columna in ws.columns:
        letra = columna[0].column_letter
        largo_max = max((len(str(c.value)) for c in columna if c.value is not None), default=10)
        ws.column_dimensions[letra].width = min(max(largo_max + 2, 10), 40)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename={nombre_archivo}.xlsx'
    wb.save(response)
    return response


def _pdf_response(nombre_archivo, titulo, resultado):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    from movimientos.templatetags.movimientos_extras import separador_miles

    columnas = resultado['columnas']
    filas = resultado['filas']
    columnas_numericas = resultado.get('columnas_numericas', set())
    columnas_fecha = resultado.get('columnas_fecha', set())
    anchos_relativos = resultado.get('anchos') or [1] * len(columnas)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename={nombre_archivo}.pdf'

    margen = 1 * cm
    doc = SimpleDocTemplate(
        response, pagesize=landscape(A4),
        topMargin=margen, bottomMargin=margen, leftMargin=margen, rightMargin=margen,
    )
    ancho_disponible = landscape(A4)[0] - doc.leftMargin - doc.rightMargin

    estilos = getSampleStyleSheet()
    estilo_celda = ParagraphStyle('celda_reporte_mc', parent=estilos['Normal'], fontSize=7, leading=8.5)
    estilo_encabezado = ParagraphStyle(
        'encabezado_reporte_mc', parent=estilo_celda, textColor=colors.white, fontName='Helvetica-Bold',
    )

    elementos = [Paragraph(titulo, estilos['Title']), Spacer(1, 0.4 * cm)]

    def formatear_valor(indice, valor):
        if valor is None or valor == '':
            return ''
        if indice in columnas_numericas:
            return separador_miles(valor)
        if indice in columnas_fecha and hasattr(valor, 'strftime'):
            return valor.strftime('%d-%m-%Y')
        return str(valor)

    fila_encabezado = [Paragraph(str(col), estilo_encabezado) for col in columnas]
    datos = [fila_encabezado]
    for fila in filas:
        datos.append([
            Paragraph(formatear_valor(indice, valor), estilo_celda)
            for indice, valor in enumerate(fila)
        ])

    total_relativo = sum(anchos_relativos) or 1
    col_widths = [ancho_disponible * (peso / total_relativo) for peso in anchos_relativos]

    tabla = Table(datos, colWidths=col_widths, repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#343a40')),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]))
    elementos.append(tabla)
    doc.build(elementos)
    return response


def movimiento_caja_reporte_excel(request):
    _form, movimientos, _filtros_activos = _movimientos_reporte_filtrados(request)
    resultado = _filas_movimiento_caja_reporte(movimientos)
    return _excel_response('reporte_movimientos_caja', resultado)


def movimiento_caja_reporte_pdf(request):
    _form, movimientos, _filtros_activos = _movimientos_reporte_filtrados(request)
    resultado = _filas_movimiento_caja_reporte(movimientos)
    return _pdf_response('reporte_movimientos_caja', 'Reporte de movimientos de caja', resultado)


# --- Ranking de entidades (por monto total, sólo filtrando por emisión) -----

# Entidad que representa a la propia empresa (Fontana); mismo id que usa
# liquidaciones.views.ENTIDAD_PROPIA_ID y MovimientoCaja.emisor.
ENTIDAD_PROPIA_ID = 100


def _movimientos_ranking_filtrados(request):
    """Aplica a MovimientoCaja los filtros de RankingEntidadesForm.

    Devuelve (form, queryset, filtros_activos). Centraliza el filtrado para
    que la pantalla y las exportaciones (Excel / PDF) usen siempre los
    mismos criterios.
    """
    form = RankingEntidadesForm(request.GET or None)
    movimientos = MovimientoCaja.objects.filter(receptor__isnull=False)

    filtros_activos = False

    if form.is_valid():
        fecha_desde = form.cleaned_data.get('fecha_desde')
        fecha_hasta = form.cleaned_data.get('fecha_hasta')
        excluir_fontana = form.cleaned_data.get('excluir_fontana')
        if fecha_desde:
            movimientos = movimientos.filter(emision__gte=fecha_desde)
        if fecha_hasta:
            movimientos = movimientos.filter(emision__lte=fecha_hasta)
        if excluir_fontana:
            movimientos = movimientos.exclude(receptor_id=ENTIDAD_PROPIA_ID)
        filtros_activos = bool(fecha_desde or fecha_hasta or excluir_fontana)

    return form, movimientos, filtros_activos


def _calcular_ranking(movimientos):
    """A partir de un queryset de MovimientoCaja, arma el ranking de
    entidades receptoras por monto total (de mayor a menor) y el total
    general. Devuelve (ranking, total_general).
    """
    ranking = list(
        movimientos.values('receptor_id', 'receptor__nombre')
        .annotate(total_monto=Sum('monto'), cantidad=Count('id'))
        .order_by('-total_monto')
    )

    total_general = sum((fila['total_monto'] for fila in ranking), Decimal('0'))
    for posicion, fila in enumerate(ranking, start=1):
        fila['posicion'] = posicion
        fila['porcentaje'] = (fila['total_monto'] / total_general * 100) if total_general else Decimal('0')

    return ranking, total_general


@requiere_grupo('Rankings')
def movimiento_caja_ranking_entidades(request):
    """Ranking de entidades receptoras según la suma de montos de sus
    movimientos de caja, de mayor a menor, filtrando opcionalmente por un
    rango de fecha de emisión (y excluyendo, si se pide, a Fontana).
    """
    form, movimientos, filtros_activos = _movimientos_ranking_filtrados(request)
    ranking, total_general = _calcular_ranking(movimientos)
    ranking = aplicar_orden_lista(request, ranking, {
        'posicion': lambda f: f['posicion'],
        'entidad': lambda f: (f['receptor__nombre'] or '').lower(),
        'cantidad': lambda f: f['cantidad'],
        'monto': lambda f: f['total_monto'],
        'porcentaje': lambda f: f['porcentaje'],
    })

    return render(request, 'movimientos_caja/movimiento_caja_ranking_entidades.html', {
        'form': form,
        'ranking': ranking,
        'total_general': total_general,
        'filtros_activos': filtros_activos,
    })


def _filas_ranking_entidades(ranking):
    columnas = ['#', 'Entidad', 'Movimientos', 'Monto total', 'Participación %']
    filas = [
        [
            fila['posicion'],
            fila['receptor__nombre'] or 'Sin nombre',
            fila['cantidad'],
            _numero_o_none(fila['total_monto']),
            _numero_o_none(fila['porcentaje']),
        ]
        for fila in ranking
    ]
    return {
        'columnas': columnas,
        'filas': filas,
        'columnas_numericas': {3, 4},  # Monto total, Participación %
        'anchos': [0.4, 2.2, 1.0, 1.2, 1.2],
    }


@requiere_grupo('Rankings')
def movimiento_caja_ranking_entidades_excel(request):
    _form, movimientos, _filtros_activos = _movimientos_ranking_filtrados(request)
    ranking, _total_general = _calcular_ranking(movimientos)
    resultado = _filas_ranking_entidades(ranking)
    return _excel_response('ranking_entidades', resultado)


@requiere_grupo('Rankings')
def movimiento_caja_ranking_entidades_pdf(request):
    _form, movimientos, _filtros_activos = _movimientos_ranking_filtrados(request)
    ranking, _total_general = _calcular_ranking(movimientos)
    resultado = _filas_ranking_entidades(ranking)
    return _pdf_response('ranking_entidades', 'Ranking de entidades por monto', resultado)


# ---------------------------------------------------------------------------
# Estado de caja: saldo disponible de una o más cajas a una fecha elegida,
# calculado a partir del saldo inicial del último libro de cada una.
# ---------------------------------------------------------------------------

def _ultimo_libro_de_caja(caja):
    """El libro 'vigente' de una caja: el de fecha_creacion más reciente.
    Los libros que todavía no tienen fecha_creacion cargada (pendientes del
    comando de gestión backfill_fecha_creacion_libros) se ordenan al final,
    y se usa el id como desempate."""
    return (
        LibroCaja.objects.filter(caja=caja)
        .order_by(F('fecha_creacion').desc(nulls_last=True), '-id')
        .first()
    )


def _movimientos_firmes_de_libro(libro, fecha):
    """Movimientos cargados en ESE libro que ya impactan el saldo a la fecha
    dada: los que no tienen diferido (se consideran siempre firmes) o cuyo
    diferido ya llegó (diferido <= fecha)."""
    return MovimientoCaja.objects.filter(asiento_libro__libro=libro).filter(
        Q(movimientocajadiferido__isnull=True)
        | Q(movimientocajadiferido__diferido__isnull=True)
        | Q(movimientocajadiferido__diferido__lte=fecha)
    )


def _calcular_estado_caja(caja, fecha):
    """Arma el estado de una caja a una fecha dada: último libro, su saldo
    inicial, el saldo resultante a esa fecha y, para poder proyectar hacia
    adelante, los movimientos de ese mismo libro que todavía están
    pendientes (con diferido posterior a la fecha elegida), agrupados por
    día. Devuelve un dict lista para el template y para armar la
    exportación."""
    libro = _ultimo_libro_de_caja(caja)
    movimientos_sin_libro = MovimientoCaja.objects.filter(caja=caja, asiento_libro__isnull=True).count()

    if libro is None:
        return {
            'caja': caja,
            'libro': None,
            'saldo_inicial': None,
            'saldo_a_fecha': None,
            'movimientos_sin_libro': movimientos_sin_libro,
            'proyeccion': [],
        }

    saldo_inicial = libro.saldo_inicial if libro.saldo_inicial is not None else Decimal('0')
    total_firme = _movimientos_firmes_de_libro(libro, fecha).aggregate(total=Sum('monto'))['total'] or Decimal('0')
    saldo_a_fecha = saldo_inicial - total_firme

    pendientes = (
        MovimientoCaja.objects.filter(asiento_libro__libro=libro, movimientocajadiferido__diferido__gt=fecha)
        .values('movimientocajadiferido__diferido')
        .annotate(total_dia=Sum('monto'))
        .order_by('movimientocajadiferido__diferido')
    )

    proyeccion = []
    saldo_corriendo = saldo_a_fecha
    for fila in pendientes:
        saldo_corriendo = saldo_corriendo - fila['total_dia']
        proyeccion.append({
            'fecha': fila['movimientocajadiferido__diferido'],
            'monto_dia': fila['total_dia'],
            'saldo': saldo_corriendo,
        })

    return {
        'caja': caja,
        'libro': libro,
        'saldo_inicial': saldo_inicial,
        'saldo_a_fecha': saldo_a_fecha,
        'movimientos_sin_libro': movimientos_sin_libro,
        'proyeccion': proyeccion,
    }


def _datos_estado_caja(request):
    """Lee fecha y cajas de la query string (compartido por la pantalla y
    las exportaciones) y calcula el estado de cada caja seleccionada.
    Devuelve (form, fecha, resultados)."""
    if request.GET:
        form = EstadoCajaForm(request.GET)
        if form.is_valid():
            fecha = form.cleaned_data.get('fecha') or timezone.localdate()
            cajas_seleccionadas = list(form.cleaned_data.get('cajas') or [])
        else:
            fecha = timezone.localdate()
            cajas_seleccionadas = []
    else:
        form = EstadoCajaForm(initial={'fecha': timezone.localdate()})
        fecha = timezone.localdate()
        cajas_seleccionadas = []

    resultados = [_calcular_estado_caja(caja, fecha) for caja in cajas_seleccionadas]
    return form, fecha, resultados


def _filas_proyeccion_pantalla(resultados):
    """Arma las filas de la tabla de proyección para la pantalla: una lista
    de {'fecha':, 'celdas': [...]}, con una celda (o None, si esa caja no
    tuvo movimiento ese día) por cada caja seleccionada, en el mismo orden
    que 'resultados' -- para poder recorrerlas en el template con un simple
    for anidado, sin necesidad de buscar por fecha ahí."""
    proyeccion_por_caja = [{p['fecha']: p for p in r['proyeccion']} for r in resultados]
    fechas_futuras = sorted({p['fecha'] for r in resultados for p in r['proyeccion']})

    filas = []
    for dia in fechas_futuras:
        celdas = [proyeccion_por_caja[indice].get(dia) for indice in range(len(resultados))]
        filas.append({'fecha': dia, 'celdas': celdas})
    return filas


@requiere_grupo('Rankings')
def movimiento_caja_estado(request):
    form, fecha, resultados = _datos_estado_caja(request)
    calculado = bool(request.GET) and form.is_valid() and bool(resultados)

    total_saldo = sum(
        (r['saldo_a_fecha'] for r in resultados if r['saldo_a_fecha'] is not None), Decimal('0'),
    )

    return render(request, 'movimientos_caja/movimiento_caja_estado.html', {
        'form': form,
        'fecha': fecha,
        'resultados': resultados,
        'calculado': calculado,
        'total_saldo': total_saldo,
        'sin_seleccion': bool(request.GET) and form.is_valid() and not resultados,
        'filas_proyeccion': _filas_proyeccion_pantalla(resultados) if calculado else [],
    })


def _tabla_estado_caja(fecha, resultados):
    """Arma una única tabla (para Excel/PDF) con la fecha compartida en la
    primera columna y, por cada caja seleccionada, un par de columnas
    Movimiento/Saldo: una fila con el saldo a la fecha elegida y, después,
    una fila por cada día posterior que tenga algún movimiento diferido
    pendiente en alguna de las cajas (parecido a la planilla de referencia,
    con la fecha compartida y un par Movimiento/Saldo por cuenta)."""
    columnas = ['Fecha']
    for r in resultados:
        nombre = str(r['caja'])
        columnas.append(f'{nombre} - Movimiento')
        columnas.append(f'{nombre} - Saldo')

    fila_inicial = [fecha]
    saldo_corriente = []
    for r in resultados:
        fila_inicial.append(None)
        saldo_a_fecha = r['saldo_a_fecha']
        fila_inicial.append(_numero_o_none(saldo_a_fecha) if saldo_a_fecha is not None else None)
        saldo_corriente.append(saldo_a_fecha)

    filas = [fila_inicial]

    proyeccion_por_caja = [{p['fecha']: p for p in r['proyeccion']} for r in resultados]
    fechas_futuras = sorted({p['fecha'] for r in resultados for p in r['proyeccion']})

    for dia in fechas_futuras:
        fila = [dia]
        for indice in range(len(resultados)):
            entrada = proyeccion_por_caja[indice].get(dia)
            if entrada:
                saldo_corriente[indice] = entrada['saldo']
                fila.append(_numero_o_none(-entrada['monto_dia']))
                fila.append(_numero_o_none(entrada['saldo']))
            else:
                saldo_actual = saldo_corriente[indice]
                fila.append(None)
                fila.append(_numero_o_none(saldo_actual) if saldo_actual is not None else None)
        filas.append(fila)

    columnas_numericas = set(range(1, len(columnas)))
    return {
        'columnas': columnas,
        'filas': filas,
        'columnas_numericas': columnas_numericas,
        'columnas_fecha': {0},
        'anchos': [0.9] + [1.1] * (len(columnas) - 1),
    }


@requiere_grupo('Rankings')
def movimiento_caja_estado_excel(request):
    _form, fecha, resultados = _datos_estado_caja(request)
    resultado = _tabla_estado_caja(fecha, resultados)
    return _excel_response('estado_de_caja', resultado)


@requiere_grupo('Rankings')
def movimiento_caja_estado_pdf(request):
    _form, fecha, resultados = _datos_estado_caja(request)
    resultado = _tabla_estado_caja(fecha, resultados)
    return _pdf_response('estado_de_caja', 'Estado de caja', resultado)
